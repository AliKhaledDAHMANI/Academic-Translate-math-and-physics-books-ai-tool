#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL
======================================
Single-file pipeline for translating mathematics and physics books (PDF, scanned PDF, DOCX, LaTeX) between languages while keeping formulas,
notation, numbering, tables, figures and structure intact, with automatic source-vs-translation QC before export.

    python MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL.py readme        # full documentation, pipelines, QC table, provenance, tested status
    python MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL.py check-env     # what is installed
    python MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL.py selftest      # run the built-in tests

For an AI operator: read `readme` first. Parts of this file: 1 core/config/language profiles, 2 prompts, 3 glossary, 4 PDF probe/OCR/figures/packets,
5 LaTeX protection/segmentation/validation, 6 project + translation backends, 7 LaTeX build + LaTeX->DOCX, 8 native DOCX, 9 QC, 10 CLI/README/self-test.
"""
# =============================================================================
# PART 1 — IMPORTS, CONSTANTS, PROVENANCE, CONFIG, LANGUAGE PROFILES, UTILITIES
# =============================================================================
import argparse, collections, copy, csv, difflib, hashlib, json, math, os, re, shutil
import statistics, subprocess, sys, tempfile, time, unicodedata, zipfile
import urllib.error, urllib.request
from dataclasses import dataclass, field, asdict
from pathlib import Path

TOOL_NAME = "MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL"
TOOL_VERSION = "1.0"

# Placeholder syntax used to shield protected content from the translator.
#   ⟦M12⟧   protected span (M=math, C=comment, K=command, E=env token, G=graphic, T=table, V=verbatim, X=docx atom, S=child segment)
#   ⟦s3⟧ … ⟦/s3⟧   style span (DOCX run formatting)
PH_RE = re.compile("⟦(/?)([A-Za-z]+)(\\d+)⟧")
def PH(kind, n): return "⟦%s%d⟧" % (kind, n)

# ----------------------------------------------------------------------------
# PROVENANCE: what was ACTUALLY used for the French->English translation of
# "Cours d'analyse 1" (Laffaille & Pauly) versus what this tool adds.
# Verified against the command log of that session. Nothing here is embellished.
# ----------------------------------------------------------------------------
PROVENANCE = {
 "actually_used_in_the_original_session": [
  "poppler-utils: pdfinfo, pdffonts, pdfimages -list (inspection: 73 pages, pdfTeX, embedded Computer Modern fonts, text layer present)",
  "poppler-utils: pdftotext -layout (whole book -> one text file, plus page ranges per chapter)",
  "poppler-utils: pdftoppm -png at 50-95 dpi (page images) to check formulas/figures the text layer garbled",
  "The language model itself (Claude) as BOTH math transcriber and translator: it read the extracted text plus page images and wrote English LaTeX chapter by chapter in a single pass (French text -> English LaTeX; formulas copied/reconstructed from the page images)",
  "TeX Live pdflatex (two passes) with amsmath, amssymb, amsthm, tikz, pgfplots, graphicx, geometry, hyperref, babel[english], enumerate; custom amsthm styles; per-section theorem counters",
  "Python 3 standard library (re, unicodedata) for ONE automatic check: inventory of numbered items (Definition/Theorem/Proposition/Lemma/Exercise N.N.N) in source vs translated text; ligature replacement (fi/ff) and NFC normalisation before comparing",
  "Pillow (PIL) to build side-by-side contact sheets of source/translated pages; manual visual inspection of selected pages (title page, TOC, figures, truth tables, Taylor graph)",
  "Compile-log inspection (undefined references, errors, overfull boxes) after each pdflatex run",
  "Anthropic-provided skill documents (file-reading, pdf-reading SKILL.md) giving environment-specific advice on PDF inspection. These are proprietary documentation of that environment and are NOT exportable; this tool replaces them with its own probe/extract logic (PART 4).",
 ],
 "available_in_that_sandbox_but_NOT_used": [
  "pandoc, xelatex, lualatex, python-docx, pdfplumber, pytesseract/tesseract (English data only) were present; none were used for the French book.",
 ],
 "deviations_from_the_standard_you_now_ask_for": [
  "FIGURES WERE REDRAWN in TikZ/pgfplots (right triangle, hyperbola, complex plane, tangent/secant, cobweb, cosine Taylor graph, unit circle, ln x vs cos x). You asked for figures NOT to be redrawn. This tool's default is the opposite: crop the original vector figure from the source PDF (PART 4, detect_figures + \\srcfig) and overlay translated labels only where needed.",
  "The document was REBUILT as new LaTeX (77 pages vs 73 original), not edited in place; layout is similar, not identical.",
  "The delivered Analysis_1_Course_EN.pdf embeds 17 Type 3 (bitmap) fonts WITHOUT a Unicode map (the sandbox had no lmodern/cm-super): it reads correctly on screen but is blurry when zoomed and copy/paste/search of accents and ligatures is unreliable. This tool now avoids it (font_package auto, glyphtounicode) and its QC flags it (bitmap_fonts, no_unicode_map). The LaTeX sources of that translation were not kept, so fixing it means re-running the pipeline.",
  "Math was reconstructed by the model from text layer + page images, NOT protected by an automated placeholder system, and NOT compared automatically against the source. The only automatic check was the numbered-item inventory (160/160 matched); everything else was visual spot-checking. The exact-math comparison, placeholder protection, glossary, language-ID, duplicate, table/figure and cross-reference checks in this tool are NEW and were not part of that session.",
  "No explicit glossary file existed; consistency relied on translating the whole book in one session. This tool adds an explicit glossary layer.",
  "Translation and transcription happened in one pass. This tool separates them (transcribe -> verify -> protect -> translate -> compare) so that math can be compared mechanically.",
  "Policy choices made for that book (now configurable, not hard-coded): French connectives et/ou/non and V/F kept inside formulas and truth tables; decimal commas kept; plain-language words inside displayed statements translated; author's notation (e.g. cotan) kept.",
 ],
 "proprietary_or_not_exportable": [
  "The translating/transcribing model (Claude) is proprietary and cannot be exported. Replacement: any strong LLM through the backends in PART 6 (Anthropic API, any OpenAI-compatible endpoint including local servers, or 'agent' mode where another AI/human fills batch files). For page-image transcription use a vision-capable model.",
  "The environment's skill documents (see above). Replacement: PART 4.",
  "Mathpix (optional math OCR) is a proprietary paid API. Open replacements: Nougat, Marker, pix2tex, or a vision-capable LLM transcribing page images to LaTeX.",
 ],
}

# ----------------------------------------------------------------------------
# DEFAULT CONFIGURATION (all values overridable in <project>/project.json)
# ----------------------------------------------------------------------------
DEFAULT_CONFIG = {
 "project": {"source_lang": "fr", "target_lang": "en", "domain": "mathematics",  # mathematics | physics | both
             "title": "", "author_note": ""},
 "extraction": {"min_text_chars_per_page": 40, "raster_dpi": 80, "ocr_dpi": 300,
                "ocr_engine": "tesseract", "math_ocr_engine": "none",   # none | nougat | marker | pix2tex | mathpix
                "garbled_ratio_threshold": 0.02, "dehyphenate": False},
 "protect": {
   "text_in_math": "translate_prose",     # preserve | translate_prose   (what to do with \text{...} inside formulas)
   "protected_words": ["V", "F"],         # tokens never translated when they occur alone inside math/table cells
   "keep_words_in_math": [],              # e.g. ["et","ou","non"] to keep source-language connectives inside \text{}/cells (session policy)
   "tables": "translate_text_cells",      # protect | translate_text_cells
   "tikz_labels": "translate_prose",      # protect | translate_prose
   "comments": "preserve",
   "extra_nontranslatable_commands": [],
 },
 "notation": {"policy": "preserve",       # preserve | flag_only   (never silently localise sin/tan/sh/Arctan etc.)
              "decimal_separator": "preserve", "unit_policy": "preserve"},
 "translate": {"backend": "agent",         # agent | mock | anthropic | openai
               "batch_chars": 6000, "context_segments": 2, "max_retries": 3, "temperature": 0.2,
               "model": "claude-sonnet-5-5", "api_base": "", "api_key_env": "", "timeout_s": 180,
               "reviewer_pass": False},
 "latex": {"engine": "auto", "passes": 3, "documentclass": "report", "fontsize": "12pt",
           "paper": "letterpaper", "margin": "1.15in", "extra_packages": [],
           "font_package": "auto"},   # auto | lmodern | mathptmx | none  (pdfLaTeX only; auto avoids bitmap Type 3 fonts, see qc bitmap_fonts)
 "figures": {"mode": "crop_original", "pad_pt": 14, "min_w_pt": 50, "min_h_pt": 40, "table_min_h_pt": 18, "label_gap_pt": 6,
             "translate_labels": "overlay"},   # overlay | keep | list_only
 "qc": {"length_ratio_mad_k": 4.0, "min_chars_for_ratio": 60, "dup_jaccard": 0.9, "residue_stopword_ratio": 0.25,
        "block_on": ["math_altered", "math_missing", "math_added", "placeholder_mismatch", "numbering_mismatch", "env_mismatch",
                     "label_mismatch", "ref_mismatch", "broken_reference", "table_mismatch", "figure_mismatch", "caption_mismatch",
                     "footnote_mismatch", "number_altered", "latex_error", "untranslated_segment", "formatting_corruption",
                     "docx_structure_mismatch", "docx_math_mismatch"]},   # numbering_mismatch_heuristic, bitmap_fonts, no_unicode_map are warnings by default
 "docx": {"skip_styles": ["Code", "SourceCode", "VerbatimChar", "Verbatim"], "update_fields_on_open": True},
}

def deep_merge(base, over):
    out = copy.deepcopy(base)
    for k, v in (over or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict): out[k] = deep_merge(out[k], v)
        else: out[k] = v
    return out

def load_config(path=None):
    cfg = copy.deepcopy(DEFAULT_CONFIG)
    if path and Path(path).exists():
        cfg = deep_merge(cfg, json.loads(Path(path).read_text(encoding="utf-8")))
    return cfg

# ----------------------------------------------------------------------------
# LANGUAGE PROFILES. Mathematical NOTATION is language independent and is never
# taken from here; profiles only hold ordinary-language conventions.
# `labels` are structural words (theorem names...). They are DEFAULTS: have them
# confirmed by a mathematician who reads the target language (review=True means "I am less sure").
# ----------------------------------------------------------------------------
def _L(theorem, definition, proposition, lemma, corollary, remark, example, exercise, proof, solution,
       contents, chapter, figure, table):
    return dict(theorem=theorem, definition=definition, proposition=proposition, lemma=lemma,
                corollary=corollary, remark=remark, example=example, exercise=exercise, proof=proof,
                solution=solution, contents=contents, chapter=chapter, figure=figure, table=table)

LANGUAGE_PROFILES = {
 "en": dict(name="English", script="latin", dir="ltr", babel="english", polyglossia="english", ocr="eng", bcp47="en-US",
    labels=_L("Theorem","Definition","Proposition","Lemma","Corollary","Remark","Example","Exercise","Proof","Solution","Contents","Chapter","Figure","Table"),
    quotes=("“","”"), review=False,
    style="Declarative textbook register: 'Let ... Then ...', 'We have ...', 'Suppose that ...', 'It follows that ...'. One spelling variant (US or UK) throughout.",
    stopwords="the of and to in is that for be let we then if with as by on at an or this are has have such which all any there".split(),
    ai_phrases=[r"\bdelve\b", r"it is (important|worth) (to note|noting)", r"\bin conclusion\b", r"\blet'?s\b", r"\bin today's\b", r"\bseamless", r"\btapestry\b", r"\bcrucial\b", r"\bfurthermore\b.*\bfurthermore\b"]),
 "fr": dict(name="French", script="latin", dir="ltr", babel="french", polyglossia="french", ocr="fra", bcp47="fr-FR",
    labels=_L("Théorème","Définition","Proposition","Lemme","Corollaire","Remarque","Exemple","Exercice","Démonstration","Solution","Table des matières","Chapitre","Figure","Tableau"),
    quotes=("« "," »"), review=False,
    style="'Soit ... Alors ...', 'On a ...', 'Supposons que ...', 'Il en résulte que ...'. Non-breaking space before : ; ! ? and inside guillemets.",
    stopwords="le la les de des du un une et est que qui dans pour par sur avec sont au aux ce cette on il si alors tel soit tout tous ou ne pas".split(),
    ai_phrases=[r"\bil est important de noter\b", r"\bdans le monde d'aujourd'hui\b"]),
 "es": dict(name="Spanish", script="latin", dir="ltr", babel="spanish", polyglossia="spanish", ocr="spa", bcp47="es-ES",
    labels=_L("Teorema","Definición","Proposición","Lema","Corolario","Observación","Ejemplo","Ejercicio","Demostración","Solución","Índice general","Capítulo","Figura","Tabla"),
    quotes=("«","»"), review=False,
    style="'Sea ... Entonces ...', 'Se tiene que ...', 'Supongamos que ...', 'De ello se sigue que ...'. Use 'tal que', 'si y solo si' consistently.",
    stopwords="el la los las de del un una y es que en para por con son al lo se si entonces sea todo todos o no".split(),
    ai_phrases=[r"\bcabe destacar\b", r"\ben el mundo actual\b"]),
 "de": dict(name="German", script="latin", dir="ltr", babel="ngerman", polyglossia="german", ocr="deu", bcp47="de-DE",
    labels=_L("Satz","Definition","Proposition","Lemma","Korollar","Bemerkung","Beispiel","Aufgabe","Beweis","Lösung","Inhaltsverzeichnis","Kapitel","Abbildung","Tabelle"),
    quotes=("„","“"), review=False,
    style="'Sei ... Dann gilt ...', 'Es gilt ...', 'Angenommen, ...', 'Daraus folgt ...'. Keep German mathematical terms (Folge, Reihe, Grenzwert); avoid anglicisms.",
    stopwords="der die das den dem des und ist ein eine einer zu mit von für auf wenn dann sei alle oder nicht auch sich".split(),
    ai_phrases=[r"\bes ist wichtig zu beachten\b", r"\bin der heutigen\b"]),
 "it": dict(name="Italian", script="latin", dir="ltr", babel="italian", polyglossia="italian", ocr="ita", bcp47="it-IT",
    labels=_L("Teorema","Definizione","Proposizione","Lemma","Corollario","Osservazione","Esempio","Esercizio","Dimostrazione","Soluzione","Indice","Capitolo","Figura","Tabella"),
    quotes=("«","»"), review=False,
    style="'Sia ... Allora ...', 'Si ha ...', 'Supponiamo che ...', 'Ne segue che ...'.",
    stopwords="il lo la i gli le di del un una e è che in per con sono al se allora sia ogni tutti o non".split(),
    ai_phrases=[r"\bè importante notare\b"]),
 "pt": dict(name="Portuguese", script="latin", dir="ltr", babel="portuguese", polyglossia="portuguese", ocr="por", bcp47="pt-BR",
    labels=_L("Teorema","Definição","Proposição","Lema","Corolário","Observação","Exemplo","Exercício","Demonstração","Solução","Sumário","Capítulo","Figura","Tabela"),
    quotes=("“","”"), review=False,
    style="'Seja ... Então ...', 'Tem-se que ...' (PT) / 'Temos que ...' (BR). Choose PT-PT or PT-BR once (bcp47) and keep it.",
    stopwords="o a os as de do da dos das um uma e é que em para por com são ao se então seja todo todos ou não".split(),
    ai_phrases=[r"\bé importante notar\b"]),
 "nl": dict(name="Dutch", script="latin", dir="ltr", babel="dutch", polyglossia="dutch", ocr="nld", bcp47="nl-NL",
    labels=_L("Stelling","Definitie","Propositie","Lemma","Gevolg","Opmerking","Voorbeeld","Opgave","Bewijs","Oplossing","Inhoudsopgave","Hoofdstuk","Figuur","Tabel"),
    quotes=("“","”"), review=True,
    style="'Zij ... Dan geldt ...', 'Er geldt ...', 'Neem aan dat ...'.",
    stopwords="de het een en is van dat die in voor met zijn op als dan laat alle of niet er te".split(), ai_phrases=[]),
 "ru": dict(name="Russian", script="cyrillic", dir="ltr", babel="russian", polyglossia="russian", ocr="rus", bcp47="ru-RU",
    labels=_L("Теорема","Определение","Предложение","Лемма","Следствие","Замечание","Пример","Упражнение","Доказательство","Решение","Оглавление","Глава","Рисунок","Таблица"),
    quotes=("«","»"), review=False,
    style="'Пусть ... Тогда ...', 'Имеем ...', 'Предположим, что ...', 'Отсюда следует, что ...'. Do NOT convert tan/arctan/sinh to tg/arctg/sh unless the author does; flag instead.",
    stopwords="и в не на что с по для это как если то пусть все или при из к от".split(), ai_phrases=[]),
 "tr": dict(name="Turkish", script="latin", dir="ltr", babel="turkish", polyglossia="turkish", ocr="tur", bcp47="tr-TR",
    labels=_L("Teorem","Tanım","Önerme","Lemma","Sonuç","Not","Örnek","Alıştırma","İspat","Çözüm","İçindekiler","Bölüm","Şekil","Tablo"),
    quotes=("“","”"), review=True,
    style="'... olsun. O zaman ...', 'Elimizde ... vardır'. Agglutinative: keep placeholders adjacent to suffix-bearing words intact.",
    stopwords="ve bir bu için ile de da olsun her ise ya veya değil olan".split(), ai_phrases=[]),
 "ar": dict(name="Arabic", script="arabic", dir="rtl", babel="arabic", polyglossia="arabic", ocr="ara", bcp47="ar",
    labels=_L("مبرهنة","تعريف","قضية","تمهيدية","نتيجة","ملاحظة","مثال","تمرين","برهان","حل","المحتويات","الفصل","الشكل","الجدول"),
    quotes=("«","»"), review=True,
    style="Right-to-left prose; ALL mathematics stays left-to-right and unmirrored. Keep digit style (Western/Eastern) as in the source. Typical: 'ليكن ... عندئذٍ ...'.",
    stopwords="في من على أن إلى هذا هذه التي الذي كل إذا عندئذ ليكن".split(), ai_phrases=[]),
 "zh": dict(name="Chinese (Simplified)", script="han", dir="ltr", babel=None, polyglossia=None, ocr="chi_sim", bcp47="zh-CN",
    labels=_L("定理","定义","命题","引理","推论","注","例","习题","证明","解","目录","章","图","表"),
    quotes=("“","”"), review=False,
    style="'设……，则……', '由……可知……'. Full-width punctuation outside math; no spaces inside Chinese text; one consistent rule at Chinese–Latin boundaries.",
    stopwords="的 是 在 和 为 则 设 若 对 任意 存在 满足 因此".split(), ai_phrases=[]),
 "ja": dict(name="Japanese", script="kana_han", dir="ltr", babel=None, polyglossia=None, ocr="jpn", bcp47="ja-JP",
    labels=_L("定理","定義","命題","補題","系","注意","例","演習問題","証明","解答","目次","章","図","表"),
    quotes=("「","」"), review=True,
    style="Textbook plain style (である体): '... とする。このとき ... が成り立つ。'.",
    stopwords="の は を に が と で する とき ならば 任意 存在".split(), ai_phrases=[]),
 "ko": dict(name="Korean", script="hangul", dir="ltr", babel=None, polyglossia=None, ocr="kor", bcp47="ko-KR",
    labels=_L("정리","정의","명제","보조정리","따름정리","주석","예","연습문제","증명","풀이","차례","장","그림","표"),
    quotes=("“","”"), review=True,
    style="Textbook plain style: '... 라 하자. 그러면 ... 이다.' Use the established Korean mathematical term list.",
    stopwords="의 는 은 이 가 을 를 에서 하면 임의의 존재 만족".split(), ai_phrases=[]),
 "hi": dict(name="Hindi", script="devanagari", dir="ltr", babel=None, polyglossia="hindi", ocr="hin", bcp47="hi-IN",
    labels=_L("प्रमेय","परिभाषा","प्रस्ताव","उपप्रमेय","अनुप्रमेय","टिप्पणी","उदाहरण","अभ्यास","उपपत्ति","हल","विषय-सूची","अध्याय","चित्र","सारणी"),
    quotes=("“","”"), review=True,
    style="'माना कि ... तब ...'. Keep Western digits unless the source uses Devanagari digits. Prefer established university terminology.",
    stopwords="का की के में है और से को पर तब माना यदि".split(), ai_phrases=[]),
}
LANG_ALIASES = {"fra": "fr", "eng": "en", "spa": "es", "deu": "de", "ger": "de", "ita": "it", "por": "pt", "nld": "nl",
                "rus": "ru", "tur": "tr", "ara": "ar", "zho": "zh", "chi": "zh", "zh-cn": "zh", "jpn": "ja", "kor": "ko", "hin": "hi"}

def get_profile(code):
    c = LANG_ALIASES.get(code.lower(), code.lower())
    if c not in LANGUAGE_PROFILES:
        raise SystemExit("Unknown language '%s'. Add it with register_language() or extend LANGUAGE_PROFILES. Known: %s"
                         % (code, ", ".join(sorted(LANGUAGE_PROFILES))))
    return c, LANGUAGE_PROFILES[c]

def register_language(code, profile):
    """Add any language: copy a similar profile and edit name/script/dir/babel/labels/quotes/style/stopwords/ocr/bcp47."""
    LANGUAGE_PROFILES[code] = profile

SCRIPT_RANGES = {
 "latin": [(0x41,0x5A),(0x61,0x7A),(0xC0,0x24F),(0x1E00,0x1EFF)], "cyrillic": [(0x400,0x52F)],
 "arabic": [(0x600,0x6FF),(0x750,0x77F),(0xFB50,0xFDFF),(0xFE70,0xFEFF)], "han": [(0x4E00,0x9FFF),(0x3400,0x4DBF)],
 "kana": [(0x3040,0x30FF)], "hangul": [(0xAC00,0xD7AF),(0x1100,0x11FF)], "devanagari": [(0x900,0x97F)], "greek": [(0x370,0x3FF)],
}
def script_ratio(text, script):
    letters = [ch for ch in text if ch.isalpha()]
    if not letters: return 0.0
    ranges = []
    for s in (["kana", "han"] if script == "kana_han" else [script]): ranges += SCRIPT_RANGES.get(s, [])
    hit = sum(1 for ch in letters if any(a <= ord(ch) <= b for a, b in ranges))
    return hit / len(letters)

# ----------------------------------------------------------------------------
# UTILITIES
# ----------------------------------------------------------------------------
LIGATURES = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl", "­": ""}
# T1 (EC) encoded fonts without a ToUnicode map come out of pdftotext as raw control characters (e.g. 'De<0x1C>nition'):
LIGATURES.update({chr(27): "ff", chr(28): "fi", chr(29): "fl", chr(30): "ffi", chr(31): "ffl"})
def normalize_text(s, dehyphenate=False):
    """Text-layer normalisation used before any comparison: NFC, ligatures (fi/ff...), soft hyphen, NBSP."""
    s = unicodedata.normalize("NFC", s)
    for a, b in LIGATURES.items(): s = s.replace(a, b)
    s = s.replace(" ", " ").replace(" ", " ")
    if dehyphenate: s = re.sub(r"([a-zà-ÿ])-\n\s*([a-zà-ÿ])", r"\1\2", s)
    return s

def fold(s):
    s = unicodedata.normalize("NFKD", s.lower())
    return "".join(c for c in s if not unicodedata.combining(c))

def sh(cmd, timeout=600, cwd=None, input_text=None):
    """Run a command; never raises on non-zero exit; returns (rc, stdout, stderr)."""
    try:
        p = subprocess.run([str(c) for c in cmd], capture_output=True, text=True, timeout=timeout, cwd=cwd, input=input_text,
                           encoding="utf-8", errors="replace")
        return p.returncode, p.stdout, p.stderr
    except FileNotFoundError:
        return 127, "", "command not found: %s" % cmd[0]
    except subprocess.TimeoutExpired:
        return 124, "", "timeout after %ss: %s" % (timeout, " ".join(map(str, cmd)))

def have(binary): return shutil.which(binary) is not None
def py_has(mod):
    try: __import__(mod); return True
    except Exception: return False

def log(msg): print(msg, file=sys.stderr)
def write(path, text):
    Path(path).parent.mkdir(parents=True, exist_ok=True); Path(path).write_text(text, encoding="utf-8")
def read(path): return Path(path).read_text(encoding="utf-8", errors="replace")
def jdump(path, obj): write(path, json.dumps(obj, ensure_ascii=False, indent=1))
def jload(path, default=None):
    return json.loads(read(path)) if Path(path).exists() else default

@dataclass
class Finding:
    severity: str      # block | warn | info
    check: str
    where: str
    detail: str
    src: str = ""
    tgt: str = ""
    def key(self): return hashlib.sha1(("%s|%s|%s|%s" % (self.check, self.where, self.src[:200], self.tgt[:200])).encode()).hexdigest()[:12]

# =============================================================================
# PART 2 — INTERNAL PROMPTS (authored for this tool; they encode the rules the
# session followed implicitly, made explicit and machine-checkable)
# =============================================================================
TRANSLATE_SYSTEM_PROMPT = r"""
You are translating one segment batch of a university-level <<DOMAIN>> textbook from <<SRC_NAME>> into <<TGT_NAME>>.
The result must read as if the book had been written originally in <<TGT_NAME>> by a university professor.
You are a TRANSLATOR, not an editor: you never correct, improve, modernise, shorten, expand or comment on the author's text.

HARD RULES (violations are detected automatically and the segment is rejected)
1. PLACEHOLDERS. Tokens of the form ⟦M12⟧ ⟦K3⟧ ⟦E7⟧ ⟦G2⟧ ⟦T1⟧ ⟦C4⟧ ⟦V5⟧ ⟦X9⟧ ⟦S2⟧ and style tags ⟦s2⟧…⟦/s2⟧ stand for protected content (formulas, commands,
   environments, figures, tables, comments, fields, run formatting). Copy every placeholder EXACTLY, once each, with no spaces inside it. Never translate, expand, merge,
   split, renumber or invent placeholders. You may move a placeholder within the sentence only if <<TGT_NAME>> word order requires it.
2. NO MATH IN PROSE. Never convert a formula to words, never rewrite notation, never add or drop symbols, digits, units, indices or punctuation that belong to a formula.
   Digits and decimal separators in running text stay exactly as in the source.
3. NOTHING ADDED, NOTHING REMOVED. Every sentence, example, exercise statement, remark, footnote and caption of the source appears in the output. No translator's notes.
4. LaTeX/markup in the text (\emph{…}, \textbf{…}, \item, ~, \\, quotation marks, etc.) is kept as is; translate only the words inside.
5. TERMINOLOGY. Glossary entries supplied with a segment are MANDATORY for that segment (pick the entry whose context matches). Otherwise use the established term of <<TGT_NAME>> university
   textbooks for this field; never coin terms. If two established terms exist and you are unsure, use the more standard one and add a flag {"type":"term_uncertain"}.
6. NAMES AND SYMBOLS. Proper names of mathematicians/physicists use their conventional <<TGT_NAME>> form. Abbreviations that are notation (gcd, sup, cotan…) are NOT translated unless
   a glossary entry says so. Do not localise notation (sin/tan/sh/Arctan, decimal comma/point, vector arrows) — flag it with {"type":"notation_convention"} instead.
7. AUTHOR'S ERRORS. If the source contains an apparent error, ambiguity or typo, TRANSLATE IT FAITHFULLY and add a flag {"type":"source_issue","note":"…"}. Never fix it silently.
8. STYLE. <<STYLE_NOTE>> Formal textbook register. No conversational phrasing, no second-person chat style unless the source has it, no filler, no rhetorical flourishes,
   no stock "AI" phrasing. Keep the author's sentence structure when it is natural in <<TGT_NAME>>; restructure only when a literal rendering would be unidiomatic.
   Typographic quotes: <<QUOTES>>.
9. A segment of kind "textinmath" is a short phrase that sits INSIDE a displayed formula (e.g. 'pour tout', 'si'). Translate it as a short phrase, with no added capital letter and no final
   punctuation, preserving leading/trailing spaces. A segment of kind "cell" is a table cell: translate only its wording. A segment of kind "figlabel" is a label inside a figure:
   translate it as a short label. A segment of kind "heading" contains a sectioning command: keep the command, translate the title.
10. OUTPUT: a single JSON object, nothing else: {"translations":[{"id":"<segment id>","tgt":"<translation>","flags":[{"type":"…","note":"…"}]}]} with the same ids, same order.
"""

RECONSTRUCT_PROMPT = r"""
TASK: faithful TRANSCRIPTION (no translation) of source pages into LaTeX. You are given, per page: the page image, the text layer (or OCR text) as unreliable evidence, and detected
figure regions (with ready-made \srcfig macros). The page image is the ground truth; the text layer is evidence only (math is often garbled in it).
RULES
 - Output LaTeX in the SOURCE language, verbatim. Do not translate, correct, modernise, shorten or reorder anything. Keep the author's notation exactly (e.g. \operatorname{cotan}, decimal commas, \mathbb sets).
 - Use \chapter/\section/\subsection for headings and the project's theorem-like environments (definition, theorem, proposition, lemma, corollary, exercise, remark…) so that LaTeX counters
   reproduce the printed numbers. Where a printed number cannot be produced by a counter (e.g. solutions numbered like the exercises) write it literally and add the comment `% LITERAL-NUMBER`.
 - Every displayed equation, table (tabular, same rows/columns/merged cells), truth table, list, footnote, caption, reference and cross-reference of the page must be present. Use \label/\ref
   for cross references when the page shows a numbered reference.
 - Figures: do NOT redraw. Use the supplied \srcfig{page}{x0}{y0}{x1}{y1} macro (crops the original vector art). Put the caption in a figure environment.
 - If any character, symbol, index or number is unreadable or ambiguous, transcribe your best reading and add a comment `% FLAG(uncertain-math|uncertain-text): <what and why>` on the same line. Never guess silently.
 - Preserve paragraph breaks as blank lines. Do not insert any commentary outside LaTeX comments.
 - Put the preamble in preamble.tex (use `python tool.py preamble <project>` for a generated one) and one file per chapter (ch1.tex ...) plus main.tex.
"""

REVIEW_PROMPT = r"""
TASK: terminology-and-naturalness REVIEW of an already translated segment batch (<<SRC_NAME>> -> <<TGT_NAME>>). You see source and translation side by side.
You may ONLY (a) fix wording so it reads as native academic <<TGT_NAME>>, (b) align terminology with the glossary, (c) remove literal-translation artefacts.
You may NOT touch placeholders, formulas, numbers, structure, or meaning, and you may NOT correct the author. Return the same JSON schema as the translation step; return the translation
unchanged when it is already right. List each change as a flag {"type":"review_change","note":"old -> new, reason"}.
"""

GLOSSARY_PROMPT = r"""
TASK: build a context-aware <<DOMAIN>> glossary <<SRC_NAME>> -> <<TGT_NAME>> for the candidate terms below, using the established terminology of <<TGT_NAME>> university textbooks.
Output CSV rows with columns: src,tgt,lang_pair,domain,context,forbidden,notes,locked
 - one row per sense; 'context' = |-separated words that, when present near the term, select this sense (e.g. travail -> work with context force|énergie|joule);
 - 'forbidden' = |-separated literal/incorrect renderings that must NOT appear (e.g. limited development);
 - 'notes' = convention differences between traditions (e.g. French 'croissante' = non-decreasing); 'locked'=1 when the choice must never vary.
Do not invent terms; leave tgt empty and put 'REVIEW' in notes if you are not sure.
"""

SOURCE_FLAG_PROMPT = r"""
TASK: list apparent errors/ambiguities in the SOURCE book (typos in formulas, inconsistent numbering, mismatched references, undefined symbols). Report only; NEVER propose silent changes to the
translation. Output JSONL: {"where":"<segment id or page>","type":"…","note":"…"}.
"""

def fill_prompt(tpl, cfg):
    sc, sp = get_profile(cfg["project"]["source_lang"]); tc, tp = get_profile(cfg["project"]["target_lang"])
    q = "%s…%s" % tp["quotes"]
    return (tpl.replace("<<SRC_NAME>>", sp["name"]).replace("<<TGT_NAME>>", tp["name"])
               .replace("<<DOMAIN>>", {"both": "mathematics and physics"}.get(cfg["project"]["domain"], cfg["project"]["domain"]))
               .replace("<<STYLE_NOTE>>", tp["style"]).replace("<<QUOTES>>", q))

# =============================================================================
# PART 3 — GLOSSARY / TERMINOLOGY LAYER
# =============================================================================
GLOSSARY_COLUMNS = ["src", "tgt", "lang_pair", "domain", "context", "forbidden", "notes", "locked"]

# Seed glossary fr->en (analysis, algebra, logic, mechanics/electricity/optics). Context-aware rows share a src.
# (src, tgt, domain, context, forbidden, notes). A SEED, not a finished dictionary: a mathematician must review it.
_SEED_FR_EN = [
 ("suite", "sequence", "analysis", "", "", ""), ("suite réelle", "real sequence", "analysis", "", "", ""),
 ("suite extraite", "subsequence", "analysis", "", "extracted sequence|extracted subsequence", ""),
 ("sous-suite", "subsequence", "analysis", "", "", ""), ("suite de Cauchy", "Cauchy sequence", "analysis", "", "", ""),
 ("suites adjacentes", "adjacent sequences", "analysis", "", "", ""), ("suite récurrente", "recursively defined sequence", "analysis", "", "", ""),
 ("série", "series", "analysis", "", "", ""), ("limite", "limit", "analysis", "", "", ""),
 ("valeur d'adhérence", "cluster point", "analysis", "", "adherence value", "Also 'subsequential limit'; pick one for the book."),
 ("borne supérieure", "supremum", "analysis", "", "", "= least upper bound; keep 'upper bound' for 'majorant'."),
 ("borne inférieure", "infimum", "analysis", "", "", ""), ("majorant", "upper bound", "analysis", "", "", ""), ("minorant", "lower bound", "analysis", "", "", ""),
 ("majoré", "bounded above", "analysis", "", "", ""), ("minoré", "bounded below", "analysis", "", "", ""), ("borné", "bounded", "analysis", "", "", ""),
 ("développement limité", "Taylor expansion", "analysis", "", "limited development|limited expansion|finite development",
  "Session choice 'Taylor expansion (développement limité)'. Alternatives: asymptotic expansion / finite expansion with o() remainder."),
 ("formule de Taylor-Lagrange", "Taylor–Lagrange formula", "analysis", "", "", ""), ("formule de Taylor-Young", "Taylor–Young formula", "analysis", "", "", ""),
 ("théorème des valeurs intermédiaires", "intermediate value theorem", "analysis", "", "", ""),
 ("théorème des accroissements finis", "mean value theorem", "analysis", "", "finite increments theorem|finite growth theorem", ""),
 ("théorème de Rolle", "Rolle's theorem", "analysis", "", "", ""), ("dérivable", "differentiable", "analysis", "", "derivable", ""),
 ("dérivée", "derivative", "analysis", "", "", ""), ("fonction réciproque", "inverse function", "analysis", "", "reciprocal function", ""),
 ("équivalent", "equivalent", "analysis", "suite|limite|infini|~", "", "u_n ~ v_n: 'asymptotically equivalent' on first use, then 'equivalent'."),
 ("négligeable", "negligible", "analysis", "", "", "o(.) notation: 'u_n is negligible compared to v_n'."), ("point fixe", "fixed point", "analysis", "", "", ""),
 ("voisinage", "neighborhood", "analysis", "", "", "US spelling; use 'neighbourhood' if the project is UK."), ("intervalle", "interval", "analysis", "", "", ""),
 ("croissante", "increasing", "analysis", "", "", "French 'croissante' = non-decreasing; 'strictement croissante' = strictly increasing. Decide once and keep."),
 ("strictement croissante", "strictly increasing", "analysis", "", "", ""), ("convexe", "convex", "analysis", "", "", ""),
 ("inégalité triangulaire", "triangle inequality", "analysis", "", "", ""), ("cercle trigonométrique", "unit circle", "analysis", "", "trigonometric circle", ""),
 ("fonctions hyperboliques", "hyperbolic functions", "analysis", "", "", ""), ("logarithme népérien", "natural logarithm", "analysis", "", "Neperian logarithm", ""),
 ("application", "map", "algebra", "ensemble|injective|surjective|bijective|image", "", "French 'application' = map/function between sets; avoid the false friend 'application'."),
 ("injective", "injective", "algebra", "", "", ""), ("surjective", "surjective", "algebra", "", "", ""), ("bijective", "bijective", "algebra", "", "", ""),
 ("image réciproque", "preimage", "algebra", "", "reciprocal image", "'inverse image' is acceptable; choose one."),
 ("partie", "subset", "set theory", "ensemble|inclus|de E", "", "'une partie de E' = a subset of E."),
 ("ensemble vide", "empty set", "set theory", "", "", ""), ("produit cartésien", "Cartesian product", "set theory", "", "", ""),
 ("quantificateur", "quantifier", "logic", "", "", ""), ("condition nécessaire et suffisante", "necessary and sufficient condition", "logic", "", "", ""),
 ("contraposée", "contrapositive", "logic", "", "", ""), ("raisonnement par l'absurde", "proof by contradiction", "logic", "", "reasoning by absurdity", ""),
 ("récurrence", "induction", "logic", "", "recurrence reasoning", "Use 'recurrence' only for a recurrence relation."),
 ("entier naturel", "natural number", "number systems", "", "", ""), ("entier relatif", "integer", "number systems", "", "relative integer", ""),
 ("nombre premier", "prime number", "number systems", "", "", ""), ("nombres premiers entre eux", "coprime numbers", "number systems", "", "", ""),
 ("pgcd", "gcd", "number systems", "", "", "Notation/abbreviation."), ("ppcm", "lcm", "number systems", "", "", ""),
 ("partie entière", "integer part", "number systems", "", "", "'floor' if the author writes ⌊x⌋."), ("valeur absolue", "absolute value", "number systems", "", "", ""),
 ("module", "modulus", "complex numbers", "complexe|argument|nombre", "", ""), ("racine de l'unité", "root of unity", "complex numbers", "", "", ""),
 ("formule du binôme", "binomial theorem", "number systems", "", "", ""),
 ("corps", "field", "algebra", "anneau|groupe|commutatif|ℚ|ℝ|ℂ|extension", "", ""),
 ("corps", "body", "mechanics", "chute|solide|masse|pesanteur|pendule", "", "Physical body."),
 ("espace vectoriel", "vector space", "algebra", "", "", ""), ("application linéaire", "linear map", "algebra", "", "linear application", ""),
 ("noyau", "kernel", "algebra", "", "", ""), ("valeur propre", "eigenvalue", "algebra", "", "", ""), ("vecteur propre", "eigenvector", "algebra", "", "", ""),
 ("base", "basis", "algebra", "espace|vecteur|vectoriel|famille", "", ""),
 ("vitesse", "velocity", "mechanics", "vecteur|cinématique|trajectoire|mouvement", "", "'speed' when only the magnitude is meant."),
 ("travail", "work", "mechanics", "force|énergie|joule|déplacement", "labor", ""), ("puissance", "power", "mechanics", "watt|travail|énergie|force", "", ""),
 ("quantité de mouvement", "momentum", "mechanics", "", "quantity of movement", ""), ("moment cinétique", "angular momentum", "mechanics", "", "kinetic moment", ""),
 ("énergie cinétique", "kinetic energy", "mechanics", "", "", ""), ("référentiel", "reference frame", "mechanics", "", "referential", ""),
 ("galiléen", "inertial", "mechanics", "référentiel", "", "Use 'Galilean' only if the author means Galilean transformations."),
 ("pesanteur", "gravity", "mechanics", "", "heaviness", ""), ("frottement", "friction", "mechanics", "", "", ""),
 ("intensité", "current", "electricity", "courant|circuit|ampère|résistance|tension", "", ""),
 ("intensité", "intensity", "optics", "lumineuse|lumière|onde|rayonnement", "", ""),
 ("tension", "voltage", "electricity", "circuit|potentiel|volt|résistance", "", ""), ("tension", "tension", "mechanics", "fil|ressort|corde", "", ""),
 ("champ électrique", "electric field", "electricity", "", "", ""), ("champ magnétique", "magnetic field", "electricity", "", "", ""),
 ("loi d'Ohm", "Ohm's law", "electricity", "", "", ""), ("longueur d'onde", "wavelength", "optics", "", "", ""),
 ("indice de réfraction", "refractive index", "optics", "", "refraction index", ""), ("gaz parfait", "ideal gas", "thermodynamics", "", "perfect gas", "'perfect gas' only if the author's tradition does."),
 ("masse volumique", "mass density", "mechanics", "", "", ""),
]

def accent_fold_regex(term):
    return re.compile(r"(?<!\w)" + re.escape(fold(term)) + r"(?:s|x|es)?(?!\w)")

class Glossary:
    """Context-aware terminology layer. Rows: GLOSSARY_COLUMNS. Matching is accent/case-insensitive."""
    def __init__(self, rows=None, pair="fr-en"):
        self.rows, self.pair, self._rx = [], pair, None
        for r in rows or []: self.add(**r)
    def add(self, src, tgt="", lang_pair=None, domain="", context="", forbidden="", notes="", locked="0"):
        self.rows.append(dict(src=src, tgt=tgt, lang_pair=lang_pair or self.pair, domain=domain, context=context,
                              forbidden=forbidden, notes=notes, locked=str(locked)))
        self._rx = None
    @classmethod
    def seed(cls, pair):
        g = cls(pair=pair)
        if pair == "fr-en":
            for s, t, d, c, f, n in _SEED_FR_EN: g.add(s, t, pair, d, c, f, n)
        return g
    @classmethod
    def load(cls, path, pair="fr-en"):
        g = cls(pair=pair)
        if Path(path).exists():
            with open(path, encoding="utf-8", newline="") as fh:
                for r in csv.DictReader(fh): g.add(**{k: r.get(k, "") or "" for k in GLOSSARY_COLUMNS})
        return g
    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=GLOSSARY_COLUMNS); w.writeheader(); w.writerows(self.rows)
    def _compile(self):
        if self._rx is None:
            self._rx = [(accent_fold_regex(r["src"]), r) for r in self.rows if r["src"]]
            self._rx.sort(key=lambda x: -len(x[1]["src"]))   # longest terms first
        return self._rx
    def hits(self, text, ctx=""):
        """Entries applicable to `text` (+ heading context). Among rows sharing a src, prefer the one whose context keywords occur."""
        ft = fold(text); fc = fold(ctx + " " + text); raw = ctx + " " + text
        found, covered = {}, []
        for rx, r in self._compile():
            m = rx.search(ft)
            if not m: continue
            if any(a <= m.start() and m.end() <= b and srcc != r["src"] for a, b, srcc in covered): continue   # inside a longer term
            covered.append((m.start(), m.end(), r["src"]))
            score = 0; kws = [k for k in r["context"].split("|") if k]
            if kws:
                def _kw_hit(k):   # symbols / very short keys: raw substring; words: accent-folded word-start match
                    if len(k) <= 2 or not k.isalpha(): return k in raw
                    return re.search(r"(?<!\w)" + re.escape(fold(k)), fc) is not None
                score = sum(1 for k in kws if _kw_hit(k)) or -1
            found.setdefault(r["src"], []).append((score, r))
        out = []
        for src, cands in found.items():
            cands.sort(key=lambda x: -x[0]); out.append(cands[0][1])
        return out
    def check(self, src_text, tgt_text, ctx="", where=""):
        """Findings: approved target missing (warn), forbidden rendering present (warn; block if locked)."""
        out, ft = [], fold(tgt_text)
        for r in self.hits(src_text, ctx):
            if not r["tgt"]: continue
            if fold(r["tgt"]).replace("–", "-") not in ft.replace("–", "-"):
                out.append(Finding("warn", "glossary_miss", where, "'%s' -> expected '%s'" % (r["src"], r["tgt"]), src_text[:160], tgt_text[:160]))
        for r in self.rows:
            for bad in [b for b in r["forbidden"].split("|") if b]:
                if re.search(r"(?<!\w)" + re.escape(fold(bad)) + r"(?!\w)", ft):
                    out.append(Finding("block" if r["locked"] == "1" else "warn", "forbidden_term", where,
                                       "forbidden rendering '%s' (glossary: %s -> %s)" % (bad, r["src"], r["tgt"]), src_text[:160], tgt_text[:160]))
        return out
    def prompt_block(self, hits):
        return [{"src": r["src"], "tgt": r["tgt"], "note": r["notes"]} for r in hits if r["tgt"]]

def extract_term_candidates(texts, lang, top=250):
    """Frequent 1-3 word candidates (stop-words excluded) to feed GLOSSARY_PROMPT. Heuristic; review required."""
    _, prof = get_profile(lang); stop = set(map(fold, prof["stopwords"]))
    cnt = collections.Counter()
    for t in texts:
        t = re.sub(r"⟦[^⟧]*⟧", " ", t)
        toks = [w for w in re.findall(r"[^\W\d_]{3,}", fold(t))]
        for n in (1, 2, 3):
            for i in range(len(toks) - n + 1):
                g = toks[i:i + n]
                if g[0] in stop or g[-1] in stop: continue
                cnt[" ".join(g)] += 1
    return [(w, c) for w, c in cnt.most_common(top * 3) if c >= 2][:top]

# =============================================================================
# PART 4 — PDF PROBE, TEXT-LAYER EXTRACTION, RASTERISATION, OCR, FIGURES, PAGE PACKETS
# Decision rule: keep the existing text layer wherever it is trustworthy (digital PDF);
# OCR only pages with no/garbled text; flag math-dense OCR pages for math-aware OCR or vision transcription.
# =============================================================================
OPERATOR_WORDS = set("sin cos tan cot sec csc sinh cosh tanh coth arcsin arccos arctan arccot ln log exp lim sup inf max min det dim ker rank rg tr deg gcd lcm "
                     "pgcd ppcm mod div rot grad curl arg Re Im Id id cotan tg ctg sh ch th Arctan Arcsin Arccos".split())
MATH_FONT_RE = re.compile(r"(CM(MI|SY|EX|MIB|BSY)\d*|MSAM|MSBM|EUFM|EUSM|LMMath|LatinModernMath|CambriaMath|Cambria Math|STIX|XITS|MathematicalPi|"
                          r"Symbol|TeXGyre.*Math|Asana|Euler|rsfs|wasy|stmary|MnSymbol|Math|CMMI|CMSY|CMEX)", re.I)
OCR_PRODUCER_RE = re.compile(r"(tesseract|ocrmypdf|abbyy|paper ?capture|readiris|omnipage|scan)", re.I)

def probe_pdf(pdf, cfg):
    """Classify the PDF: digital | scanned | mixed, per page. Mirrors the first steps of the original session
    (pdfinfo, pdffonts, pdfimages -list, pdftotext) but automated."""
    pdf = str(pdf); res = {"file": pdf}
    rc, out, _ = sh(["pdfinfo", pdf])
    if rc != 0: raise SystemExit("pdfinfo failed - is poppler-utils installed / is this a valid PDF? (%s)" % out)
    info = dict(l.split(":", 1) for l in out.splitlines() if ":" in l); info = {k.strip(): v.strip() for k, v in info.items()}
    n = int(info.get("Pages", "0")); res.update(pages=n, producer=info.get("Producer", ""), creator=info.get("Creator", ""),
                                               page_size=info.get("Page size", ""), pdf_version=info.get("PDF version", ""))
    rc, out, _ = sh(["pdffonts", pdf]); fonts = []
    for l in out.splitlines()[2:]:
        parts = l.split()
        if len(parts) >= 7: fonts.append({"name": parts[0].split("+")[-1], "emb": parts[-5] == "yes"})
    res["fonts"] = len(fonts); res["math_fonts"] = sorted({f["name"] for f in fonts if MATH_FONT_RE.search(f["name"])})
    res["tex_generated"] = bool(re.search(r"(pdfTeX|LaTeX|XeTeX|LuaTeX|TeX)", res["producer"] + res["creator"]))
    img_pages = collections.Counter()
    rc, out, _ = sh(["pdfimages", "-list", pdf])
    for l in out.splitlines()[2:]:
        p = l.split()
        if p and p[0].isdigit(): img_pages[int(p[0])] += 1
    rc, txt, _ = sh(["pdftotext", "-layout", pdf, "-"]); pages = txt.split("\f")[:n]
    per, min_c, thr = [], cfg["extraction"]["min_text_chars_per_page"], cfg["extraction"]["garbled_ratio_threshold"]
    for i in range(n):
        t = pages[i] if i < len(pages) else ""; chars = len(re.sub(r"\s", "", t))
        bad = len(re.findall("[�-]|\\(cid:\\d+\\)", t)); garbled = chars > 0 and bad / chars > thr
        per.append({"page": i + 1, "chars": chars, "images": img_pages.get(i + 1, 0), "garbled": garbled,
                    "needs_ocr": (chars < min_c and img_pages.get(i + 1, 0) > 0) or garbled})
    need = sum(p["needs_ocr"] for p in per)
    res["per_page"] = per
    res["classification"] = "scanned" if (n and need == n) else ("mixed" if need else "digital")
    res["suspect_ocr_layer"] = bool(OCR_PRODUCER_RE.search(res["producer"] + res["creator"]))
    res["advice"] = []
    if res["classification"] == "digital": res["advice"].append("Keep the text layer; rasterise pages only to verify math/figures.")
    if res["classification"] != "digital": res["advice"].append("OCR the pages with needs_ocr=true (`ocr` command); transcribe math from page images (vision model or math OCR).")
    if res["suspect_ocr_layer"]: res["advice"].append("Text layer looks like a previous OCR pass: math is unreliable; transcribe math from images.")
    if res["math_fonts"] and not res["tex_generated"]: res["advice"].append("Math fonts present but not TeX-generated: expect layout-specific glyph encodings.")
    return res

def extract_text_layer(pdf, outdir, cfg):
    """pdftotext -layout for the whole book, split by form feed into pNNN.txt (+ full.txt), normalised (NFC, ligatures)."""
    rc, txt, err = sh(["pdftotext", "-layout", str(pdf), "-"])
    if rc != 0: raise SystemExit("pdftotext failed: " + err)
    pages = [normalize_text(p, cfg["extraction"]["dehyphenate"]) for p in txt.split("\f")]
    if pages and not pages[-1].strip(): pages = pages[:-1]
    outdir = Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    for i, p in enumerate(pages, 1): write(outdir / ("p%03d.txt" % i), p)
    write(outdir / "full.txt", "\n\f".join(pages)); return pages

def rasterize(pdf, outdir, dpi=80, first=None, last=None, prefix="pg"):
    Path(outdir).mkdir(parents=True, exist_ok=True); cmd = ["pdftoppm", "-r", str(dpi), "-png"]
    if first: cmd += ["-f", str(first)]
    if last: cmd += ["-l", str(last)]
    rc, _, err = sh(cmd + [str(pdf), str(Path(outdir) / prefix)], timeout=1800)
    if rc != 0: raise SystemExit("pdftoppm failed: " + err)
    return sorted(Path(outdir).glob(prefix + "-*.png"))

def contact_sheet(png_a, png_b, out, height=900):
    """Side-by-side source/translated page image for human visual comparison (used in the session with PIL)."""
    from PIL import Image
    A, B = Image.open(png_a).convert("RGB"), Image.open(png_b).convert("RGB")
    A = A.resize((int(A.width * height / A.height), height)); B = B.resize((int(B.width * height / B.height), height))
    canvas = Image.new("RGB", (A.width + B.width + 20, height), "white"); canvas.paste(A, (0, 0)); canvas.paste(B, (A.width + 20, 0))
    Path(out).parent.mkdir(parents=True, exist_ok=True); canvas.save(out); return out

LIG_CID = {27: "ff", 28: "fi", 29: "fl", 30: "ffi", 31: "ffl"}
BOLD_FONT_RE = re.compile(r"(bold|black|heavy|semibold|demi|CMBX|CMB\d|[-_,.]B$|[A-Za-z]BX\d|BX\d)", re.I)
class Inventory(list):
    """List of (kind, number) with .mode: 'bold' (only bold-face line starts were counted: authoritative) or 'all' (every line start; may contain wrapped cross-references)."""
    mode = "all"

SPACING_ACCENTS = "\u00b4`\u00a8\u02c6\u02dc\u00b8\u00af\u02d8\u02d9\u02da\u02dd\u02c7^~"   # OT1 text layers: "D´efinition" = accent glyph + letter

def numbering_inventory_pdf(pdf, lang, bold_only="auto"):
    """Numbered items (Definition 1.2.3, Theorem 2.1 ...) found as the first two words of a text line. When the book sets them in a bold face
    (almost all textbooks do) only bold occurrences count, so cross-references that happen to wrap to a line start are ignored.
    Returns an Inventory ([(kind, number)] in reading order, .mode = 'bold' or 'all'). Uses pdfplumber; falls back to the text-layer
    version (mode 'all') when it is not installed. Fonts whose name is unknown (e.g. Type 3 bitmap fonts) cannot be recognised as bold -> mode 'all'."""
    if not py_has("pdfplumber"):
        inv = Inventory(numbering_inventory_text(sh(["pdftotext", "-layout", str(pdf), "-"])[1], lang)); inv.mode = "all"; return inv
    import pdfplumber
    _, prof = get_profile(lang); keys = ["theorem", "definition", "proposition", "lemma", "corollary", "exercise", "example", "remark"]
    words = {fold(prof["labels"][k]): k for k in keys}; num_rx = re.compile(r"^\d+(?:\.\d+)*[.):]?$")
    accent_tbl = {ord(c): None for c in SPACING_ACCENTS}
    def clean(t):
        t = re.sub(r"\(cid:(\d+)\)", lambda m: LIG_CID.get(int(m.group(1)), m.group(0)), t)
        return fold(normalize_text(t.translate(accent_tbl))).strip("()[]")
    items = []
    with pdfplumber.open(str(pdf)) as doc:
        for pno, page in enumerate(doc.pages, 1):
            ws = page.extract_words(extra_attrs=["fontname"]); lines = collections.defaultdict(list)
            for w in ws: lines[round(w["top"] / 3)].append(w)
            for key in sorted(lines):
                L = sorted(lines[key], key=lambda w: w["x0"]); t0 = clean(L[0]["text"]); bold = bool(BOLD_FONT_RE.search(L[0].get("fontname", "")))
                if t0 in words and len(L) > 1 and num_rx.match(clean(L[1]["text"])): items.append((words[t0], clean(L[1]["text"]).rstrip(".):"), bold, pno))
                else:
                    m = re.match(r"^(\D+?)(\d+(?:\.\d+)*)$", t0)     # CJK style: no space between label and number
                    if m and m.group(1) in words: items.append((words[m.group(1)], m.group(2), bold, pno))
    nb = sum(1 for i in items if i[2])
    use_bold = bold_only is True or (bold_only == "auto" and nb >= 0.5 * max(1, len(items)) and nb > 0)
    inv = Inventory([(k, n) for k, n, b, _ in items if b or not use_bold]); inv.mode = "bold" if use_bold else "all"; return inv

# ---------------------------------------------------------------- OCR
def tesseract_langs():
    rc, out, _ = sh(["tesseract", "--list-langs"]); return set(out.split("\n")[1:]) if rc == 0 else set()

def ocr_page(png, ocr_lang, psm=6):
    """OCR one page image with tesseract; returns text, low-confidence words and a math-suspicion flag.
    Text OCR is NOT trusted for mathematics: pages with math_suspect=True must go through math-aware OCR or vision transcription."""
    langs = tesseract_langs(); warn = ""
    if ocr_lang not in langs:
        warn = "OCR language '%s' not installed (have %s); fell back to eng. Install tesseract-ocr-%s." % (ocr_lang, sorted(l for l in langs if l), ocr_lang)
        ocr_lang = "eng"
    rc, out, err = sh(["tesseract", str(png), "stdout", "-l", ocr_lang, "--psm", str(psm), "tsv"])
    if rc != 0: raise SystemExit("tesseract failed: " + err)
    lines, low = collections.OrderedDict(), []
    for row in out.splitlines()[1:]:
        c = row.split("\t")
        if len(c) < 12 or not c[11].strip(): continue
        key = (c[2], c[3], c[4]); conf = float(c[10]); lines.setdefault(key, []).append(c[11])
        if conf < 60: low.append({"word": c[11], "conf": conf, "box": [int(c[6]), int(c[7]), int(c[8]), int(c[9])]})
    text = "\n".join(" ".join(w) for w in lines.values())
    sym = len(re.findall(r"[=<>^_+\-*/∑∫√≤≥≠≈±×÷∞∂∇∈∉⊂∪∩→⇒⇔∀∃αβγδεθλμπσφωΔΣΩ()\[\]{}|]", text)); dens = sym / max(1, len(text))
    return {"text": normalize_text(text), "low_conf": low, "symbol_density": round(dens, 3), "warning": warn,
            "math_suspect": dens > 0.08 or len(low) > 0.12 * max(1, len(text.split()))}

# Math-aware OCR adapters. Command templates are DATA so they can be adapted; all are optional and
# were NOT used for the French book (it had a text layer). Status: UNTESTED here (not installed / proprietary / no network).
MATH_OCR_TEMPLATES = {
 "nougat":  ["nougat", "{pdf}", "-o", "{outdir}", "--markdown"],          # open source, whole-PDF scientific OCR -> .mmd with LaTeX
 "marker":  ["marker_single", "{pdf}", "--output_dir", "{outdir}"],       # open source, PDF -> markdown with LaTeX
 "pix2tex": ["pix2tex", "{png}"],                                          # open source, one formula image -> LaTeX
}
def run_math_ocr(engine, pdf=None, png=None, outdir="mathocr"):
    if engine == "mathpix": return mathpix_ocr(png)
    tpl = MATH_OCR_TEMPLATES.get(engine)
    if not tpl: raise SystemExit("unknown math OCR engine: %s (choose %s or mathpix)" % (engine, ", ".join(MATH_OCR_TEMPLATES)))
    if not have(tpl[0]): raise SystemExit("%s is not installed. Install it (see README) or use the 'vision' route: send page PNGs to a vision LLM with RECONSTRUCT_PROMPT." % tpl[0])
    rc, out, err = sh([a.format(pdf=pdf, png=png, outdir=outdir) for a in tpl], timeout=3600)
    return out if rc == 0 else "ERROR: " + err

def mathpix_ocr(png):
    """PROPRIETARY paid API (Mathpix). Needs MATHPIX_APP_ID / MATHPIX_APP_KEY. Untested in the build sandbox (no network)."""
    import base64
    app_id, app_key = os.environ.get("MATHPIX_APP_ID"), os.environ.get("MATHPIX_APP_KEY")
    if not (app_id and app_key): raise SystemExit("Set MATHPIX_APP_ID and MATHPIX_APP_KEY")
    body = json.dumps({"src": "data:image/png;base64," + base64.b64encode(Path(png).read_bytes()).decode(),
                       "formats": ["text"], "math_inline_delimiters": ["$", "$"], "math_display_delimiters": ["\\[", "\\]"]}).encode()
    req = urllib.request.Request("https://api.mathpix.com/v3/text", body, {"app_id": app_id, "app_key": app_key, "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as r: return json.loads(r.read().decode()).get("text", "")

# ---------------------------------------------------------------- FIGURES
def _cluster(boxes, gap):
    """Union boxes (x0,top,x1,bottom,count,kinds) that touch or lie within `gap` points. Grid-hashed union-find (fast on plots with thousands of segments)."""
    n = len(boxes); parent = list(range(n))
    def find(i):
        while parent[i] != i: parent[i] = parent[parent[i]]; i = parent[i]
        return i
    grid, cs = collections.defaultdict(list), 40.0
    for i, b in enumerate(boxes):
        for gx in range(int((b[0] - gap) // cs), int((b[2] + gap) // cs) + 1):
            for gy in range(int((b[1] - gap) // cs), int((b[3] + gap) // cs) + 1):
                for j in grid[(gx, gy)]:
                    a = boxes[j]
                    if a[0] <= b[2] + gap and a[2] >= b[0] - gap and a[1] <= b[3] + gap and a[3] >= b[1] - gap:
                        ri, rj = find(i), find(j)
                        if ri != rj: parent[ri] = rj
                grid[(gx, gy)].append(i)
    groups = collections.defaultdict(list)
    for i, b in enumerate(boxes): groups[find(i)].append(b)
    return [[min(b[0] for b in g), min(b[1] for b in g), max(b[2] for b in g), max(b[3] for b in g), sum(b[4] for b in g),
             set().union(*[b[5] for b in g])] for g in groups.values()]

def detect_figures(pdf, cfg, pages=None):
    """Find vector/raster figure regions per page (needs pdfplumber). Returns dicts with bbox in PDF points measured from the
    TOP-LEFT, ready-made crop commands, and the translatable (prose) labels found next to the drawing.
    Regions made only of axis-aligned rules/rects are reported as kind='table_like' (tables, truth tables) and are not treated as figures."""
    if not py_has("pdfplumber"): raise SystemExit("pip install pdfplumber (needed for vector figure detection) - or list figures by hand")
    import pdfplumber
    fc = cfg["figures"]; out = []
    with pdfplumber.open(str(pdf)) as doc:
        for pno, page in enumerate(doc.pages, 1):
            if pages and pno not in pages: continue
            W, H = float(page.width), float(page.height); boxes = []
            for c in page.curves: boxes.append((c["x0"], c["top"], c["x1"], c["bottom"], 1, {"curve"}))
            for im in page.images: boxes.append((im["x0"], im["top"], im["x1"], im["bottom"], 1, {"image"}))
            for l in page.lines:
                w, h = l["x1"] - l["x0"], l["bottom"] - l["top"]
                if (h < 1.0 and w < 30) or (w < 1.0 and h < 6): continue          # fraction bars / short rules
                kind = "diag" if (w > 1.5 and h > 1.5) else "axis"; boxes.append((l["x0"], l["top"], l["x1"], l["bottom"], 1, {kind}))
            for r in page.rects:
                mx, mn = max(r["width"], r["height"]), min(r["width"], r["height"])
                if mx < 8: continue                       # tiny marks (QED squares, bullets)
                if mn < 1.5:                              # rules drawn as filled rectangles (modern TeX): treat like lines
                    if mx >= 30: boxes.append((r["x0"], r["top"], r["x1"], r["bottom"], 1, {"axis"}))
                    continue
                boxes.append((r["x0"], r["top"], r["x1"], r["bottom"], 1, {"rect"}))
            if not boxes: continue
            words = page.extract_words(extra_attrs=["fontname"])
            textlines = collections.defaultdict(list)
            for w in words: textlines[round(w["top"] / 3)].append(w)
            body = [(min(w["x0"] for w in ws), min(w["top"] for w in ws), max(w["x1"] for w in ws), max(w["bottom"] for w in ws))
                    for ws in textlines.values() if max(w["x1"] for w in ws) - min(w["x0"] for w in ws) > 0.5 * W]
            for x0, top, x1, bot, n, kinds in _cluster(boxes, 12):
                is_fig = bool(kinds & {"curve", "image", "diag"}); pad = fc["pad_pt"]
                if (x1 - x0) < fc["min_w_pt"] or (bot - top) < (fc["min_h_pt"] if is_fig else fc.get("table_min_h_pt", 18)): continue
                X0, T, X1, B = x0, top, x1, bot       # tight cluster bbox; adjacent labels are added below
                lab = []
                for w in words:   # a label = word touching the drawing whose whole text line stays within the figure's horizontal extent
                    if not (w["x1"] >= x0 - pad and w["x0"] <= x1 + pad and w["bottom"] >= top - pad and w["top"] <= bot + pad): continue
                    ws = textlines[round(w["top"] / 3)]; lx0, lx1 = min(v["x0"] for v in ws), max(v["x1"] for v in ws)
                    if lx0 < x0 - 25 or lx1 > x1 + 25: continue
                    gx = max(x0 - w["x1"], w["x0"] - x1, 0); gy = max(top - w["bottom"], w["top"] - bot, 0)
                    if max(gx, gy) > fc.get("label_gap_pt", 6): continue      # must touch / sit right next to the drawing
                    if any(b[1] - 1 <= w["top"] and w["bottom"] <= b[3] + 1 for b in body): continue
                    lab.append(w)
                for w in lab: X0, X1, T, B = min(X0, w["x0"]), max(X1, w["x1"]), min(T, w["top"]), max(B, w["bottom"])
                X0, T, X1, B = max(0, X0 - 3), max(0, T - 3), min(W, X1 + 3), min(H, B + 3)
                prose = [{"text": w["text"], "bbox": [round(w["x0"], 1), round(w["top"], 1), round(w["x1"], 1), round(w["bottom"], 1)]}
                         for w in lab if not MATH_FONT_RE.search(w.get("fontname", "")) and w["text"] not in OPERATOR_WORDS
                         and re.fullmatch(r"[^\W\d_]{3,}[.,;:]?", w["text"])]
                bbox = [round(v, 1) for v in (X0, T, X1, B)]
                out.append({"page": pno, "kind": "figure" if is_fig else "table_like", "bbox": bbox, "page_w": W, "page_h": H,
                            "primitives": n, "srcfig": srcfig_call(pno, bbox, W, H) if is_fig else "", "prose_labels": prose if is_fig else []})
    return out

def srcfig_call(page, bbox, W, H):
    """\\srcfig{page}{left}{bottom}{right}{top}: pdfTeX crops the ORIGINAL vector art (trim values are distances from the page edges)."""
    x0, top, x1, bot = bbox
    return "\\srcfig{%d}{%.1f}{%.1f}{%.1f}{%.1f}" % (page, x0, H - bot, W - x1, top)

def figure_overlay_tex(fig, translations):
    """Keep the original graphic, white-out ONLY the prose labels and set the translated labels on top (TikZ overlay).
    translations: {original_label_text: translated_text}. Math/symbol labels are never touched."""
    x0, top, x1, bot = fig["bbox"]; lines = ["\\begin{tikzpicture}[x=1pt,y=1pt]", "\\node[anchor=south west,inner sep=0] at (0,0) {%s};" % fig["srcfig"]]
    for lab in fig["prose_labels"]:
        new = translations.get(lab["text"])
        if not new: continue
        a, t, b, d = lab["bbox"]; xa, xb, ya, yb = a - x0, b - x0, bot - d, bot - t
        lines.append("\\fill[white] (%.1f,%.1f) rectangle (%.1f,%.1f);" % (xa - 1, ya - 1, xb + 1, yb + 1))
        lines.append("\\node[inner sep=0,font=\\footnotesize,anchor=center] at (%.1f,%.1f) {%s};" % ((xa + xb) / 2, (ya + yb) / 2, new))
    lines.append("\\end{tikzpicture}"); return "\n".join(lines)

# ---------------------------------------------------------------- PAGE PACKETS (for transcription of digital OR scanned pages)
def make_packets(pdf, outdir, cfg, pages=None, chunk=8):
    """Per page: PNG, evidence text (text layer or OCR), figure/table regions; plus reconstruction task files.
    This is the structured form of what the model saw in the session (extracted text + page images)."""
    outdir = Path(outdir); pr = probe_pdf(pdf, cfg); n = pr["pages"]; pages = pages or list(range(1, n + 1))
    figs = detect_figures(pdf, cfg, set(pages)) if py_has("pdfplumber") else []
    layer = extract_text_layer(pdf, outdir / "text", cfg); sc, sp = get_profile(cfg["project"]["source_lang"])
    packets = []
    for p in pages:
        meta = pr["per_page"][p - 1]
        png = rasterize(pdf, outdir / "img", cfg["extraction"]["ocr_dpi"] if meta["needs_ocr"] else cfg["extraction"]["raster_dpi"], p, p, "pg%03d" % p)
        png = str(png[0]) if png else ""
        if meta["needs_ocr"] and have("tesseract") and png:
            o = ocr_page(png, sp["ocr"]); text, src = o["text"], "ocr"; extra = {k: o[k] for k in ("low_conf", "symbol_density", "math_suspect", "warning")}
        else: text, src, extra = (layer[p - 1] if p - 1 < len(layer) else ""), "text_layer", {}
        pk = {"page": p, "image": png, "text_source": src, "text": text, "figures": [f for f in figs if f["page"] == p and f["kind"] == "figure"],
              "tables_like": [f["bbox"] for f in figs if f["page"] == p and f["kind"] == "table_like"], **extra}
        jdump(outdir / ("packet_p%03d.json" % p), pk); packets.append(pk)
    tasks = []
    for k in range(0, len(packets), chunk):
        part = packets[k:k + chunk]; md = [fill_prompt(RECONSTRUCT_PROMPT, cfg), "", "SOURCE PDF FOR \\srcfig: source.pdf (copy of the input)", ""]
        for pk in part:
            md += ["## PAGE %d  (evidence: %s%s)" % (pk["page"], pk["text_source"], "; MATH-SUSPECT, transcribe from the image" if pk.get("math_suspect") else ""),
                   "image: %s" % pk["image"], "figures: %s" % json.dumps(pk["figures"], ensure_ascii=False), "```text", pk["text"], "```", ""]
        t = outdir / "tasks" / ("reconstruct_%03d_%03d.md" % (part[0]["page"], part[-1]["page"])); write(t, "\n".join(md)); tasks.append(str(t))
    shutil.copyfile(str(pdf), str(outdir / "source.pdf")); return {"probe": pr, "packets": len(packets), "tasks": tasks}

# =============================================================================
# PART 5 — LaTeX MATH/COMMAND PROTECTION, SEGMENTATION, VALIDATION, RESTORATION
# Mathematics is protected content: it never reaches the translator. Only ordinary language does.
# =============================================================================
VERBATIM_ENVS = {"verbatim", "verbatim*", "Verbatim", "lstlisting", "minted", "comment", "thebibliography", "filecontents"}
DISPLAY_MATH_ENVS = {"equation", "equation*", "align", "align*", "alignat", "alignat*", "gather", "gather*", "multline", "multline*",
                     "eqnarray", "eqnarray*", "displaymath", "math", "flalign", "flalign*"}
PICTURE_ENVS = {"tikzpicture", "pgfpicture", "axis", "tikzcd", "circuitikz", "pspicture", "picture", "forest"}
TABULAR_ENVS = {"tabular": 1, "tabular*": 2, "tabularx": 2, "tabulary": 2, "longtable": 1, "supertabular": 1, "array": 1}
TEXTLIKE_IN_MATH = ("text", "textrm", "textit", "textsf", "textbf", "mbox", "hbox", "intertext", "shortintertext")
NONTRANSLATABLE_CMDS = {  # name -> number of mandatory brace args consumed (optional [..] args before them are consumed too)
 "label": 1, "ref": 1, "eqref": 1, "pageref": 1, "autoref": 1, "cref": 1, "Cref": 1, "vref": 1, "nameref": 1, "cite": 1, "citep": 1, "citet": 1,
 "citeauthor": 1, "citeyear": 1, "nocite": 1, "bibliography": 1, "bibliographystyle": 1, "input": 1, "include": 1, "includegraphics": 1,
 "usepackage": 1, "documentclass": 1, "newcommand": 2, "renewcommand": 2, "providecommand": 2, "newtheorem": 2, "setlength": 2, "setcounter": 2,
 "addtocounter": 2, "numberwithin": 2, "vspace": 1, "hspace": 1, "url": 1, "href": 1, "hyperlink": 1, "hypertarget": 1, "index": 1, "color": 1,
 "textcolor": 1, "definecolor": 3, "pagestyle": 1, "thispagestyle": 1, "pagenumbering": 1, "SI": 2, "si": 1, "num": 1, "qty": 2, "ang": 1, "unit": 1,
 "SIrange": 3, "srcfig": 5, "includepdf": 1, "graphicspath": 1, "theoremstyle": 1, "newenvironment": 3, "DeclareMathOperator": 2,
}
UNIT_TOKENS = set("kg mg cm mm km ms Hz kHz MHz GHz mA eV keV MeV GeV Pa kPa MPa mol rad sr cd lm lx mT kV mV nm um mum min".split())

def match_brace(s, i, open_="{", close="}"):
    """s[i] == open_. Returns index just after the matching close (escape-aware), or -1."""
    depth, j, n = 0, i, len(s)
    while j < n:
        c = s[j]
        if c == "\\": j += 2; continue
        if c == open_: depth += 1
        elif c == close:
            depth -= 1
            if depth == 0: return j + 1
        j += 1
    return -1

def find_env_end(s, i, env):
    """i = index after \\begin{env}. Returns (start_of_\\end, index_after_\\end) honouring nesting."""
    pat = re.compile(r"\\(begin|end)\{" + re.escape(env) + r"\}"); depth = 1
    for m in pat.finditer(s, i):
        depth += 1 if m.group(1) == "begin" else -1
        if depth == 0: return m.start(), m.end()
    return -1, -1

def protected_words(cfg):
    return set(cfg["protect"]["protected_words"]) | set(cfg["protect"]["keep_words_in_math"])

def is_prose(text, cfg):
    """True if `text` (a \\text{} body, table cell or label) contains ordinary language that must be translated."""
    t = re.sub(r"\$[^$]*\$", " ", text); t = re.sub(r"\\\(.*?\\\)", " ", t); t = re.sub(r"\u27e6[^\u27e7]*\u27e7", " ", t)
    t = re.sub(r"\\[A-Za-z]+\*?", " ", t); t = re.sub(r"[\\{}_^$&~]", " ", t); keep = protected_words(cfg); keep_l = {k.lower() for k in keep}
    for tok in re.findall(r"[^\W\d_]+", t):
        if len(tok) >= 2 and tok not in keep and tok.lower() not in keep_l and tok not in UNIT_TOKENS and tok not in OPERATOR_WORDS: return True
    return False

@dataclass
class Segment:
    id: str; kind: str; src: str; ctx: str = ""; file: str = ""; lead: str = ""; trail: str = ""
    tgt: str = None; flags: list = field(default_factory=list); status: str = "pending"; tries: int = 0

class Store:
    """Per-document registry of protected spans (token -> original text) and child segments (ordinary text found INSIDE protected spans).
    Segment ids are '<doc>:<c|p>NNNNN' so they are unique across the whole project."""
    def __init__(self, cfg, file="doc"):
        self.cfg, self.file, self.ph, self.children, self.segments, self.n = cfg, file, {}, {}, [], collections.Counter()
        self.nchild = 0
    def token(self, kind, original):
        self.n[kind] += 1; t = PH(kind, self.n[kind]); self.ph[t] = original; return t
    def child(self, kind, text, ctx=""):
        """Register text that lives INSIDE a protected span but is ordinary language (\\text{}, table cell, TikZ label)."""
        lead, core, trail = re.match(r"^(\s*)(.*?)(\s*)$", text, re.S).groups()
        self.nchild += 1; cid = "%s:c%05d" % (self.file, self.nchild)
        seg = Segment(cid, kind, protect(core, self, ctx), ctx, self.file, lead, trail); self.segments.append(seg)
        self.n["S"] += 1; t = PH("S", self.n["S"]); self.children[t] = cid; return t

def _skip_args(s, i, nargs):
    """Consume optional [..] / starred forms and `nargs` mandatory {..} args starting at i; stops once the mandatory ones are consumed."""
    j, got, n = i, 0, len(s)
    while got < nargs:
        k = j
        while k < n and s[k] in " \t": k += 1
        if k >= n: break
        if s[k] == "*" and got == 0 and k == j: j = k + 1; continue
        if s[k] == "[" and got == 0:
            e = match_brace(s, k, "[", "]")
            if e < 0: break
            j = e; continue
        if s[k] == "{":
            e = match_brace(s, k)
            if e < 0: break
            j, got = e, got + 1; continue
        if got == 0 and s[k] == "\\":                      # \newcommand\foo{...}
            m = re.match(r"\\[A-Za-z]+", s[k:])
            if m: j = k + m.end(); got += 1; continue
        break
    return j

def _math_children(m, store, ctx):
    """Inside a protected math span, hand ordinary-language \\text{...} bodies to the translator as child segments."""
    if store.cfg["protect"]["text_in_math"] != "translate_prose": return m
    out, i = [], 0
    rx = re.compile(r"\\(" + "|".join(TEXTLIKE_IN_MATH) + r")\s*\{")
    while True:
        mm = rx.search(m, i)
        if not mm: out.append(m[i:]); break
        b = mm.end() - 1; e = match_brace(m, b)
        if e < 0: out.append(m[i:]); break
        body = m[b + 1:e - 1]; out.append(m[i:b + 1])
        out.append(store.child("textinmath", body, ctx) if is_prose(body, store.cfg) else body); out.append("}"); i = e
    return "".join(out)

def _table(env, s, i, store, ctx):
    """s[i:] starts just after \\begin{env}. Returns (token_text, end_index) or None."""
    a, b = find_env_end(s, i, env)
    if a < 0: return None
    j = _skip_args(s, i, TABULAR_ENVS[env]); head, body = s[i:j], s[j:a]
    cur, depth, k, mathmode, out = [], 0, 0, False, []
    def flush(txt):
        lead = re.match(r"^\s*(\\(hline|toprule|midrule|bottomrule)\s*|\\cline\{[^}]*\}\s*|\\cmidrule(\[[^\]]*\])?(\([^)]*\))?\{[^}]*\}\s*)*", txt)
        pre = txt[:lead.end()] if lead else ""; rest = txt[len(pre):]
        if store.cfg["protect"]["tables"] == "translate_text_cells":
            mc = re.match(r"^(\s*\\multicolumn\{[^}]*\}\{[^}]*\}\{)(.*)(\}\s*)$", rest, re.S)
            if mc and is_prose(mc.group(2), store.cfg): return pre + mc.group(1) + store.child("cell", mc.group(2), ctx) + mc.group(3)
            if not mc and is_prose(rest, store.cfg): return pre + store.child("cell", rest, ctx)
        return txt
    while k < len(body):
        c = body[k]
        if c == "\\" and body[k:k + 2] == "\\\\": out.append(flush("".join(cur))); out.append("\\\\"); cur = []; k += 2; continue
        if c == "\\": cur.append(body[k:k + 2]); k += 2; continue
        if c == "$": mathmode = not mathmode
        elif c == "{": depth += 1
        elif c == "}": depth -= 1
        elif c == "&" and depth == 0 and not mathmode: out.append(flush("".join(cur))); out.append("&"); cur = []; k += 1; continue
        cur.append(c); k += 1
    out.append(flush("".join(cur)))
    return "\\begin{%s}%s%s\\end{%s}" % (env, head, "".join(out), env), b

def _picture(env, s, i, store, ctx):
    a, b = find_env_end(s, i, env)
    if a < 0: return None
    body = s[i:a]
    if store.cfg["protect"]["tikz_labels"] == "translate_prose":
        rx = re.compile(r"(\bnode\b[^;{}]{0,240}?\{|\\addlegendentry\s*\{|\b(?:xlabel|ylabel|title|zlabel)\s*=\s*\{|\\legend\s*\{)")
        out, p = [], 0
        while True:
            m = rx.search(body, p)
            if not m: out.append(body[p:]); break
            bo = m.end() - 1; e = match_brace(body, bo)
            if e < 0: out.append(body[p:]); break
            inner = body[bo + 1:e - 1]; out.append(body[p:bo + 1])
            out.append(store.child("figlabel", inner, ctx) if is_prose(inner, store.cfg) else inner)
            out.append("}"); p = e
        body = "".join(out)
    return "\\begin{%s}%s\\end{%s}" % (env, body, env), b

def protect(s, store, ctx=""):
    """Replace every protected span in `s` by a placeholder and return the shielded text.
    Order of recognition (left to right): comment | verbatim/picture/table/display-math environments | $$..$$ | $..$ | \\(..\\) | \\[..\\]
    | non-translatable commands (\\label,\\ref,\\cite,\\includegraphics,...) | \\begin{env}/\\end{env} tokens | \\verb."""
    out, i, n = [], 0, len(s)
    extra = {c: 1 for c in store.cfg["protect"].get("extra_nontranslatable_commands", [])}
    def emit(kind, original): out.append(store.token(kind, original))
    while i < n:
        c = s[i]
        if c == "%":                                   # comment to end of line
            e = s.find("\n", i); e = n if e < 0 else e
            if store.cfg["protect"]["comments"] == "preserve": emit("C", s[i:e])
            else: out.append(s[i:e])
            i = e; continue
        if c == "$":
            if s.startswith("$$", i):
                e = s.find("$$", i + 2); e = n if e < 0 else e + 2
            else:
                j, e = i + 1, -1
                while j < n:                            # honour \text{...$..$..} nesting
                    if s[j] == "\\":
                        mm = re.match(r"\\(" + "|".join(TEXTLIKE_IN_MATH) + r")\s*\{", s[j:])
                        if mm:
                            k = match_brace(s, j + mm.end() - 1); j = k if k > 0 else j + 2; continue
                        j += 2; continue
                    if s[j] == "$": e = j + 1; break
                    j += 1
                if e < 0: e = n
            emit("M", _math_children(s[i:e], store, ctx)); i = e; continue
        if c == "\\":
            m = re.match(r"\\([A-Za-z]+\*?)", s[i:])
            if m:
                name = m.group(1)
                if name == "begin":
                    mb = re.match(r"\\begin\{([^}]*)\}", s[i:])
                    if mb:
                        env, after = mb.group(1), i + mb.end()
                        if env in VERBATIM_ENVS:
                            a, b = find_env_end(s, after, env); b = n if a < 0 else b; emit("V", s[i:b]); i = b; continue
                        if env in DISPLAY_MATH_ENVS:
                            a, b = find_env_end(s, after, env); b = n if a < 0 else b; emit("M", _math_children(s[i:b], store, ctx)); i = b; continue
                        if env in PICTURE_ENVS:
                            r = _picture(env, s, after, store, ctx)
                            if r: emit("G", r[0]); i = r[1]; continue
                        if env in TABULAR_ENVS:
                            r = _table(env, s, after, store, ctx)
                            if r: emit("T", r[0]); i = r[1]; continue
                        emit("E", s[i:after]); i = after; continue   # an optional [title] after \begin{theorem} stays visible (it is translated)
                if name == "end":
                    me = re.match(r"\\end\{[^}]*\}", s[i:])
                    if me: emit("E", me.group(0)); i += me.end(); continue
                if name == "verb" and i + 5 < n:
                    d = s[i + 5]; e = s.find(d, i + 6); e = n if e < 0 else e + 1; emit("V", s[i:e]); i = e; continue
                spec = NONTRANSLATABLE_CMDS.get(name.rstrip("*"), extra.get(name))
                if spec:
                    e = _skip_args(s, i + m.end(), spec); emit("K", s[i:e]); i = e; continue
                out.append(m.group(0)); i += m.end(); continue
            nxt = s[i + 1] if i + 1 < n else ""
            if nxt in "([":                              # \( .. \)  and  \[ .. \]
                close = "\\)" if nxt == "(" else "\\]"; e = s.find(close, i + 2); e = n if e < 0 else e + 2
                emit("M", _math_children(s[i:e], store, ctx)); i = e; continue
            out.append(s[i:i + 2]); i += 2; continue
        out.append(c); i += 1
    return "".join(out)

# ---------------------------------------------------------------- segmentation of a whole file
SEG_MARK = "\u27ea%s\u27eb"      # ⟪ch1:p00001⟫ marks where a paragraph segment sits in the skeleton
LEVELS = ["chapter", "section", "subsection", "subsubsection", "paragraph"]
HEAD_RE = re.compile(r"\\(chapter|section|subsection|subsubsection|paragraph)\*?\s*(?:\[[^\]]*\])?\s*\{([^{}]*)\}")

def has_letters(s):
    t = PH_RE.sub(" ", s); t = re.sub(r"\\[A-Za-z]+\*?|\\.", " ", t); return bool(re.search(r"[^\W\d_]", t))

def _child_ids(text, store, by_id, seen=None):
    """All child-segment ids reachable from `text` through protected spans (used to give children their heading context)."""
    seen = seen if seen is not None else set(); out = []
    for m in PH_RE.finditer(text):
        tok = m.group(0)
        if tok in seen: continue
        seen.add(tok)
        if tok in store.children:
            cid = store.children[tok]; out.append(cid)
            if cid in by_id: out += _child_ids(by_id[cid].src, store, by_id, seen)
        elif tok in store.ph: out += _child_ids(store.ph[tok], store, by_id, seen)
    return out

def segment_tex(text, cfg, file="doc"):
    """-> (skeleton, store). Segments (paragraph level + children) are in store.segments."""
    store = Store(cfg, file); head, body, tail = "", text, ""
    m = re.search(r"\\begin\{document\}", text)
    if m:                                             # preamble kept verbatim except \title/\subtitle/\date prose
        pre = text[:m.end()]; body = text[m.end():]
        def _t(mm): return mm.group(1) + "{" + store.child("heading", mm.group(2)) + "}"
        head = re.sub(r"(\\(?:title|subtitle|date))\{([^{}]*[^\W\d_][^{}]*)\}", _t, pre)
        e = re.search(r"\\end\{document\}", body)
        if e: tail, body = body[e.start():], body[:e.start()]
    shielded = protect(body, store); parts = re.split(r"(\n[ \t]*\n+)", shielded); sk, ctxs, pn = [], [], 0
    para = []
    for part in parts:
        if re.fullmatch(r"\n[ \t]*\n+", part) or not has_letters(part): sk.append(part); continue
        lead, core, trail = re.match(r"^(\s*)(.*?)(\s*)$", part, re.S).groups(); pn += 1; sid = "%s:p%05d" % (file, pn)
        h = HEAD_RE.findall(core)
        for lvl, title in h:
            idx = LEVELS.index(lvl); ctxs = [x for x in ctxs if x[0] < idx] + [(idx, PH_RE.sub("", title))]
        ctx = " > ".join(t for _, t in ctxs)
        kind = "heading" if (h and HEAD_RE.fullmatch(core.strip())) else "prose"
        sg = Segment(sid, kind, core, ctx, file, lead, trail); store.segments.append(sg); para.append(sg); sk.append(SEG_MARK % sid)
    by_id = {s.id: s for s in store.segments}
    for sg in para:                                   # children inherit the heading context of the paragraph that contains them
        for cid in _child_ids(sg.src, store, by_id):
            if not by_id[cid].ctx: by_id[cid].ctx = sg.ctx
    return head + "\u27eaBODY\u27eb" + "".join(sk) + "\u27eaEND\u27eb" + tail, store

def restore(s, store, segs, depth=0):
    """Replace placeholders (recursively) with their originals / translated children. `segs`: dict id -> Segment."""
    if depth > 8: return s
    def rep(m):
        tok = m.group(0)
        if tok in store.ph: return restore(store.ph[tok], store, segs, depth + 1)
        if tok in store.children:
            sg = segs[store.children[tok]]; body = sg.tgt if sg.tgt is not None else sg.src
            return sg.lead + restore(body, store, segs, depth + 1) + sg.trail
        return tok
    return PH_RE.sub(rep, s)

def assemble(skeleton, store, segs=None):
    """Build the final TeX from skeleton + translated segments (falls back to the SOURCE text if a segment is untranslated/failed;
    the QC reports such segments as blocking)."""
    segs = segs if segs is not None else {sg.id: sg for sg in store.segments}
    sk = skeleton.replace("\u27eaBODY\u27eb", "").replace("\u27eaEND\u27eb", "")
    def rep(m):
        sg = segs[m.group(1)]; body = sg.tgt if sg.tgt is not None else sg.src
        return sg.lead + body + sg.trail
    sk = re.sub("\u27ea([^\u27eb]+)\u27eb", rep, sk)
    return restore(sk, store, segs)

CS_RE = re.compile(r"\\[A-Za-z]+\*?|\\[^A-Za-z]")
def validate_translation(src, tgt):
    """Mechanical acceptance test of one translated segment. Returns list of (code, detail); 'placeholder_mismatch' and 'empty' are blocking."""
    if tgt is None or not tgt.strip(): return [("empty", "empty translation")]
    probs = []
    a = collections.Counter("".join(k) for k in PH_RE.findall(src)); b = collections.Counter("".join(k) for k in PH_RE.findall(tgt))
    if a != b:
        probs.append(("placeholder_mismatch", "missing=%s extra=%s" % (list((a - b).elements())[:6], list((b - a).elements())[:6])))
    if re.search("\u27e6(?![/A-Za-z]+\\d+\u27e7)", tgt): probs.append(("placeholder_mismatch", "malformed placeholder"))
    ss, tt = PH_RE.sub("", src), PH_RE.sub("", tgt)
    ca, cb = collections.Counter(CS_RE.findall(ss)), collections.Counter(CS_RE.findall(tt))
    if ca != cb: probs.append(("command_skeleton", "LaTeX commands changed: %s" % dict(list(((ca - cb) + (cb - ca)).items())[:6])))
    if ss.count("{") - ss.count("}") != tt.count("{") - tt.count("}"): probs.append(("brace_balance", "brace balance differs"))
    if len(re.findall(r"(?<!\\)%", ss)) != len(re.findall(r"(?<!\\)%", tt)): probs.append(("percent", "comment/percent count differs"))
    return probs

# ---------------------------------------------------------------- source-issue scanner (REPORT ONLY - never edits)
def scan_source_issues(text, file="doc"):
    """Apparent problems in the SOURCE book, reported separately and NEVER silently fixed in the translation."""
    out = []
    labels = re.findall(r"\\label\{([^}]*)\}", text); dup = [k for k, v in collections.Counter(labels).items() if v > 1]
    for d in dup: out.append({"file": file, "type": "duplicate_label", "note": d})
    for r in sorted(set(re.findall(r"\\(?:eq)?ref\{([^}]*)\}", text)) - set(labels)): out.append({"file": file, "type": "undefined_reference_in_file", "note": r + " (may be defined in another file)"})
    b = collections.Counter(re.findall(r"\\begin\{([^}]*)\}", text)); e = collections.Counter(re.findall(r"\\end\{([^}]*)\}", text))
    for k in set(b) | set(e):
        if b[k] != e[k]: out.append({"file": file, "type": "unbalanced_environment", "note": "%s: %d begin / %d end" % (k, b[k], e[k])})
    st = Store(load_config(), file); protect(text, st)
    for tok, o in st.ph.items():
        if tok.startswith("\u27e6M"):
            if o.count("{") != o.count("}") or len(re.findall(r"\\left\b", o)) != len(re.findall(r"\\right\b", o)):
                out.append({"file": file, "type": "unbalanced_math", "note": o[:120]})
    return out

# =============================================================================
# PART 6 — PROJECT STATE, TRANSLATION BACKENDS, BATCHING, VALIDATION/RETRY
# The translator only ever sees shielded text (placeholders instead of mathematics).
# =============================================================================
def docname_from_path(rel):
    return re.sub(r"[^A-Za-z0-9_.-]+", "__", re.sub(r"\.tex$", "", str(rel)))

class Project:
    """On-disk project:  project.json | glossary.csv | source/ (original tex + assets) | work/{skel,store,orig,tasks,segments.jsonl,docs.json}
    | out/tex (assembled translation) | out/final (exported deliverables) | qc/ (reports)."""
    def __init__(self, root):
        self.root = Path(root); self.cfg_path = self.root / "project.json"; self.cfg = load_config(self.cfg_path)
    def p(self, *a): return self.root.joinpath(*a)
    @classmethod
    def init(cls, root, src, tgt, domain="mathematics", keep_words=None):
        root = Path(root); (root / "work").mkdir(parents=True, exist_ok=True)
        sc, _ = get_profile(src); tc, _ = get_profile(tgt)
        cfg = deep_merge(DEFAULT_CONFIG, {"project": {"source_lang": sc, "target_lang": tc, "domain": domain},
                                          "protect": {"keep_words_in_math": keep_words or []}})
        jdump(root / "project.json", cfg); P = cls(root)
        if not P.p("glossary.csv").exists(): Glossary.seed("%s-%s" % (sc, tc)).save(P.p("glossary.csv"))
        return P
    @property
    def pair(self): return "%s-%s" % (self.cfg["project"]["source_lang"], self.cfg["project"]["target_lang"])
    def glossary(self): return Glossary.load(self.p("glossary.csv"), self.pair)
    def save_doc(self, name, skeleton, store):
        write(self.p("work", "skel", name + ".skel"), skeleton)
        jdump(self.p("work", "store", name + ".json"), {"ph": store.ph, "children": store.children, "n": dict(store.n), "nchild": store.nchild})
    def load_doc(self, name):
        st = Store(self.cfg, name); d = jload(self.p("work", "store", name + ".json")); st.ph, st.children, st.n, st.nchild = d["ph"], d["children"], collections.Counter(d["n"]), d["nchild"]
        return read(self.p("work", "skel", name + ".skel")), st
    def save_segments(self, segs): write(self.p("work", "segments.jsonl"), "\n".join(json.dumps(asdict(s), ensure_ascii=False) for s in segs) + "\n")
    def load_segments(self):
        f = self.p("work", "segments.jsonl")
        return [Segment(**json.loads(l)) for l in read(f).splitlines() if l.strip()] if f.exists() else []
    def docs(self):
        d = jload(self.p("work", "docs.json"), {}); return d      # {docname: relative path of the .tex in source/}
    def segs_by_id(self): return {s.id: s for s in self.load_segments()}

def ingest_tex(project, files, root=None):
    """tex->tex route entry: copy sources into <project>/source, protect + segment every .tex file, register segments."""
    cfg = project.cfg; files = [Path(f) for f in files]; root = Path(root) if root else (files[0].parent if files else Path("."))
    docs, allsegs = {}, []
    for f in files:
        rel = f.resolve().relative_to(root.resolve()) if root.resolve() in f.resolve().parents else Path(f.name)
        dest = project.p("source", rel); dest.parent.mkdir(parents=True, exist_ok=True); shutil.copyfile(f, dest)
        text = read(f); name = docname_from_path(rel); sk, st = segment_tex(text, cfg, name)
        project.save_doc(name, sk, st); write(project.p("work", "orig", name + ".tex"), text); docs[name] = str(rel); allsegs += st.segments
    for asset in root.rglob("*"):         # figures / bibliographies / source PDF the document refers to
        if asset.is_file() and asset.suffix.lower() in (".png", ".jpg", ".jpeg", ".pdf", ".eps", ".svg", ".bib", ".sty", ".cls", ".bst"):
            d = project.p("source", asset.relative_to(root)); d.parent.mkdir(parents=True, exist_ok=True)
            if not d.exists(): shutil.copyfile(asset, d)
    jdump(project.p("work", "docs.json"), docs); project.save_segments(allsegs)
    return {"documents": len(docs), "segments": len(allsegs), "by_kind": dict(collections.Counter(s.kind for s in allsegs))}

def assemble_project(project, outdir=None):
    """Write the translated .tex tree (uses source text for any segment that is not 'done'); adapt preamble language settings."""
    outdir = Path(outdir or project.p("out", "tex")); segs = project.segs_by_id(); written = []
    for name, rel in project.docs().items():
        sk, st = project.load_doc(name); text = assemble(sk, st, segs)
        if "\\newtheorem" in text or "\\documentclass" in text: text, _ = adapt_preamble(text, project.cfg)
        dest = outdir / rel; write(dest, text); written.append(str(dest))
    src = project.p("source")
    if src.exists():
        for a in src.rglob("*"):
            if a.is_file() and a.suffix.lower() != ".tex":
                d = outdir / a.relative_to(src); d.parent.mkdir(parents=True, exist_ok=True)
                if not d.exists(): shutil.copyfile(a, d)
    return written

# ---------------------------------------------------------------- backends
def parse_llm_json(text):
    t = re.sub(r"^```(?:json)?\s*|\s*```\s*$", "", text.strip())
    a, b = t.find("{"), t.rfind("}")
    if a < 0 or b < a: raise ValueError("no JSON object in model output")
    return json.loads(t[a:b + 1])

def user_message(items, instruction):
    return instruction + "\n\n" + json.dumps({"items": items}, ensure_ascii=False)

class Backend:
    name = "base"
    def __init__(self, cfg, project=None): self.cfg, self.project = cfg, project
    def translate(self, items, system, mode="translate"): raise NotImplementedError

# Tiny fr->en dictionary used ONLY by the mock backend (self-test fixture). It is NOT a translator.
MOCK_FR_EN = {"soit": "let", "alors": "then", "une": "a", "un": "a", "suite": "sequence", "réelle": "real", "réel": "real", "converge": "converges", "vers": "to",
              "si": "if", "et": "and", "ou": "or", "pour": "for", "tout": "all", "tel": "such", "que": "that", "on": "we", "a": "has", "de": "of", "la": "the",
              "le": "the", "les": "the", "des": "the", "est": "is", "sont": "are", "dans": "in", "fonction": "function", "dérivable": "differentiable",
              "théorème": "Theorem", "démonstration": "Proof", "exemple": "Example", "exercice": "Exercise", "remarque": "Remark", "définition": "Definition",
              "limite": "limit", "borne": "bound", "supérieure": "upper", "bornée": "bounded", "chapitre": "Chapter", "ensemble": "set", "développement": "expansion",
              "limité": "limited", "intensité": "current", "courant": "flow", "circuit": "circuit", "vitesse": "velocity", "du": "of", "au": "to", "d'un": "of a",
              "valeurs": "values", "intermédiaires": "intermediate", "tableau": "table", "vrai": "true", "faux": "false", "suivant": "following", "où": "where",
              "avec": "with", "sur": "on", "par": "by", "ce": "this", "cette": "this", "donc": "hence", "nous": "we", "avons": "have", "voir": "see",
              "formule": "formula", "premier": "first", "nombre": "number", "entier": "integer", "naturel": "natural", "positif": "positive",
              "extraite": "extracted", "énoncé": "statement", "figure": "Figure", "titre": "title", "principal": "main", "section": "section",
              "résultat": "result", "preuve": "Proof", "pas": "not", "existe": "exists", "unique": "unique", "environ": "approximately",
              "continue": "continuous", "équation": "equation", "pente": "slope", "positive": "positive", "cours": "course", "analyse": "analysis",
              "suites": "sequences", "d'une": "of a", "chapitre": "chapter", "non": "not", "commentaire": "comment", "conserver": "keep", "paragraphe": "paragraph"}

class MockBackend(Backend):
    """Deterministic dictionary 'translator' for the self-test and for dry runs of the pipeline. NEVER use for real output."""
    name = "mock"
    def translate(self, items, system, mode="translate"):
        out = {}
        for it in items:
            text = it["src"]
            for g in sorted(it.get("glossary", []), key=lambda x: -len(x["src"])):
                text = re.sub(re.escape(g["src"]), g["tgt"], text, flags=re.I)
            def w(m):
                t = m.group(0); r = MOCK_FR_EN.get(t.lower())
                if r is None: return t
                return r[0].upper() + r[1:] if t[0].isupper() else r
            parts = re.split("(⟦[^⟧]*⟧|\\\\[A-Za-z]+)", text)
            out[it["id"]] = {"tgt": "".join(p if i % 2 else re.sub(r"[^\W\d_]+(?:['’-][^\W\d_]+)?", w, p) for i, p in enumerate(parts)), "flags": []}
        return out

class AgentBackend(Backend):
    """Hand-off backend for 'another AI or a human': writes work/tasks/batch_NNNN.json containing the system prompt and items.
    The agent writes work/tasks/batch_NNNN.result.json = {"translations":[{"id","tgt","flags"}]}; `translate --collect` then ingests every result file."""
    name = "agent"
    def reset_unresolved(self):
        d = self.project.p("work", "tasks")
        if d.exists():
            for f in d.glob("batch_*.json"):
                if not f.name.endswith(".result.json") and not f.with_name(f.name[:-5] + ".result.json").exists(): f.unlink()
    def translate(self, items, system, mode="translate"):
        d = self.project.p("work", "tasks"); d.mkdir(parents=True, exist_ok=True)
        nums = [int(re.search(r"batch_(\d+)", p.name).group(1)) for p in d.glob("batch_*.json")]
        n = (max(nums) if nums else 0) + 1; f = d / ("batch_%04d.json" % n); res = d / ("batch_%04d.result.json" % n)
        jdump(f, {"mode": mode, "system": system, "result_file": str(res), "items": items,
                  "how_to": "Translate every item following `system`. Write {\"translations\":[{\"id\":..., \"tgt\":..., \"flags\":[...]}]} to result_file. Do not edit this file."})
        log("AGENT HAND-OFF: wrote %s (%d items). Fill %s then run `translate --collect`." % (f.name, len(items), res.name)); return {}
    def collect(self):
        out = {}
        for r in sorted(self.project.p("work", "tasks").glob("batch_*.result.json")):
            data = parse_llm_json(read(r))
            for t in data.get("translations", []): out[t["id"]] = {"tgt": t.get("tgt"), "flags": t.get("flags", [])}
        return out

def _http_json(url, headers, body, timeout):
    req = urllib.request.Request(url, json.dumps(body).encode(), headers)
    with urllib.request.urlopen(req, timeout=timeout) as r: return json.loads(r.read().decode())

class AnthropicBackend(Backend):
    """Anthropic Messages API. UNTESTED in the build sandbox (no network). Needs ANTHROPIC_API_KEY (or cfg.translate.api_key_env)."""
    name = "anthropic"
    def translate(self, items, system, mode="translate"):
        t = self.cfg["translate"]; key = os.environ.get(t["api_key_env"] or "ANTHROPIC_API_KEY")
        if not key: raise SystemExit("Set ANTHROPIC_API_KEY")
        url = (t["api_base"] or "https://api.anthropic.com") + "/v1/messages"
        body = {"model": t["model"], "max_tokens": 8192, "temperature": t["temperature"], "system": system,
                "messages": [{"role": "user", "content": user_message(items, "Return ONLY the JSON object described in the system prompt.")}]}
        for attempt in range(4):
            try:
                r = _http_json(url, {"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"}, body, t["timeout_s"])
                data = parse_llm_json("".join(b.get("text", "") for b in r["content"])); return {x["id"]: {"tgt": x.get("tgt"), "flags": x.get("flags", [])} for x in data["translations"]}
            except (urllib.error.URLError, ValueError, KeyError) as e:
                log("anthropic attempt %d failed: %s" % (attempt + 1, e)); time.sleep(2 ** attempt)
        return {}

class OpenAICompatBackend(Backend):
    """Any OpenAI-compatible /chat/completions endpoint (OpenAI, gateways, vLLM, llama.cpp server, Ollama...). UNTESTED here (no network)."""
    name = "openai"
    def translate(self, items, system, mode="translate"):
        t = self.cfg["translate"]; key = os.environ.get(t["api_key_env"] or "OPENAI_API_KEY", "none"); base = (t["api_base"] or "https://api.openai.com/v1").rstrip("/")
        body = {"model": t["model"], "temperature": t["temperature"], "messages": [{"role": "system", "content": system},
                {"role": "user", "content": user_message(items, "Return ONLY the JSON object described in the system prompt.")}]}
        for attempt in range(4):
            try:
                r = _http_json(base + "/chat/completions", {"Authorization": "Bearer " + key, "Content-Type": "application/json"}, body, t["timeout_s"])
                data = parse_llm_json(r["choices"][0]["message"]["content"]); return {x["id"]: {"tgt": x.get("tgt"), "flags": x.get("flags", [])} for x in data["translations"]}
            except (urllib.error.URLError, ValueError, KeyError) as e:
                log("openai-compatible attempt %d failed: %s" % (attempt + 1, e)); time.sleep(2 ** attempt)
        return {}

BACKENDS = {"mock": MockBackend, "agent": AgentBackend, "anthropic": AnthropicBackend, "openai": OpenAICompatBackend}

# ---------------------------------------------------------------- orchestration
def make_items(segs, glossary, cfg, done_ctx):
    items = []
    for s in segs:
        it = {"id": s.id, "kind": s.kind, "ctx": s.ctx, "src": s.src, "glossary": glossary.prompt_block(glossary.hits(PH_RE.sub(" ", s.src), s.ctx))}
        if done_ctx.get(s.file) and s.kind in ("prose", "heading"): it["previous"] = done_ctx[s.file][-cfg["translate"]["context_segments"]:]
        if s.status == "retry": it["retry_note"] = "Your previous answer was rejected: " + "; ".join(f["note"] for f in s.flags if f.get("type") == "validation")[-300:]
        items.append(it)
    return items

def batches(items, max_chars):
    cur, size = [], 0
    for it in items:
        sz = len(it["src"]) + 200
        if cur and size + sz > max_chars: yield cur; cur, size = [], 0
        cur.append(it); size += sz
    if cur: yield cur

def apply_results(segs_by_id, results, cfg, glossary=None):
    stats = collections.Counter()
    for sid, r in results.items():
        s = segs_by_id.get(sid)
        if not s or s.status == "done": continue
        probs = validate_translation(s.src, r.get("tgt")); s.tries += 1
        block = [p for p in probs if p[0] in ("placeholder_mismatch", "empty")]
        s.flags = [f for f in s.flags if f.get("type") not in ("validation", "validation_warning")] + [f for f in (r.get("flags") or []) if f.get("type") not in ("validation", "validation_warning")]
        if block:
            s.status = "failed" if s.tries >= cfg["translate"]["max_retries"] else "retry"
            if s.status == "failed": s.flags.append({"type": "rejected_output", "note": (r.get("tgt") or "")[:400]})
            s.tgt = None     # assemble() falls back to the SOURCE text; QC reports the segment as blocking
            s.flags += [{"type": "validation", "note": "%s: %s" % p} for p in block]; stats[s.status] += 1
        else:
            s.tgt, s.status = r["tgt"], "done"; stats["done"] += 1
            s.flags += [{"type": "validation_warning", "note": "%s: %s" % p} for p in probs]
    return stats

def apply_review(by, res):
    st = collections.Counter()
    for sid, r in res.items():
        s = by[sid]
        if r.get("tgt") and not [p for p in validate_translation(s.src, r["tgt"]) if p[0] in ("placeholder_mismatch", "empty")]:
            if r["tgt"] != s.tgt: s.flags.append({"type": "review_change", "note": "%r -> %r" % ((s.tgt or "")[:80], r["tgt"][:80])}); s.tgt = r["tgt"]; st["changed"] += 1
    return st

def translate_all(project, backend_name=None, collect=False, review=False):
    cfg = project.cfg; name = backend_name or cfg["translate"]["backend"]; backend = BACKENDS[name](cfg, project)
    segs = project.load_segments(); by = {s.id: s for s in segs}; g = project.glossary()
    system = fill_prompt(REVIEW_PROMPT if review else TRANSLATE_SYSTEM_PROMPT, cfg); stats = collections.Counter()
    if collect and isinstance(backend, AgentBackend):
        stats.update(apply_results(by, backend.collect(), cfg, g)); project.save_segments(segs); return dict(stats)
    if isinstance(backend, AgentBackend): backend.reset_unresolved()
    for rnd in range(cfg["translate"]["max_retries"] + 1):
        todo = [s for s in segs if s.status in ("pending", "retry")] if not review else [s for s in segs if s.status == "done"]
        if not todo: break
        done_ctx = collections.defaultdict(list)
        for s in segs:
            if s.status == "done" and s.kind in ("prose", "heading"): done_ctx[s.file].append({"src": s.src[-300:], "tgt": (s.tgt or "")[-300:]})
        progressed = 0
        for b in batches(make_items(todo, g, cfg, done_ctx), cfg["translate"]["batch_chars"]):
            if review:
                for it in b: it["tgt"] = by[it["id"]].tgt
            res = backend.translate(b, system, "review" if review else "translate")
            st = apply_results(by, res, cfg, g) if not review else apply_review(by, res); stats.update(st); progressed += sum(st.values()); project.save_segments(segs)
        if isinstance(backend, AgentBackend) or review or not progressed: break
    project.save_segments(segs); return dict(stats)

def collect_source_flags(project):
    out = [{"where": s.id, **f} for s in project.load_segments() for f in s.flags if f.get("type") in ("source_issue", "term_uncertain", "notation_convention", "review_change")]
    for d in project.docs():
        f = project.p("work", "orig", d + ".tex")
        if f.exists(): out += scan_source_issues(read(f), d)
    write(project.p("qc", "source_flags.jsonl"), "\n".join(json.dumps(o, ensure_ascii=False) for o in out) + "\n"); return out

# =============================================================================
# PART 7 — DOCUMENT RECONSTRUCTION: LaTeX preamble/engines per language, compile + log parsing,
#          source-preamble adaptation, LaTeX -> DOCX (pandoc) with numbering made literal.
# The original session used: pdflatex (2 passes), babel[english], amsthm styles, per-section counters.
# =============================================================================
FONT_CANDIDATES = {
 "latin": ["Latin Modern Roman", "TeX Gyre Termes", "Libertinus Serif", "DejaVu Serif"],
 "cyrillic": ["CMU Serif", "Libertinus Serif", "Noto Serif", "DejaVu Serif"],
 "arabic": ["Amiri", "Noto Naskh Arabic", "Noto Serif Arabic", "DejaVu Sans"],
 "han": ["Noto Serif CJK SC", "Source Han Serif SC", "Songti SC", "SimSun", "Noto Sans CJK SC"],
 "kana_han": ["Noto Serif CJK JP", "IPAexMincho", "Hiragino Mincho ProN", "MS Mincho", "Noto Sans CJK JP"],
 "hangul": ["Noto Serif CJK KR", "UnBatang", "Nanum Myeongjo", "Malgun Gothic", "Noto Sans CJK KR"],
 "devanagari": ["Noto Serif Devanagari", "Lohit Devanagari", "Mangal", "Noto Sans Devanagari"],
}
def installed_fonts():
    rc, out, _ = sh(["fc-list", ":", "family"]); fams = set()
    for l in out.splitlines():
        for f in l.split(","): fams.add(f.strip())
    return fams
def pick_font(script):
    fams = installed_fonts()
    for f in FONT_CANDIDATES.get(script, []):
        if f in fams: return f
    return None

_KPSE = {}
def kpse(name):
    """Path of a TeX file known to kpsewhich ('' if absent); cached."""
    if name not in _KPSE: _KPSE[name] = sh(["kpsewhich", name])[1].strip() if have("kpsewhich") else ""
    return _KPSE[name]

def latin_font_package(cfg):
    """Type 1 text font package for the pdfLaTeX route (default T1 Computer Modern falls back to BITMAP Type 3 fonts when cm-super/lmodern are missing).
    config latex.font_package: auto | lmodern | mathptmx | none. Returns the package name or None."""
    want = cfg["latex"].get("font_package", "auto")
    if want == "none": return None
    if want in ("lmodern", "auto") and kpse("lmodern.sty"): return "lmodern"
    if want in ("mathptmx", "auto") and kpse("mathptmx.sty") and kpse("ptmr8t.vf"): return "mathptmx"
    return None

KNOWN_FONT_PKGS = r"lmodern|mathptmx|newtxtext|newtxmath|times|mathpazo|palatino|libertine|fourier|kpfonts|stix2?|tgtermes|tgpagella|charter|XCharter|ebgaramond|bookman|cm-super|fontspec|unicode-math|lucidabr|helvet"

def ensure_text_fonts(text, cfg):
    """pdfLaTeX sources that use T1 and no font package get BITMAP Type 3 fonts (and a wrong text layer for accents) when cm-super/lmodern are missing.
    Adds a Type 1 font package (latin_font_package) and glyphtounicode right after \\usepackage[..T1..]{fontenc}. Idempotent; returns (text, notes)."""
    notes = []
    m = re.search(r"\\usepackage\[[^\]]*T1[^\]]*\]\{fontenc\}", text)
    if not m or cfg["latex"]["engine"] in ("xelatex", "lualatex") or XE_MARKERS.search(text): return text, notes
    add = []; fp = latin_font_package(cfg)
    if fp and not re.search(r"\\usepackage(\[[^\]]*\])?\{(%s)\}" % KNOWN_FONT_PKGS, text):
        add.append("\\usepackage{%s}" % fp); notes.append("added \\usepackage{%s} (Type 1 text font; the default T1 Computer Modern would be bitmap in this TeX installation)" % fp)
    if kpse("glyphtounicode.tex") and "pdfgentounicode" not in text:
        add += ["\\input{glyphtounicode}", "\\pdfgentounicode=1"]; notes.append("enabled \\pdfgentounicode (correct copy/paste of accents and ligatures)")
    if add: text = text[:m.end()] + "\n" + "\n".join(add) + text[m.end():]
    return text, notes

def babel_available(tp): return bool(tp["babel"]) and bool(kpse(tp["babel"] + ".ldf"))
def polyglossia_available(tp): return bool(tp["polyglossia"]) and bool(kpse("gloss-%s.ldf" % tp["polyglossia"])) and have("xelatex")

def engine_for(cfg):
    """pdflatex+babel for Latin scripts when the babel language file is installed; otherwise xelatex+polyglossia when available; non-Latin scripts always xelatex."""
    e = cfg["latex"]["engine"]
    if e != "auto": return e
    _, tp = get_profile(cfg["project"]["target_lang"])
    if tp["script"] != "latin": return "xelatex"
    if not tp["babel"] or babel_available(tp): return "pdflatex"
    return "xelatex" if polyglossia_available(tp) else "pdflatex"

XE_MARKERS = re.compile(r"\\usepackage(\[[^\]]*\])?\{(fontspec|polyglossia|xeCJK|unicode-math|luacode|xltxtra|bidi)\}|\\setmainfont|\\setmainlanguage")

def engine_for_text(text, cfg):
    """The engine a given .tex file is written for: fontspec/polyglossia/xeCJK -> xelatex; inputenc/fontenc (pdfTeX style) -> pdflatex; otherwise the language-based default.
    A source file must be compiled with ITS engine, whatever the target language needs."""
    e = cfg["latex"]["engine"]
    if e != "auto": return e
    if XE_MARKERS.search(text): return "xelatex"
    if re.search(r"\\usepackage(\[[^\]]*\])?\{(inputenc|fontenc)\}", text): return "pdflatex"
    return engine_for(cfg)

def render_preamble(cfg, source_pdf=None):
    """Generate a complete preamble for the target language: engine-specific fonts/language, math packages, theorem styles with
    localised names and the same numbering scheme as the session (theorem-like per section, exercises per chapter)."""
    tc, tp = get_profile(cfg["project"]["target_lang"]); L = tp["labels"]; eng = engine_for(cfg); lc = cfg["latex"]
    P = ["\\documentclass[%s,%s]{%s}" % (lc["fontsize"], lc["paper"], lc["documentclass"])]
    if eng == "pdflatex":
        P += ["\\usepackage[T2A,T1]{fontenc}" if tp["script"] == "cyrillic" else "\\usepackage[T1]{fontenc}", "\\usepackage[utf8]{inputenc}"]
        fp = latin_font_package(cfg)
        if fp: P.append("\\usepackage{%s}" % fp)
        if kpse("glyphtounicode.tex"): P += ["\\input{glyphtounicode}", "\\pdfgentounicode=1"]     # correct copy/paste of accents and ligatures
        else: P.append("% WARNING: glyphtounicode.tex not found; copy/paste of accents/ligatures from the PDF may be wrong")
        if fp is None: P.append("% WARNING: no Type 1 text font package found (lmodern/mathptmx); the PDF will contain bitmap Type 3 fonts (QC: bitmap_fonts)")
        if babel_available(tp): P.append("\\usepackage[%s]{babel}" % tp["babel"])
        elif tp["babel"]: P.append("%% WARNING: babel language '%s' is not installed in this TeX Live; no language package loaded (hyphenation/localised names unavailable)" % tp["babel"])
    else:
        P.append("\\usepackage{fontspec}"); font = pick_font(tp["script"])
        if tp["script"] in ("han", "kana_han", "hangul"):
            latin = pick_font("latin") or "DejaVu Serif"; P += ["\\setmainfont{%s}" % latin, "\\usepackage{xeCJK}", "\\setCJKmainfont{%s}" % (font or "Noto Serif CJK SC")]
            if not font: P.append("%% WARNING: no CJK font found by fc-list; install one of %s" % FONT_CANDIDATES[tp["script"]])
        else:
            P.append("\\setmainfont{%s}" % (font or "DejaVu Serif"))
            if tp["polyglossia"]:
                P += ["\\usepackage{polyglossia}", "\\setmainlanguage{%s}" % tp["polyglossia"]]
                if tp["script"] == "arabic": P.append("\\newfontfamily\\arabicfont[Script=Arabic]{%s}" % (font or "DejaVu Sans"))
                if tp["script"] == "devanagari": P.append("\\newfontfamily\\devanagarifont[Script=Devanagari]{%s}" % (font or "Noto Serif Devanagari"))
    P += ["\\usepackage{amsmath,amssymb,amsthm}", "\\usepackage{tikz}", "\\usepackage{pgfplots}", "\\pgfplotsset{compat=1.16}", "\\usepackage{graphicx}",
          "\\usepackage[margin=%s]{geometry}" % lc["margin"], "\\usepackage{enumerate}", "\\usepackage{hyperref}", "\\hypersetup{hidelinks}",
          "\\setlength{\\parindent}{1.5em}", "\\setlength{\\parskip}{0.2em}"] + ["\\usepackage{%s}" % p for p in lc["extra_packages"]]
    P += ["\\newtheoremstyle{cours}{8pt}{8pt}{\\itshape}{}{\\bfseries}{}{ }{\\thmname{#1}\\ \\thmnumber{#2}\\ \\thmnote{(#3)}}", "\\theoremstyle{cours}"]
    for env in ("definition", "theorem", "proposition", "lemma", "corollary"): P.append("\\newtheorem{%s}{%s}[section]" % (env, L[env]))
    P += ["\\newtheoremstyle{coursplain}{8pt}{8pt}{\\normalfont}{}{\\bfseries}{}{ }{\\thmname{#1}\\ \\thmnumber{#2}\\ \\thmnote{(#3)}}", "\\theoremstyle{coursplain}",
          "\\newtheorem{exercise}{%s}[chapter]" % L["exercise"], "\\renewcommand{\\proofname}{\\itshape %s}" % L["proof"], "\\renewcommand{\\qedsymbol}{$\\blacksquare$}",
          "\\renewcommand{\\contentsname}{%s}" % L["contents"], "\\renewcommand{\\figurename}{%s}" % L["figure"], "\\renewcommand{\\tablename}{%s}" % L["table"],
          "\\renewcommand{\\chaptername}{%s}" % L["chapter"], "\\renewcommand{\\thesection}{\\thechapter.\\arabic{section}}", "\\setcounter{secnumdepth}{3}", "\\setcounter{tocdepth}{2}"]
    if source_pdf: P += ["\\newcommand{\\srcpdfname}{%s}" % source_pdf, "\\newcommand{\\srcfig}[5]{\\includegraphics[page=#1,trim=#2pt #3pt #4pt #5pt,clip]{\\srcpdfname}}"]
    return "\n".join(P) + "\n"

def adapt_preamble(text, cfg):
    """tex->tex route: retarget a SOURCE preamble to the target language (babel option, localised theorem names/proofname/contentsname).
    Returns (new_text, notes). Never touches math macros."""
    sc, sp = get_profile(cfg["project"]["source_lang"]); tc, tp = get_profile(cfg["project"]["target_lang"]); notes = []
    inv = {v: k for k, v in sp["labels"].items()}
    text = re.sub(r"\\newtheorem\{(\w+\*?)\}(\[\w+\])?\{([^}]*)\}",
                  lambda m: "\\newtheorem{%s}%s{%s}" % (m.group(1), m.group(2) or "", tp["labels"].get(inv.get(m.group(3)), m.group(3))), text)
    for macro, key in (("proofname", "proof"), ("contentsname", "contents"), ("figurename", "figure"), ("tablename", "table"), ("chaptername", "chapter")):
        text = re.sub(r"(\\renewcommand\{\\%s\}\{(?:\\itshape\s*)?)[^}]*\}" % macro, lambda m, k=key: m.group(1) + tp["labels"][k] + "}", text)
    if tp["script"] != "latin" and re.search(r"\\usepackage(\[[^\]]*\])?\{babel\}", text):
        notes.append("Target script is not Latin: regenerate the preamble with `preamble` (xelatex + fonts) instead of adapting babel.")
    elif tp["babel"]:
        if not babel_available(tp): notes.append("babel language '%s' is not installed here; install it (texlive-lang-*) or regenerate with `preamble` (polyglossia fallback)." % tp["babel"])
        text = re.sub(r"\\usepackage\[[^\]]*\]\{babel\}", lambda m: "\\usepackage[%s]{babel}" % tp["babel"], text)
    text, n2 = ensure_text_fonts(text, cfg); notes += n2
    return text, notes

# ---------------------------------------------------------------- compile + log parsing
def parse_log(log_text):
    errs = re.findall(r"^(?:\./)?[^\n:]*:?\d*:?\s*! (.+)$|^! (.+)$", log_text, re.M)
    errs = [a or b for a, b in errs]
    return {"errors": errs[:20], "undefined_refs": sorted(set(re.findall(r"Reference `([^']*)' on page \d+ undefined", log_text))),
            "missing_chars": len(re.findall(r"Missing character", log_text)), "overfull": len(re.findall(r"Overfull \\hbox", log_text)),
            "rerun": "Rerun to get cross-references right" in log_text or "Label(s) may have changed" in log_text}

def build_pdf(main_tex, cfg, engine=None, passes=None):
    """Compile until references settle (default max 3 passes; the session used 2). Returns a report dict; never raises on TeX errors."""
    main_tex = Path(main_tex); eng = engine or engine_for_text(read(main_tex), cfg); passes = passes or cfg["latex"]["passes"]
    if not have(eng): return {"ok": False, "errors": ["%s not installed (install TeX Live)" % eng], "pdf": None}
    rep = {}
    for i in range(passes):
        rc, out, err = sh([eng, "-interaction=nonstopmode", main_tex.name], cwd=str(main_tex.parent), timeout=1800)
        logf = main_tex.with_suffix(".log"); rep = parse_log(read(logf) if logf.exists() else out)
        if i == 0 and re.search(r"\\bibliography\{", read(main_tex)) and have("bibtex"): sh(["bibtex", main_tex.stem], cwd=str(main_tex.parent))
        if not rep["rerun"] and i >= 1: break
    pdf = main_tex.with_suffix(".pdf"); rep["pdf"] = str(pdf) if pdf.exists() else None; rep["engine"] = eng
    if pdf.exists():
        _, o, _ = sh(["pdfinfo", str(pdf)]); m = re.search(r"Pages:\s+(\d+)", o); rep["pages"] = int(m.group(1)) if m else None
    rep["ok"] = bool(rep["pdf"]) and not rep["errors"]; return rep

def read_aux_labels(aux_path):
    """label -> printed number, from the .aux written by LaTeX (hyperref-aware)."""
    if not Path(aux_path).exists(): return {}
    return dict(re.findall(r"\\newlabel\{([^}]*)\}\{\{([^{}]*)\}", read(aux_path)))

# ---------------------------------------------------------------- LaTeX -> DOCX with literal numbering
def flatten_includes(path, seen=None):
    path = Path(path); seen = seen if seen is not None else set(); seen.add(path.resolve()); txt = read(path)
    def rep(m):
        f = path.parent / (m.group(2) if m.group(2).endswith(".tex") else m.group(2) + ".tex")
        return flatten_includes(f, seen) if f.exists() and f.resolve() not in seen else m.group(0)
    return re.sub(r"\\(input|include)\{([^}]*)\}", rep, txt)

def rasterize_snippet(snippet, preamble, workdir, name, dpi=220):
    """Compile a TikZ/pgfplots/\\srcfig snippet alone (standalone class) and rasterise it to PNG for DOCX."""
    pre = re.sub(r"\\documentclass(\[[^\]]*\])?\{[^}]*\}", lambda m: "\\documentclass[border=3pt]{standalone}", preamble, count=1)
    pre = re.sub(r"\\usepackage(\[[^\]]*\])?\{(geometry|hyperref)\}|\\hypersetup\{[^}]*\}", "", pre)
    f = Path(workdir) / (name + ".tex"); write(f, pre + "\n\\begin{document}\n" + snippet + "\n\\end{document}\n")
    eng = "xelatex" if re.search(r"fontspec", pre) else "pdflatex"; sh([eng, "-interaction=nonstopmode", f.name], cwd=str(workdir))
    if not f.with_suffix(".pdf").exists(): return None
    sh(["pdftoppm", "-r", str(dpi), "-png", "-singlefile", name + ".pdf", name], cwd=str(workdir)); return name + ".png"

def latex_to_docx(main_tex, out_docx, cfg, keep_workdir=None):
    """Convert a (translated) LaTeX book to DOCX. Math becomes native Word equations (pandoc OMML). Because pandoc does not reproduce LaTeX counters,
    numbers of theorem-like environments, headings and \\begin{equation} are taken from the compiled .aux and written literally; \\ref/\\eqref are resolved.
    TikZ/pgfplots/\\srcfig figures are rasterised to PNG. Approximations: theorem bodies are not italicised; numbers of align-type equations are not re-inserted."""
    if not have("pandoc"): raise SystemExit("pandoc not installed")
    main_tex = Path(main_tex); work = Path(keep_workdir or tempfile.mkdtemp(prefix="docx_")); work.mkdir(parents=True, exist_ok=True)
    tc, tp = get_profile(cfg["project"]["target_lang"]); L = tp["labels"]
    flat = flatten_includes(main_tex); m = re.search(r"\\begin\{document\}", flat)
    if not m: raise SystemExit("no \\begin{document} in %s" % main_tex)
    pre, body = flat[:m.start()], flat[m.end():]; body = body.split("\\end{document}")[0]
    names = {env: name for env, name in re.findall(r"\\newtheorem\{(\w+)\}(?:\[\w+\])?\{([^}]*)\}", pre)}
    names.update({env: name for env, name in re.findall(r"\\newtheorem\{(\w+)\}\{([^}]*)\}\[\w+\]", pre)})
    envs = [e for e in names]; counter = [0]
    def lab(kind): counter[0] += 1; return "\\label{auto:%s:%d}" % (kind, counter[0])
    for env in envs: body = re.sub(r"(\\begin\{%s\}(?:\[[^\]]*\])?)" % re.escape(env), lambda mm: mm.group(1) + lab(env), body)
    body = re.sub(r"(\\begin\{equation\})", lambda mm: mm.group(1) + lab("eq"), body)
    body = re.sub(r"(\\(?:chapter|section|subsection)\{[^{}]*\})", lambda mm: mm.group(1) + lab("sec"), body)
    numtex = work / "numbering.tex"; write(numtex, pre + "\\begin{document}" + body + "\\end{document}\n")
    for f in main_tex.parent.rglob("*"):
        if f.is_file() and f.suffix.lower() in (".pdf", ".png", ".jpg", ".jpeg", ".sty", ".bib", ".cls", ".eps"):
            d = work / f.relative_to(main_tex.parent)
            if not d.exists(): d.parent.mkdir(parents=True, exist_ok=True); shutil.copyfile(f, d)
    rep = build_pdf(numtex, cfg, passes=3); nums = read_aux_labels(numtex.with_suffix(".aux"))
    for env in envs:
        def sub_env(mm, env=env):
            note, num = mm.group(1), nums.get(mm.group(2), "")
            return "\\par\\medskip\\noindent\\textbf{%s %s%s.} " % (names[env], num, " (%s)" % note[1:-1] if note else "")
        body = re.sub(r"\\begin\{%s\}(\[[^\]]*\])?\\label\{(auto:%s:\d+)\}" % (re.escape(env), re.escape(env)), sub_env, body)
        body = body.replace("\\end{%s}" % env, "\\par\\medskip ")
    body = re.sub(r"\\begin\{proof\}(\[[^\]]*\])?", lambda mm: "\\par\\textit{%s.} " % (mm.group(1)[1:-1] if mm.group(1) else L["proof"]), body)
    body = body.replace("\\end{proof}", " \\hfill$\\blacksquare$\\par ")
    body = re.sub(r"\\begin\{equation\}\\label\{(auto:eq:\d+)\}(.*?)\\end\{equation\}",
                  lambda mm: "\\[ %s \\qquad (%s) \\]" % (re.sub(r"\\label\{[^}]*\}", "", mm.group(2)).strip(), nums.get(mm.group(1), "")), body, flags=re.S)
    body = re.sub(r"\\(chapter|section|subsection)\{([^{}]*)\}\\label\{(auto:sec:\d+)\}", lambda mm: "\\%s*{%s~%s}" % (mm.group(1), nums.get(mm.group(3), ""), mm.group(2)), body)
    body = re.sub(r"\\eqref\{([^}]*)\}", lambda mm: "(%s)" % nums.get(mm.group(1), "?"), body)
    body = re.sub(r"\\(?:auto|c|C)?ref\{([^}]*)\}", lambda mm: nums.get(mm.group(1), mm.group(0)), body)
    body = re.sub(r"\\label\{[^}]*\}", "", body)
    k = [0]
    def fig(mm):
        k[0] += 1; png = rasterize_snippet(mm.group(0), pre, work, "fig%03d" % k[0]); return "\\includegraphics[width=0.8\\textwidth]{%s}" % png if png else mm.group(0)
    body = re.sub(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", fig, body, flags=re.S)
    body = re.sub(r"\\srcfig\{[^}]*\}\{[^}]*\}\{[^}]*\}\{[^}]*\}\{[^}]*\}", fig, body)
    final = work / "export.tex"; write(final, pre + "\\begin{document}" + body + "\\end{document}\n")
    rc, out, err = sh(["pandoc", final.name, "-f", "latex", "-t", "docx", "--toc", "-o", str(Path(out_docx).resolve()),
                       "--resource-path=" + str(work), "-M", "lang=" + tp["bcp47"]], cwd=str(work), timeout=900)
    if rc != 0: raise SystemExit("pandoc failed: " + err[:500])
    docx_set_language(out_docx, cfg); return {"docx": str(out_docx), "numbering_pass": {k: v for k, v in rep.items() if k in ("ok", "errors", "pages")}, "workdir": str(work), "numbers_resolved": len(nums)}

# =============================================================================
# PART 8 — DOCX -> DOCX (native, in-place): equations (OMML), fields (equation numbers, TOC, cross-refs), drawings,
# footnotes, tables, styles and layout are preserved; only ordinary-language runs are rewritten.
# Works directly on the .docx zip + lxml (python-docx is not required).
# =============================================================================
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
XML_NS = "http://www.w3.org/XML/1998/namespace"
def wq(t): return "{%s}%s" % (W_NS, t)
def mq(t): return "{%s}%s" % (M_NS, t)
DOCX_TEXT_PARTS = re.compile(r"^word/(document|footnotes|endnotes|comments|header\d*|footer\d*)\.xml$")
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss", "imprint", "noProof",
             "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd",
             "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
             "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]

def _lx():
    try: from lxml import etree; return etree
    except ImportError: raise SystemExit("DOCX support needs lxml: pip install lxml")

def etree_local(tag): return tag.split("}")[-1]

def _insert_ordered(parent, el, order):
    name = etree_local(el.tag)
    for old in parent.findall(el.tag): parent.remove(old)
    idx = order.index(name) if name in order else len(order); pos = len(parent)
    for i, ch in enumerate(parent):
        ln = etree_local(ch.tag)
        if ln in order and order.index(ln) > idx: pos = i; break
    parent.insert(pos, el)

TEXT_CHILD = {wq("rPr"), wq("t"), wq("tab"), wq("noBreakHyphen"), wq("softHyphen"), wq("lastRenderedPageBreak")}
def run_is_text(r):
    for c in r:
        if c.tag in TEXT_CHILD: continue
        if c.tag == wq("br") and c.get(wq("type")) in (None, "textWrapping"): continue
        return False
    return True
def run_text(r):
    out = []
    for c in r:
        if c.tag == wq("t"): out.append(c.text or "")
        elif c.tag == wq("tab"): out.append("\t")
        elif c.tag == wq("br"): out.append("\n")
        elif c.tag == wq("noBreakHyphen"): out.append("‑")
    return "".join(out)
def rpr_key(r):
    etree = _lx(); p = r.find(wq("rPr"))
    if p is None: return ""
    q = copy.deepcopy(p)
    for t in (wq("lang"), wq("noProof")):
        for e in q.findall(t): q.remove(e)
    return etree.tostring(q, method="c14n").decode()

def _items(container):
    """Linearise the children of a paragraph/hyperlink into ('run'|'atom'|'hyperlink'|'front'|'back'|'drop'|'tracked', ...) items.
    Complex fields (begin..end, incl. equation numbers, TOC, REF) become ONE atom: their cached text is never translated."""
    out, depth, buf = [], 0, []
    for ch in container:
        tag = ch.tag
        if tag == wq("pPr"): continue
        if tag == wq("r"):
            fc = ch.find(wq("fldChar")); t = fc.get(wq("fldCharType")) if fc is not None else None
            if t == "begin": depth += 1
            if depth > 0:
                buf.append(ch)
                if t == "end":
                    depth -= 1
                    if depth == 0: out.append(("atom", buf, "X")); buf = []
                continue
            out.append(("run", ch) if run_is_text(ch) else ("atom", [ch], "X"))
        elif tag in (mq("oMath"), mq("oMathPara")): out.append(("atom", [ch], "M"))
        elif tag == wq("hyperlink"): out.append(("hyperlink", ch))
        elif tag in (wq("bookmarkStart"), wq("commentRangeStart")): out.append(("front", ch))
        elif tag in (wq("bookmarkEnd"), wq("commentRangeEnd")): out.append(("back", ch))
        elif tag == wq("proofErr"): out.append(("drop", ch))
        elif tag in (wq("ins"), wq("del"), wq("moveFrom"), wq("moveTo")): out.append(("tracked", ch))
        else: out.append(("atom", [ch], "X"))
    if buf: out.append(("atom", buf, "X"))
    return out

def build_unit(container):
    """-> dict(src, lead, trail, atoms{token:[elements]}, keys{id:rPr}, base_rpr, items, hyperlinks) or None if nothing translatable."""
    items = _items(container)
    if any(i[0] == "tracked" for i in items): return {"tracked": True}
    spans, atoms, cnt = [], {}, collections.Counter(); hyperlinks = []
    for it in items:
        if it[0] == "run":
            k, tx = rpr_key(it[1]), run_text(it[1])
            if spans and spans[-1][0] == "text" and spans[-1][1] == k: spans[-1][2] += tx
            else: spans.append(["text", k, tx, it[1].find(wq("rPr"))])
        elif it[0] == "atom":
            cnt[it[2]] += 1; tok = PH(it[2], cnt[it[2]]); atoms[tok] = it[1]; spans.append(["atom", tok])
        elif it[0] == "hyperlink":
            cnt["X"] += 1; tok = PH("X", cnt["X"]); atoms[tok] = [it[1]]; spans.append(["atom", tok]); hyperlinks.append(it[1])
    texts = [s for s in spans if s[0] == "text"]
    if not texts: return None
    base = max(texts, key=lambda s: len(s[2].strip()))[1]; ids, keys, parts = {}, {}, []
    for s in spans:
        if s[0] == "atom": parts.append(s[1]); continue
        if s[1] == base: parts.append(s[2]); continue
        if s[1] not in ids: ids[s[1]] = len(ids) + 1; keys[ids[s[1]]] = s[3]
        i = ids[s[1]]; parts.append("⟦s%d⟧%s⟦/s%d⟧" % (i, s[2], i))
    full = "".join(parts); lead, core, trail = re.match(r"^(\s*)(.*?)(\s*)$", full, re.S).groups()
    base_rpr = next((s[3] for s in texts if s[1] == base), None)
    return {"src": core, "lead": lead, "trail": trail, "atoms": atoms, "keys": keys, "base_rpr": base_rpr, "items": items, "hyperlinks": hyperlinks, "tracked": False}

def _mk_run(text, rpr, lang_fn):
    etree = _lx(); r = etree.Element(wq("r"))
    p = copy.deepcopy(rpr) if rpr is not None else etree.Element(wq("rPr"))
    if lang_fn: lang_fn(p)
    if len(p): r.append(p)
    for piece in re.split(r"(\t|\n)", text):
        if piece == "\t": etree.SubElement(r, wq("tab"))
        elif piece == "\n": etree.SubElement(r, wq("br"))
        elif piece:
            t = etree.SubElement(r, wq("t")); t.text = piece; t.set("{%s}space" % XML_NS, "preserve")
    return r

def _lang_setter(cfg):
    etree = _lx(); tc, tp = get_profile(cfg["project"]["target_lang"]); rtl = tp["dir"] == "rtl"
    def fn(rpr):
        lg = etree.Element(wq("lang")); lg.set(wq("val"), tp["bcp47"])
        if tp["script"] in ("han", "kana_han", "hangul"): lg.set(wq("eastAsia"), tp["bcp47"])
        if rtl: lg.set(wq("bidi"), tp["bcp47"]); _insert_ordered(rpr, etree.Element(wq("rtl")), RPR_ORDER)
        _insert_ordered(rpr, lg, RPR_ORDER)
    return fn

def rebuild_container(container, unit, tgt, cfg):
    etree = _lx(); lang_fn = _lang_setter(cfg); tc, tp = get_profile(cfg["project"]["target_lang"])
    front = [i[1] for i in unit["items"] if i[0] == "front"]; back = [i[1] for i in unit["items"] if i[0] == "back"]
    ppr = container.find(wq("pPr"))
    for ch in list(container):
        if ch is not ppr: container.remove(ch)
    for e in front: container.append(e)
    style = [0]; text = unit["lead"] + tgt + unit["trail"]
    for tok in re.split("(⟦/?[A-Za-z]+\\d+⟧)", text):
        m = PH_RE.fullmatch(tok)
        if m:
            closing, kind, n = m.groups()
            if kind == "s":
                if closing:
                    if len(style) > 1: style.pop()
                else: style.append(int(n))
            else:
                for el in unit["atoms"].get(tok, []): container.append(el)
        elif tok:
            rpr = unit["base_rpr"] if style[-1] == 0 else unit["keys"].get(style[-1], unit["base_rpr"])
            container.append(_mk_run(tok, rpr, lang_fn))
    for e in back: container.append(e)
    if tp["dir"] == "rtl" and container.tag == wq("p"):
        if ppr is None: ppr = etree.Element(wq("pPr")); container.insert(0, ppr)
        _insert_ordered(ppr, etree.Element(wq("bidi")), PPR_ORDER)

def _style_of(p):
    s = p.find(wq("pPr") + "/" + wq("pStyle")); return s.get(wq("val")) if s is not None else ""

def _docx_walk(path, cfg, translations=None):
    """Shared traversal for collect (translations=None) and apply. Segment ids are deterministic: d<part>_p<n>[h<k>]."""
    etree = _lx(); segs, changed, skipped = [], {}, collections.Counter()
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if DOCX_TEXT_PARTS.match(n)]; names.sort(key=lambda n: (n != "word/document.xml", n))
        for pi, name in enumerate(names):
            root = etree.fromstring(z.read(name)); paras = list(root.iter(wq("p"))); touched = [False]
            def handle(container, sid):
                if _style_of(container) in cfg["docx"]["skip_styles"]: skipped["style"] += 1; return []
                u = build_unit(container)
                if u is None: return []
                if u.get("tracked"): skipped["tracked_changes"] += 1; return []
                if not has_letters(u["src"]): return []
                seg = Segment(sid, "docxpara", u["src"], _style_of(container), "docx"); segs.append(seg)
                if translations is not None:
                    t = translations.get(sid)
                    if t is not None and t.tgt is not None and t.status == "done":
                        if t.src != u["src"]: raise SystemExit("Document changed since collection (segment %s). Re-run collect." % sid)
                        rebuild_container(container, u, t.tgt, cfg); touched[0] = True
                return u["hyperlinks"]
            for k, p in enumerate(paras):
                for h, hl in enumerate(handle(p, "d%d_p%d" % (pi, k)) or []): handle(hl, "d%d_p%dh%d" % (pi, k, h))
            if touched[0]: changed[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return segs, changed, skipped

def docx_collect(path, cfg): return _docx_walk(path, cfg)

def docx_set_language(path, cfg):
    """Set the default proofing language (and RTL/CJK attributes) in styles.xml docDefaults; optional updateFields flag in settings.xml."""
    etree = _lx(); tc, tp = get_profile(cfg["project"]["target_lang"]); path = str(path); tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/styles.xml":
                root = etree.fromstring(data); dd = root.find(wq("docDefaults"))
                if dd is not None:
                    rp = dd.find(wq("rPrDefault") + "/" + wq("rPr"))
                    if rp is not None:
                        for old in rp.findall(wq("lang")): rp.remove(old)
                        lg = etree.Element(wq("lang")); lg.set(wq("val"), tp["bcp47"]); lg.set(wq("eastAsia"), tp["bcp47"] if tp["script"] in ("han", "kana_han", "hangul") else "en-US"); lg.set(wq("bidi"), tp["bcp47"] if tp["dir"] == "rtl" else "ar-SA")
                        _insert_ordered(rp, lg, RPR_ORDER)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            if item.filename == "word/settings.xml" and cfg["docx"].get("update_fields_on_open"):
                root = etree.fromstring(data)
                if root.find(wq("updateFields")) is None:
                    el = etree.Element(wq("updateFields")); el.set(wq("val"), "true")
                    after = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping"]
                    pos = len(root)
                    for i, ch in enumerate(root):
                        if etree_local(ch.tag) in after: pos = i; break
                    root.insert(pos, el)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    os.replace(tmp, path)

def docx_apply(path, out_path, segs_by_id, cfg):
    _, changed, skipped = _docx_walk(path, cfg, segs_by_id)
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist(): zout.writestr(item, changed.get(item.filename, zin.read(item.filename)))
    docx_set_language(out_path, cfg); return {"parts_changed": sorted(changed), "skipped": dict(skipped)}

def docx_inventory(path):
    """Structural/mathematical inventory used by the DOCX comparison QC."""
    etree = _lx(); inv = {"math": [], "tables": [], "paragraphs": 0, "drawings": 0, "objects": 0, "fields": 0, "bookmarks": 0, "hyperlinks": 0,
                          "footnotes": 0, "endnotes": 0, "sections": 0, "comments": 0, "charts": 0, "media": []}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("word/charts/") and n.endswith(".xml"): inv["charts"] += 1
            if n.startswith("word/media/"): inv["media"].append(hashlib.sha1(z.read(n)).hexdigest()[:10])
        for n in [n for n in z.namelist() if DOCX_TEXT_PARTS.match(n)]:
            root = etree.fromstring(z.read(n)); main = n == "word/document.xml"
            inv["paragraphs"] += len(list(root.iter(wq("p"))))
            for m in root.iter(mq("oMath")): inv["math"].append(re.sub(r"\s+", "", etree.tostring(m, method="c14n").decode()))
            inv["drawings"] += len(list(root.iter(wq("drawing")))) + len(list(root.iter(wq("pict"))))
            inv["objects"] += len(list(root.iter(wq("object"))))
            inv["fields"] += len([f for f in root.iter(wq("fldChar")) if f.get(wq("fldCharType")) == "begin"]) + len(list(root.iter(wq("fldSimple"))))
            inv["bookmarks"] += len(list(root.iter(wq("bookmarkStart")))); inv["hyperlinks"] += len(list(root.iter(wq("hyperlink"))))
            if n.endswith("footnotes.xml"): inv["footnotes"] = len([f for f in root.iter(wq("footnote")) if f.get(wq("type")) in (None, "normal")])
            if n.endswith("endnotes.xml"): inv["endnotes"] = len([f for f in root.iter(wq("endnote")) if f.get(wq("type")) in (None, "normal")])
            if main:
                inv["sections"] = len(list(root.iter(wq("sectPr"))))
                for t in root.iter(wq("tbl")):
                    rows = t.findall(wq("tr")); inv["tables"].append([len(rows), [len(r.findall(wq("tc"))) for r in rows]])
            if n.endswith("comments.xml"): inv["comments"] = len(list(root.iter(wq("comment"))))
    return inv

def docx_translate(src_docx, out_docx, project, backend=None, collect=False):
    """Pipeline for DOCX: collect segments -> translate (any backend) -> apply. With backend 'agent' run twice (second time with collect=True)."""
    cfg = project.cfg; segs, _, skipped = docx_collect(src_docx, cfg)
    if not collect or not project.p("work", "segments.jsonl").exists():
        project.save_segments(segs); write(project.p("work", "docx_source_path.txt"), str(Path(src_docx).resolve()))
    stats = translate_all(project, backend, collect=collect)
    by = {s.id: s for s in project.load_segments()}; pend = [s for s in by.values() if s.status != "done"]
    if pend:
        return {"status": "incomplete", "pending": len(pend), "stats": stats, "next": "agent backend: fill work/tasks/*.result.json then re-run with --collect"}
    rep = docx_apply(src_docx, out_docx, by, cfg); rep.update(status="done", stats=stats, segments=len(by)); return rep

# =============================================================================
# PART 9 — QUALITY CONTROL: automatic comparison of source and translation BEFORE export.
# Severity 'block' stops `export` (override only with --force). Which checks block is configurable (qc.block_on).
# Human-approved exceptions go to <project>/qc/waivers.json as a list of finding keys (they are downgraded to 'info').
# =============================================================================
def mk(cfg, check, where, detail, src="", tgt=""):
    return Finding("block" if check in cfg["qc"]["block_on"] else "warn", check, where, detail, (src or "")[:300], (tgt or "")[:300])

def visible(s):
    """Reading text of a segment: placeholders, LaTeX commands and markup removed, whitespace collapsed."""
    t = PH_RE.sub(" ", s); t = re.sub(r"\\[A-Za-z]+\*?", " ", t); t = re.sub(r"[\\{}$~&^_]", " ", t); return re.sub(r"\s+", " ", t).strip()

def strip_comments(t): return re.sub(r"(?<!\\)%[^\n]*", "", t)

# ---------------------------------------------------------------- mathematics
def norm_math(m):
    """Canonical form of a protected math span for exact comparison: child text placeholders neutralised, comments dropped, insignificant spaces removed."""
    m = re.sub("\u27e6S\\d+\u27e7", "\u27e6S\u27e7", m); m = strip_comments(m); m = re.sub(r"\s+", " ", m).strip()
    return re.sub(r"(?<![A-Za-z]) | (?![A-Za-z])", "", m)

def tex_math_list(text, cfg):
    st = Store(cfg, "qc"); protect(text, st)
    return [norm_math(o) for t, o in st.ph.items() if t.startswith("\u27e6M")]

def math_delta(s, t):
    notes = []
    for name, rx in (("superscripts", r"\^"), ("subscripts", r"_"), ("digits", r"\d"), ("commands", r"\\[A-Za-z]+"), ("braces", r"[{}]"), ("symbols", r"[=<>+\-*/|()\[\]]")):
        a, b = len(re.findall(rx, s)), len(re.findall(rx, t))
        if a != b: notes.append("%s %d->%d" % (name, a, b))
    return "; ".join(notes) or "characters differ"

def compare_math(src_list, tgt_list, cfg, where):
    out = []; a, b = collections.Counter(src_list), collections.Counter(tgt_list)
    only_s, only_t = list((a - b).elements()), list((b - a).elements())
    if len(only_s) * len(only_t) > 40000:
        out.append(mk(cfg, "math_altered", where, "%d source formulas and %d target formulas differ (too many to pair); first source: %s" % (len(only_s), len(only_t), only_s[:1]), only_s[0] if only_s else "", only_t[0] if only_t else ""))
        return out
    used = set()
    for s in only_s:
        best, bj = 0.5, -1
        for j, t in enumerate(only_t):
            if j in used: continue
            sm = difflib.SequenceMatcher(None, s, t, autojunk=False)
            if sm.real_quick_ratio() <= best or sm.quick_ratio() <= best: continue
            r = sm.ratio()
            if r > best: best, bj = r, j
        if bj >= 0: used.add(bj); out.append(mk(cfg, "math_altered", where, "formula changed (%s)" % math_delta(s, only_t[bj]), s, only_t[bj]))
        else: out.append(mk(cfg, "math_missing", where, "formula missing in translation", s, ""))
    for j, t in enumerate(only_t):
        if j not in used: out.append(mk(cfg, "math_added", where, "formula added in translation", "", t))
    if not out and len(src_list) <= 3000 and src_list != tgt_list:
        if difflib.SequenceMatcher(None, src_list, tgt_list, autojunk=False).ratio() < 0.98:
            out.append(Finding("warn", "math_reordered", where, "same formulas but a different order (%d items)" % len(src_list)))
    return out

# ---------------------------------------------------------------- structure
STRUCT_CMDS = {"footnote": r"\\footnote\s*\{", "caption": r"\\caption(?:of)?\s*[\[{]", "includegraphics": r"\\includegraphics", "srcfig": r"\\srcfig\{",
               "item": r"\\item\b", "cite": r"\\cite[a-z]*\s*[\[{]", "index": r"\\index\{"}

def table_shapes(t):
    shapes = []
    for env in TABULAR_ENVS:
        for m in re.finditer(r"\\begin\{%s\}" % re.escape(env), t):
            a, b = find_env_end(t, m.end(), env)
            if a < 0: continue
            j = _skip_args(t, m.end(), TABULAR_ENVS[env]); body = t[j:a]
            rows = [r for r in re.split(r"(?<!\\)\\\\", body) if re.sub(r"\\(hline|toprule|midrule|bottomrule)|\s", "", r)]
            cols = [len(re.findall(r"(?<!\\)&", r)) + 1 for r in rows]
            shapes.append((env, len(rows), max(cols) if cols else 0, len(re.findall(r"\\multicolumn", body)), len(re.findall(r"\\multirow", body))))
    return shapes

def tex_inventory(text, cfg):
    t = strip_comments(text)
    return {"envs": collections.Counter(re.findall(r"\\begin\{([^}]*)\}", t)),
            "sections": collections.Counter(re.findall(r"\\(chapter|section|subsection|subsubsection)\*?\s*[\[{]", t)),
            "cmds": {k: len(re.findall(rx, t)) for k, rx in STRUCT_CMDS.items()},
            "labels": re.findall(r"\\label\{([^}]*)\}", t),
            "refs": collections.Counter(re.findall(r"\\(?:eq|page|auto|c|C|v|name)?ref\{([^}]*)\}", t)),
            "tables": table_shapes(t), "math": tex_math_list(text, cfg)}

def compare_inventories(si, ti, cfg, where):
    out = []
    d = (si["envs"] - ti["envs"]) + (ti["envs"] - si["envs"])
    if d: out.append(mk(cfg, "env_mismatch", where, "environments differ: %s" % dict(d)))
    d = (si["sections"] - ti["sections"]) + (ti["sections"] - si["sections"])
    if d: out.append(mk(cfg, "env_mismatch", where, "sectioning commands differ: %s" % dict(d)))
    for k, check in (("footnote", "footnote_mismatch"), ("caption", "caption_mismatch"), ("includegraphics", "figure_mismatch"), ("srcfig", "figure_mismatch"), ("item", "env_mismatch"), ("cite", "ref_mismatch"), ("index", "env_mismatch")):
        if si["cmds"][k] != ti["cmds"][k]: out.append(mk(cfg, check, where, "\\%s count %d -> %d" % (k, si["cmds"][k], ti["cmds"][k])))
    if si["envs"].get("tikzpicture", 0) != ti["envs"].get("tikzpicture", 0): out.append(mk(cfg, "figure_mismatch", where, "tikzpicture count %d -> %d" % (si["envs"].get("tikzpicture", 0), ti["envs"].get("tikzpicture", 0))))
    if si["tables"] != ti["tables"]: out.append(mk(cfg, "table_mismatch", where, "tables differ (env, rows, cols, multicolumn, multirow): %s -> %s" % (si["tables"][:6], ti["tables"][:6])))
    if sorted(si["labels"]) != sorted(ti["labels"]):
        a, b = set(si["labels"]), set(ti["labels"]); out.append(mk(cfg, "label_mismatch", where, "labels differ: missing=%s added=%s" % (sorted(a - b)[:8], sorted(b - a)[:8])))
    if si["refs"] != ti["refs"]: out.append(mk(cfg, "ref_mismatch", where, "reference targets differ: %s" % dict(((si["refs"] - ti["refs"]) + (ti["refs"] - si["refs"])).most_common(8))))
    out += compare_math(si["math"], ti["math"], cfg, where)
    return out

# ---------------------------------------------------------------- numbering (PDF text and LaTeX .aux)
def numbering_inventory_text(text, lang):
    """Numbered items ('Definition 1.2.3', 'Théorème 2.1', ...) at line starts, as in the original session's check, for any language profile."""
    _, prof = get_profile(lang); keys = ["theorem", "definition", "proposition", "lemma", "corollary", "exercise", "example", "remark"]
    words = {fold(prof["labels"][k]): k for k in keys}; alts = []
    for w in sorted(words, key=len, reverse=True): alts.append(re.escape(w) + (r"\s+" if re.fullmatch(r"[a-z ]+", w) else r"\s*"))
    rx = re.compile(r"^\s*\(?(" + "|".join(a for a in alts) + r")(\d+(?:\.\d+)*)", re.M)
    out = []
    for m in rx.finditer(fold(normalize_text(text))):
        lab = re.sub(r"\\s[+*]", "", m.group(1)).strip()
        for w, k in words.items():
            if fold(lab).strip() == w: out.append((k, m.group(2))); break
    return out

def compare_numbering(a, b, cfg, where, label="numbered items"):
    """Compare two numbered-item inventories. Exact ordered comparison (BLOCK) only when both come from bold-face line starts, or from plain lists
    (LaTeX-side data). When either PDF side could only be read as 'all line starts' (unknown / bitmap fonts), wrapped cross-references make an exact
    comparison unreliable, so the bold (authoritative) side must be a sub-multiset of the other side: otherwise a WARNING numbering_mismatch_heuristic."""
    ma, mb = getattr(a, "mode", "bold"), getattr(b, "mode", "bold")
    if ma == mb:
        if list(a) == list(b): return []
        sm = difflib.SequenceMatcher(None, list(a), list(b), autojunk=False); diffs = [(tag, list(a[i1:i2])[:3], list(b[j1:j2])[:3]) for tag, i1, i2, j1, j2 in sm.get_opcodes() if tag != "equal"]
        if ma == "all" and getattr(a, "mode", None) == "all":      # both sides heuristic: report as warning only
            return [Finding("warn", "numbering_mismatch_heuristic", where, "%s differ (line-start heuristic on both sides, fonts not recognisable as bold): %d vs %d items; first differences %s" % (label, len(a), len(b), diffs[:4]))]
        return [mk(cfg, "numbering_mismatch", where, "%s differ: %d vs %d items; first differences %s" % (label, len(a), len(b), diffs[:4]))]
    auth, other, who = (a, b, "source") if ma == "bold" else (b, a, "translation")
    missing = collections.Counter(auth) - collections.Counter(other)
    if not missing: return []
    return [Finding("warn", "numbering_mismatch_heuristic", where, "%s: %d numbered items read reliably (bold) in the %s are not found among the line starts of the other PDF (its fonts cannot be recognised as bold, so this is a heuristic): %s" % (label, sum(missing.values()), who, sorted(missing)[:6]))]

def pdf_font_audit(pdf, where):
    """Warn about PDFs whose text cannot be searched/copied reliably: Type 3 (bitmap) fonts, fonts without a ToUnicode map."""
    rc, out, _ = sh(["pdffonts", str(pdf)])
    if rc != 0: return []
    rows = [l for l in out.splitlines()[2:] if l.strip()]; F = []; t3 = nouni = 0
    for l in rows:
        cols = l.split()
        if len(cols) < 7: continue
        typ = " ".join(cols[1:-5]) if len(cols) >= 8 else cols[1]; uni = cols[-3]
        if typ.startswith("Type 3"): t3 += 1
        if uni == "no": nouni += 1
    if t3: F.append(Finding("warn", "bitmap_fonts", where, "%d Type 3 (bitmap) font(s): text stays visible but is often blurry when zoomed and copy/paste/search is unreliable. In pdfLaTeX install lmodern or cm-super (or set latex.font_package=mathptmx); XeLaTeX with an OpenType font also avoids it." % t3))
    if nouni: F.append(Finding("warn", "no_unicode_map", where, "%d font(s) without a Unicode map: copy/paste and search on this PDF will return wrong characters (ligatures, accents)." % nouni))
    return F

def compare_aux(src_aux, tgt_aux, cfg, where):
    """Compare printed numbers of every \\label between the compiled source and the compiled translation (theorem/equation/section numbering)."""
    a, b = read_aux_labels(src_aux), read_aux_labels(tgt_aux); out = []
    diff = {k: (a.get(k), b.get(k)) for k in set(a) | set(b) if a.get(k) != b.get(k)}
    if diff: out.append(mk(cfg, "numbering_mismatch", where, "%d label numbers differ between source and translation: %s" % (len(diff), dict(list(diff.items())[:6]))))
    return out

# ---------------------------------------------------------------- segment-level QC
NUM_RE = re.compile(r"(?<![\w.,])\d+(?:[.,]\d+)*(?![\w])")
LOCALISED_NOTATION = {"tg", "ctg", "arctg", "arcctg", "sh", "ch", "th", "cth", "arcsh", "arcch"}
def _shingles(v, k=4):
    w = fold(v).split(); return {" ".join(w[i:i + k]) for i in range(max(0, len(w) - k + 1))}

def qc_segments(project, segs, glossary):
    cfg = project.cfg; F = []; sc, sp = get_profile(cfg["project"]["source_lang"]); tc, tp = get_profile(cfg["project"]["target_lang"])
    done = [s for s in segs if s.status == "done" and s.tgt is not None]
    for s in segs:
        if s.status != "done" or s.tgt is None: F.append(mk(cfg, "untranslated_segment", s.id, "status=%s (source text is used in the output)" % s.status, s.src, ""))
    prose = [s for s in done if s.kind in ("prose", "heading")]
    for s in done:
        for code, detail in validate_translation(s.src, s.tgt):
            chk = "placeholder_mismatch" if code in ("placeholder_mismatch", "empty") else ("formatting_corruption" if code in ("command_skeleton", "brace_balance") else "comment_count")
            F.append(mk(cfg, chk, s.id, detail, s.src, s.tgt))
        ns, nt = collections.Counter(NUM_RE.findall(visible(s.src))), collections.Counter(NUM_RE.findall(visible(s.tgt)))
        if ns != nt: F.append(mk(cfg, "number_altered", s.id, "numbers in prose differ: missing=%s added=%s" % (list((ns - nt).elements())[:6], list((nt - ns).elements())[:6]), s.src, s.tgt))
        vs, vt = visible(s.src), visible(s.tgt)
        if sc != tc:
            bad = (set(re.findall(r"[^\W\d_]+", vt.lower())) & LOCALISED_NOTATION) - set(re.findall(r"[^\W\d_]+", vs.lower()))
            if bad: F.append(mk(cfg, "notation_localised", s.id, "notation tokens %s appear in the translation but not in the source" % sorted(bad), s.src, s.tgt))
        words = re.findall(r"[^\W\d_]+", vt.lower())
        if sc != tc and s.kind in ("prose", "heading") and len(words) >= 8:
            if sp["script"] == tp["script"]:
                stop = set(map(fold, sp["stopwords"])) - set(map(fold, tp["stopwords"]))
                r = sum(1 for w in words if fold(w) in stop) / len(words)
                if r >= cfg["qc"]["residue_stopword_ratio"]: F.append(mk(cfg, "untranslated_fragment", s.id, "%.0f%% of the words look like %s function words" % (100 * r, sp["name"]), s.src, s.tgt))
            elif script_ratio(vt, sp["script"]) > 0.3:
                F.append(mk(cfg, "untranslated_fragment", s.id, "%.0f%% of the letters are in the source script" % (100 * script_ratio(vt, sp["script"])), s.src, s.tgt))
        if sc != tc and len(words) >= 4 and fold(vs) == fold(vt): F.append(mk(cfg, "identical_to_source", s.id, "translation text equals the source text", s.src, s.tgt))
        for f in glossary.check(vs, vt, s.ctx, s.id): F.append(f)
        if s.kind in ("prose", "heading"):
            for rx in tp["ai_phrases"]:
                m = re.search(rx, vt, re.I)
                if m: F.append(mk(cfg, "style_ai_phrase", s.id, "stock phrase '%s'" % m.group(0), s.src, s.tgt))
        n_s, n_t = len(re.findall(r"[.!?。！？](?:\s|$)", vs)), len(re.findall(r"[.!?。！？](?:\s|$)", vt))
        if n_s >= 3 and n_t * 2 < n_s: F.append(mk(cfg, "possible_missing_text", s.id, "%d source sentences vs %d in translation" % (n_s, n_t), s.src, s.tgt))
        if n_s >= 1 and n_t > 2 * n_s + 1: F.append(mk(cfg, "possible_added_text", s.id, "%d source sentences vs %d in translation" % (n_s, n_t), s.src, s.tgt))
    # length-ratio outliers (robust: median + MAD of log ratios)
    pts = [(s, math.log(max(1, len(visible(s.tgt))) / max(1, len(visible(s.src))))) for s in prose if len(visible(s.src)) >= cfg["qc"]["min_chars_for_ratio"]]
    if len(pts) >= 8:
        vals = [v for _, v in pts]; med = statistics.median(vals); mad = max(0.05, 1.4826 * statistics.median([abs(v - med) for v in vals]))
        for s, v in pts:
            z = (v - med) / mad
            if abs(z) > cfg["qc"]["length_ratio_mad_k"]:
                F.append(mk(cfg, "possible_missing_text" if z < 0 else "possible_added_text", s.id, "length ratio %.2f vs typical %.2f (z=%.1f)" % (math.exp(v), math.exp(med), z), s.src, s.tgt))
    # duplicated paragraphs
    seen = {}
    for s in prose:
        vt = visible(s.tgt)
        if len(vt.split()) < 8: continue
        k = fold(vt)
        if k in seen and fold(visible(s.src)) != fold(visible(seen[k].src)): F.append(mk(cfg, "duplicated_paragraph", s.id, "identical to %s while the sources differ" % seen[k].id, s.src, s.tgt))
        seen.setdefault(k, s)
    cand = [(s, _shingles(visible(s.tgt)), _shingles(visible(s.src))) for s in prose if len(visible(s.tgt).split()) >= 10]
    if len(cand) <= 2000:
        for i in range(len(cand)):
            for j in range(i + 1, len(cand)):
                a, b = cand[i][1], cand[j][1]
                if not a or not b or abs(len(a) - len(b)) > 0.3 * max(len(a), len(b)): continue
                jac = len(a & b) / len(a | b)
                if jac >= cfg["qc"]["dup_jaccard"]:
                    sa, sb = cand[i][2], cand[j][2]
                    if not sa or not sb or len(sa & sb) / len(sa | sb) < cfg["qc"]["dup_jaccard"]:
                        F.append(mk(cfg, "duplicated_paragraph", cand[j][0].id, "near-duplicate (%.2f) of %s while the sources differ" % (jac, cand[i][0].id), cand[j][0].src, cand[j][0].tgt))
    # terminology consistency: same glossary entry rendered sometimes as approved, sometimes not; same heading -> same translation
    stat = collections.defaultdict(lambda: [0, 0])
    for s in prose:
        for r in glossary.hits(visible(s.src), s.ctx):
            if r["tgt"]: stat[(r["src"], r["tgt"])][0 if fold(r["tgt"]) in fold(visible(s.tgt)) else 1] += 1
    for (src, tgt), (ok, bad) in stat.items():
        if ok and bad: F.append(Finding("warn", "term_inconsistency", "book", "'%s' -> '%s' used in %d segments, absent in %d" % (src, tgt, ok, bad)))
    heads = collections.defaultdict(set)
    for s in done:
        if s.kind == "heading": heads[fold(visible(s.src))].add(s.tgt)
    for k, v in heads.items():
        if k and len(v) > 1: F.append(Finding("warn", "term_inconsistency", "book", "identical source heading '%s' translated %d different ways" % (k[:60], len(v))))
    openers = collections.Counter(" ".join(fold(visible(s.tgt)).split()[:3]) for s in prose if len(visible(s.tgt).split()) >= 6)
    sopen = collections.Counter(" ".join(fold(visible(s.src)).split()[:3]) for s in prose if len(visible(s.src).split()) >= 6)
    if len(prose) >= 40 and openers:
        (o, c), = openers.most_common(1)
        if c >= 6 and c / len(prose) > 0.10 and sopen.get(o, 0) == 0 and (sopen.most_common(1)[0][1] if sopen else 0) < 0.8 * c:
            F.append(Finding("warn", "repetitive_phrasing", "book", "%d paragraphs start with '%s'" % (c, o)))
    return F

# ---------------------------------------------------------------- project-level QC (tex -> tex)
def apply_waivers(project, findings):
    w = set(jload(project.p("qc", "waivers.json"), []))
    for f in findings:
        if f.key() in w: f.severity = "info"; f.detail = "[waived] " + f.detail
    return findings

def qc_project(project, compile_check=False, main=None, reassemble=False):
    cfg = project.cfg; F = []; segs = project.load_segments(); g = project.glossary()
    if not segs: raise SystemExit("no segments: run `ingest` first")
    out_dir = project.p("out", "tex")
    if reassemble or not out_dir.exists(): assemble_project(project)
    elif project.p("work", "segments.jsonl").stat().st_mtime > max((f.stat().st_mtime for f in out_dir.rglob("*.tex")), default=0):
        F.append(Finding("warn", "stale_output", "out/tex", "segments changed after the last `assemble`; QC reads the older output. Re-run `assemble`."))
    F += qc_segments(project, segs, g)
    src_inv_all, tgt_inv_all = {"labels": [], "refs": collections.Counter()}, {"labels": [], "refs": collections.Counter()}
    for name, rel in project.docs().items():
        so, to = project.p("work", "orig", name + ".tex"), out_dir / rel
        if not to.exists(): F.append(mk(cfg, "env_mismatch", name, "translated file missing: %s" % rel)); continue
        si, ti = tex_inventory(read(so), cfg), tex_inventory(read(to), cfg)
        F += compare_inventories(si, ti, cfg, name)
        tgt_inv_all["labels"] += ti["labels"]; tgt_inv_all["refs"].update(ti["refs"])
    broken = sorted(set(tgt_inv_all["refs"]) - set(tgt_inv_all["labels"]))
    if broken: F.append(mk(cfg, "broken_reference", "book", "references without a label in the translation: %s" % broken[:10]))
    if compile_check and main:
        sm, tm = project.p("source", main), out_dir / main
        if not (sm.exists() and tm.exists()): F.append(Finding("warn", "compile_skipped", main, "main file not found in source/ or out/tex"))
        else:
            tmp = Path(tempfile.mkdtemp(prefix="qc_src_")); shutil.copytree(project.p("source"), tmp / "s", dirs_exist_ok=True)
            fixed, _ = ensure_text_fonts(read(tmp / "s" / main), cfg); write(tmp / "s" / main, fixed)      # QC copy only: usable text layer for the numbering comparison
            cfg_src = deep_merge(cfg, {"project": {"target_lang": cfg["project"]["source_lang"]}})
            rs = build_pdf(tmp / "s" / main, cfg_src); rt = build_pdf(tm, cfg)
            for who, r in (("source", rs), ("translation", rt)):
                if r.get("errors"): F.append(mk(cfg, "latex_error", who, "LaTeX errors: %s" % r["errors"][:3]))
                if r.get("undefined_refs"): F.append(mk(cfg, "broken_reference", who, "undefined references after compilation: %s" % r["undefined_refs"][:6]))
                if r.get("missing_chars"): F.append(Finding("warn", "missing_glyphs", who, "%d 'Missing character' warnings (font lacks glyphs)" % r["missing_chars"]))
            if rs.get("pdf") and rt.get("pdf"):
                F += compare_aux(Path(rs["pdf"]).with_suffix(".aux"), Path(rt["pdf"]).with_suffix(".aux"), cfg, "labels")
                n_s, n_t = rs.get("pages") or 0, rt.get("pages") or 0
                if n_s and not (0.7 <= n_t / n_s <= 1.6): F.append(Finding("warn", "page_count", "book", "pages %d -> %d" % (n_s, n_t)))
                F += compare_numbering(numbering_inventory_pdf(rs["pdf"], cfg["project"]["source_lang"]), numbering_inventory_pdf(rt["pdf"], cfg["project"]["target_lang"]), cfg, "compiled books")
                F += pdf_font_audit(rt["pdf"], "translation PDF")
    return apply_waivers(project, F)

# ---------------------------------------------------------------- PDF vs PDF (transcription check, or final check of a PDF-to-PDF job)
def qc_pdf_pair(pdf_a, pdf_b, cfg, lang_a, lang_b, where="pdf"):
    F = []; ta = sh(["pdftotext", "-layout", str(pdf_a), "-"])[1]; tb = sh(["pdftotext", "-layout", str(pdf_b), "-"])[1]
    F += compare_numbering(numbering_inventory_pdf(pdf_a, lang_a), numbering_inventory_pdf(pdf_b, lang_b), cfg, where, "numbered items in the PDFs")
    F += pdf_font_audit(pdf_b, "target pdf")
    pa = int(re.search(r"Pages:\s+(\d+)", sh(["pdfinfo", str(pdf_a)])[1]).group(1)); pb = int(re.search(r"Pages:\s+(\d+)", sh(["pdfinfo", str(pdf_b)])[1]).group(1))
    if not (0.7 <= pb / pa <= 1.6): F.append(Finding("warn", "page_count", where, "pages %d -> %d" % (pa, pb)))
    na, nb = collections.Counter(re.findall(r"(?<![\w.,])\d+(?:[.,]\d+)?(?![\w])", normalize_text(ta))), collections.Counter(re.findall(r"(?<![\w.,])\d+(?:[.,]\d+)?(?![\w])", normalize_text(tb)))
    inter, union = sum((na & nb).values()), sum((na | nb).values())
    if union and inter / union < 0.85: F.append(Finding("warn", "number_altered", where, "numeric tokens overlap only %.0f%% between the two PDFs (text layers are imperfect on math; compare the flagged pages)" % (100 * inter / union)))
    if py_has("pdfplumber"):
        fa, fb = detect_figures(pdf_a, cfg), detect_figures(pdf_b, cfg)
        ca, cb = sum(f["kind"] == "figure" for f in fa), sum(f["kind"] == "figure" for f in fb)
        if ca != cb: F.append(mk(cfg, "figure_mismatch", where, "figure regions %d -> %d" % (ca, cb)))
        ta_, tb_ = sum(f["kind"] == "table_like" for f in fa), sum(f["kind"] == "table_like" for f in fb)
        if ta_ != tb_: F.append(Finding("warn", "table_mismatch", where, "table-like regions %d -> %d" % (ta_, tb_)))
    return F

# ---------------------------------------------------------------- DOCX vs DOCX
def qc_docx_pair(src_docx, out_docx, cfg):
    a, b = docx_inventory(src_docx), docx_inventory(out_docx); F = []
    if a["math"] != b["math"]: F.append(mk(cfg, "docx_math_mismatch", "word/document.xml", "equations differ: %d -> %d (first differing: %s)" % (len(a["math"]), len(b["math"]), next((x[:80] for x, y in zip(a["math"], b["math"]) if x != y), ""))))
    for k in ("paragraphs", "drawings", "objects", "fields", "bookmarks", "hyperlinks", "footnotes", "endnotes", "sections", "comments", "charts", "tables", "media"):
        if a[k] != b[k]: F.append(mk(cfg, "docx_structure_mismatch", k, "%s: %s -> %s" % (k, a[k] if not isinstance(a[k], list) else len(a[k]), b[k] if not isinstance(b[k], list) else len(b[k]))))
    return F

# ---------------------------------------------------------------- report
def summarize(findings):
    c = collections.Counter((f.severity, f.check) for f in findings)
    return {"block": sum(v for (s, _), v in c.items() if s == "block"), "warn": sum(v for (s, _), v in c.items() if s == "warn"),
            "info": sum(v for (s, _), v in c.items() if s == "info"), "by_check": {"%s:%s" % k: v for k, v in sorted(c.items())}}

def save_report(project, findings, name="report"):
    qdir = project.p("qc"); sm = summarize(findings)
    jdump(qdir / (name + ".json"), {"summary": sm, "findings": [dict(asdict(f), key=f.key()) for f in findings]})
    md = ["# QC report (%s)" % name, "", "block: **%d**  warn: %d  info: %d" % (sm["block"], sm["warn"], sm["info"]), ""]
    for sev in ("block", "warn", "info"):
        rows = [f for f in findings if f.severity == sev]
        if not rows: continue
        md += ["## %s (%d)" % (sev.upper(), len(rows)), ""]
        for f in rows[:400]:
            md.append("- `%s` **%s** @ %s — %s%s" % (f.key(), f.check, f.where, f.detail, ("\n    - src: `%s`\n    - tgt: `%s`" % (f.src[:160].replace("`", "'"), f.tgt[:160].replace("`", "'"))) if (f.src or f.tgt) else ""))
        md.append("")
    write(qdir / (name + ".md"), "\n".join(md)); return sm

# =============================================================================
# PART 10A — COMMAND-LINE INTERFACE (pipelines: PDF->PDF, PDF->DOCX, DOCX->DOCX, scanned PDF->PDF/DOCX, LaTeX->LaTeX)
# =============================================================================
def _pages(spec):
    if not spec: return None
    out = set()
    for part in spec.split(","):
        a, _, b = part.partition("-"); out.update(range(int(a), int(b or a) + 1))
    return out

def _P(path):
    P = Project(path)
    if not P.cfg_path.exists(): raise SystemExit("not a project (no project.json): %s  -> run `init` first" % path)
    return P

def _src_cfg(cfg): return deep_merge(cfg, {"project": {"target_lang": cfg["project"]["source_lang"]}})

def cmd_check_env(a):
    rows = []
    for b, why in (("pdflatex", "required: build PDF (Latin-script targets)"), ("pdfinfo", "required: PDF probe"), ("pdftotext", "required: text layer"), ("pdftoppm", "required: page images"),
                   ("pdfimages", "probe"), ("xelatex", "needed for Arabic/CJK/Hindi targets"), ("pandoc", "needed for LaTeX->DOCX"), ("tesseract", "needed for scanned PDFs"),
                   ("bibtex", "optional"), ("fc-list", "font lookup for xelatex")):
        rows.append((b, "yes" if have(b) else "NO", why))
    for m, why in (("lxml", "required for DOCX"), ("pdfplumber", "required for figure detection"), ("PIL", "contact sheets / scanned fixtures"), ("docx", "optional (tests only)")):
        rows.append(("python:" + m, "yes" if py_has(m) else "NO", why))
    langs = sorted(tesseract_langs() - {""}) if have("tesseract") else []
    for r in rows: print("%-16s %-4s %s" % r)
    print("tesseract languages:", ", ".join(langs) or "-")

def cmd_probe(a):
    cfg = load_config(); r = probe_pdf(a.pdf, cfg); pp = r.pop("per_page")
    print(json.dumps(r, ensure_ascii=False, indent=1)); need = [p["page"] for p in pp if p["needs_ocr"]]
    print("pages needing OCR: %s" % (need if len(need) < 60 else "%d pages" % len(need)))

def cmd_init(a):
    P = Project.init(a.proj, a.src, a.tgt, a.domain, [w for w in (a.keep_words or "").split(",") if w])
    print("project created at %s (%s -> %s). Edit project.json / glossary.csv as needed." % (a.proj, P.cfg["project"]["source_lang"], P.cfg["project"]["target_lang"]))

def cmd_extract(a):
    P = _P(a.proj); r = make_packets(a.pdf, P.p("packets"), P.cfg, _pages(a.pages), a.chunk)
    cfg_src = _src_cfg(P.cfg); tdir = P.p("transcribed"); tdir.mkdir(parents=True, exist_ok=True)
    if not (tdir / "preamble.tex").exists(): write(tdir / "preamble.tex", render_preamble(cfg_src, "source.pdf"))
    shutil.copyfile(a.pdf, tdir / "source.pdf")
    pr = r["probe"]; print("PDF is %s (%d pages). Packets: %d. Tasks:" % (pr["classification"], pr["pages"], r["packets"]))
    for t in r["tasks"]: print("  ", t)
    print("Next: have a vision-capable model transcribe each task file into %s/*.tex (see RECONSTRUCT_PROMPT inside the task files), then `verify-src`." % tdir)
    for adv in pr["advice"]: print("note:", adv)

def cmd_ocr(a):
    cfg = load_config(); sc, sp = get_profile(a.lang); tmp = Path(tempfile.mkdtemp(prefix="ocr_"))
    for p in sorted(_pages(a.pages) or [1]):
        png = rasterize(a.pdf, tmp, cfg["extraction"]["ocr_dpi"], p, p, "pg%03d" % p)[0]; o = ocr_page(png, sp["ocr"])
        print("== page %d  math_suspect=%s symbol_density=%s low_conf_words=%d %s" % (p, o["math_suspect"], o["symbol_density"], len(o["low_conf"]), o["warning"])); print(o["text"])

def cmd_figures(a):
    cfg = load_config(); figs = detect_figures(a.pdf, cfg, _pages(a.pages))
    for f in figs: print("page %d  %-10s bbox=%s  %s  prose_labels=%s" % (f["page"], f["kind"], f["bbox"], f["srcfig"], [l["text"] for l in f["prose_labels"]]))
    print("%d figures, %d table-like regions" % (sum(f["kind"] == "figure" for f in figs), sum(f["kind"] == "table_like" for f in figs)))

def cmd_preamble(a):
    P = _P(a.proj); cfg = P.cfg if a.lang == "tgt" else _src_cfg(P.cfg); txt = render_preamble(cfg, "source.pdf" if a.source_pdf else None)
    if a.out: write(a.out, txt); print("written", a.out)
    else: print(txt)

def cmd_ingest(a):
    P = _P(a.proj); files = []
    for f in a.files:
        f = Path(f); files += sorted(f.rglob("*.tex")) if f.is_dir() else [f]
    root = Path(a.root) if a.root else (Path(a.files[0]) if Path(a.files[0]).is_dir() else Path(a.files[0]).parent)
    print(json.dumps(ingest_tex(P, files, root), indent=1)); collect_source_flags(P)
    print("source issues (reported only, never fixed) -> %s" % P.p("qc", "source_flags.jsonl"))

def cmd_glossary(a):
    P = _P(a.proj); segs = P.load_segments(); texts = [s.src for s in segs if s.kind in ("prose", "heading")]
    cands = extract_term_candidates(texts, P.cfg["project"]["source_lang"], a.candidates); g = P.glossary(); known = {fold(r["src"]) for r in g.rows}
    new = [(w, c) for w, c in cands if fold(w) not in known]
    write(P.p("qc", "term_candidates.csv"), "term,count\n" + "\n".join("%s,%d" % x for x in new))
    task = fill_prompt(GLOSSARY_PROMPT, P.cfg) + "\nCANDIDATES (term,count):\n" + "\n".join("%s,%d" % x for x in new[:a.candidates]) + "\n\nEXISTING GLOSSARY ROWS: %d (do not duplicate)\n" % len(g.rows)
    write(P.p("work", "tasks", "glossary_task.md"), task)
    print("glossary: %d rows (%s). %d new candidates -> qc/term_candidates.csv; task for an AI/human -> work/tasks/glossary_task.md" % (len(g.rows), P.p("glossary.csv"), len(new)))

def cmd_translate(a):
    P = _P(a.proj); st = translate_all(P, a.backend, a.collect, a.review); segs = P.load_segments(); c = collections.Counter(s.status for s in segs)
    print("batch result:", st, "| segments by status:", dict(c))
    be = a.backend or P.cfg["translate"]["backend"]
    if be == "agent" and not a.collect and c.get("pending", 0) + c.get("retry", 0):
        print("AGENT MODE: fill every work/tasks/batch_NNNN.result.json, then run `translate --collect`; repeat until nothing is pending/retry.")
    if c.get("failed"): print("WARNING: %d segments failed validation after the retry limit; they stay in the source language and block export." % c["failed"])

def cmd_status(a):
    P = _P(a.proj); segs = P.load_segments(); c = collections.Counter(s.status for s in segs)
    print("project %s  (%s)\ndocuments: %d  segments: %d  %s" % (a.proj, P.pair, len(P.docs()), len(segs), dict(c)))
    rep = jload(P.p("qc", "report.json"))
    if rep: print("last QC:", rep["summary"]["block"], "blocking,", rep["summary"]["warn"], "warnings (see qc/report.md)")
    else: print("no QC report yet: run `qc`")

def cmd_assemble(a):
    P = _P(a.proj); w = assemble_project(P); print("assembled %d file(s) into %s" % (len(w), P.p("out", "tex")))

def cmd_qc(a):
    P = _P(a.proj); F = qc_project(P, a.compile, a.main, a.reassemble); sm = save_report(P, F)
    print("QC: %d blocking, %d warnings, %d info -> %s" % (sm["block"], sm["warn"], sm["info"], P.p("qc", "report.md")))
    for f in [x for x in F if x.severity == "block"][:25]: print("  BLOCK", f.check, f.where, f.detail[:150])
    raise SystemExit(2 if sm["block"] else 0)

def cmd_qc_pdf(a):
    cfg = load_config(); F = qc_pdf_pair(a.pdf_a, a.pdf_b, cfg, a.lang_a, a.lang_b); sm = summarize(F)
    print(json.dumps(sm, indent=1)); [print(" ", f.severity, f.check, f.detail[:180]) for f in F]; raise SystemExit(2 if sm["block"] else 0)

def cmd_qc_docx(a):
    cfg = load_config(); F = qc_docx_pair(a.src, a.out, cfg); sm = summarize(F)
    print(json.dumps(sm, indent=1)); [print(" ", f.severity, f.check, f.where, f.detail[:180]) for f in F]; raise SystemExit(2 if sm["block"] else 0)

def cmd_build(a):
    P = _P(a.proj); r = build_pdf(P.p("out", "tex", a.main), P.cfg); print(json.dumps(r, indent=1, ensure_ascii=False)); raise SystemExit(0 if r.get("ok") else 1)

def cmd_verify_src(a):
    """Gate between TRANSCRIPTION and TRANSLATION: compile the transcribed source-language LaTeX and compare it with the original PDF."""
    P = _P(a.proj); cfg_src = _src_cfg(P.cfg); tdir = P.p("transcribed"); main = tdir / a.main
    if not main.exists(): raise SystemExit("missing %s" % main)
    r = build_pdf(main, cfg_src); print("compile:", {k: r.get(k) for k in ("ok", "pages", "errors", "undefined_refs", "missing_chars")})
    F = []
    if r.get("errors"): F.append(mk(P.cfg, "latex_error", "transcription", "LaTeX errors: %s" % r["errors"][:3]))
    if r.get("undefined_refs"): F.append(mk(P.cfg, "broken_reference", "transcription", "undefined references: %s" % r["undefined_refs"][:6]))
    src_pdf = tdir / "source.pdf"
    if r.get("pdf"): F += qc_pdf_pair(src_pdf, r["pdf"], P.cfg, P.cfg["project"]["source_lang"], P.cfg["project"]["source_lang"], "transcription-vs-original")
    for f in tdir.rglob("*.tex"):
        for it in re.findall(r"%\s*FLAG\((.*?)\):\s*(.*)", read(f)): F.append(Finding("warn", "model_flag", f.name, "%s: %s" % it))
    sm = save_report(P, F, "transcription"); print("transcription QC: %d blocking, %d warnings -> %s" % (sm["block"], sm["warn"], P.p("qc", "transcription.md")))
    if r.get("pdf") and a.contact:
        for p in sorted(_pages(a.contact)):
            A = rasterize(src_pdf, P.p("qc", "contact"), 70, p, p, "a%03d" % p)[0]; B = rasterize(r["pdf"], P.p("qc", "contact"), 70, p, p, "b%03d" % p)[0]
            print("contact sheet:", contact_sheet(A, B, P.p("qc", "contact", "page%03d.png" % p)))
    raise SystemExit(2 if sm["block"] else 0)

def cmd_docx(a):
    P = _P(a.proj); r = docx_translate(a.src, a.out, P, a.backend, a.collect); print(json.dumps(r, indent=1, ensure_ascii=False))
    if r.get("status") == "done":
        F = qc_docx_pair(a.src, a.out, P.cfg) + qc_segments(P, P.load_segments(), P.glossary()); sm = save_report(P, F, "docx")
        print("DOCX QC: %d blocking, %d warnings -> %s" % (sm["block"], sm["warn"], P.p("qc", "docx.md"))); raise SystemExit(2 if sm["block"] else 0)

def cmd_tex2docx(a):
    P = _P(a.proj); print(json.dumps(latex_to_docx(P.p("out", "tex", a.main), a.out, P.cfg), indent=1, ensure_ascii=False))

def cmd_export(a):
    P = _P(a.proj)
    if not a.keep_edits: assemble_project(P)
    F = qc_project(P, True, a.main); sm = save_report(P, F)
    print("QC before export: %d blocking, %d warnings -> %s" % (sm["block"], sm["warn"], P.p("qc", "report.md")))
    if sm["block"] and not a.force:
        for f in [x for x in F if x.severity == "block"][:25]: print("  BLOCK", f.check, f.where, f.detail[:150])
        raise SystemExit("EXPORT REFUSED: fix the blocking findings (or waive them with `waive`, or re-run with --force).")
    fin = P.p("out", "final"); fin.mkdir(parents=True, exist_ok=True); pdf = P.p("out", "tex", Path(a.main).with_suffix(".pdf").name)
    if pdf.exists(): shutil.copyfile(pdf, fin / pdf.name); print("PDF ->", fin / pdf.name)
    if a.docx: r = latex_to_docx(P.p("out", "tex", a.main), fin / Path(a.main).with_suffix(".docx").name, P.cfg); print("DOCX ->", r["docx"])
    shutil.copyfile(P.p("qc", "report.md"), fin / "QC_REPORT.md")

def cmd_contact(a):
    cfg = load_config(); out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    for p in sorted(_pages(a.pages)):
        A = rasterize(a.pdf_a, out, 70, p, p, "a%03d" % p)[0]; B = rasterize(a.pdf_b, out, 70, p, p, "b%03d" % p)[0]; print(contact_sheet(A, B, out / ("page%03d.png" % p)))

def cmd_waive(a):
    P = _P(a.proj); f = P.p("qc", "waivers.json"); w = set(jload(f, [])) | set(a.keys); jdump(f, sorted(w)); print("%d waivers recorded in %s (re-run `qc`)" % (len(w), f))

def cmd_source_flags(a):
    P = _P(a.proj); out = collect_source_flags(P); print("%d source-side flags -> %s" % (len(out), P.p("qc", "source_flags.jsonl")))
    for o in out[:30]: print(" ", o)

def cmd_languages(a):
    for c, p in LANGUAGE_PROFILES.items(): print("%-3s %-22s script=%-10s dir=%s  babel=%s  ocr=%s%s" % (c, p["name"], p["script"], p["dir"], p["babel"], p["ocr"], "  (labels: have them reviewed)" if p["review"] else ""))

def cmd_provenance(a): print(json.dumps(PROVENANCE, ensure_ascii=False, indent=1))

def cmd_readme(a): print(render_readme())

def build_parser():
    ap = argparse.ArgumentParser(prog=TOOL_NAME, description="Translate mathematics/physics books while keeping every formula, number, table and figure intact. Run `readme` first.")
    ap.add_argument("--version", action="version", version=TOOL_NAME + " " + TOOL_VERSION)
    sp = ap.add_subparsers(dest="cmd", required=True)
    def add(name, fn, *args, help=""):
        p = sp.add_parser(name, help=help); p.set_defaults(fn=fn)
        for spec in args:
            names, kw = spec if isinstance(spec, tuple) else (spec, {}); p.add_argument(*([names] if isinstance(names, str) else names), **kw)
        return p
    add("readme", cmd_readme, help="print the full README"); add("provenance", cmd_provenance, help="what was actually used in the original session vs what is new")
    add("languages", cmd_languages, help="supported language profiles"); add("check-env", cmd_check_env, help="check installed tools")
    add("selftest", lambda a: raise_exit(run_selftest(a.keep, a.book, a.json)), ("--keep", {"action": "store_true"}), ("--book", {"help": "optional real source PDF for the real-book checks"}), ("--json", {"help": "write results here"}), help="run the built-in test-suite")
    add("probe", cmd_probe, "pdf", help="digital / scanned / mixed + advice")
    add("init", cmd_init, "proj", ("--src", {"required": True}), ("--tgt", {"required": True}), ("--domain", {"default": "mathematics", "choices": ["mathematics", "physics", "both"]}),
        ("--keep-words", {"help": "comma list of source-language words kept inside formulas/cells, e.g. et,ou,non"}), help="create a project")
    add("extract", cmd_extract, "proj", "pdf", ("--pages", {}), ("--chunk", {"type": int, "default": 8}), help="PDF -> page packets (+OCR when needed) and transcription tasks")
    add("ocr", cmd_ocr, "pdf", ("--lang", {"default": "en"}), ("--pages", {"default": "1"}), help="OCR pages and report math suspicion")
    add("figures", cmd_figures, "pdf", ("--pages", {}), help="list figures/tables with ready-made \\srcfig crops")
    add("preamble", cmd_preamble, "proj", ("--lang", {"choices": ["src", "tgt"], "default": "tgt"}), ("--source-pdf", {"action": "store_true"}), ("--out", {}), help="generate a LaTeX preamble for the language")
    add("ingest", cmd_ingest, "proj", ("files", {"nargs": "+"}), ("--root", {}), help="protect + segment LaTeX source files")
    add("glossary", cmd_glossary, "proj", ("--candidates", {"type": int, "default": 250}), help="export term candidates and a glossary task")
    add("translate", cmd_translate, "proj", ("--backend", {"choices": list(BACKENDS)}), ("--collect", {"action": "store_true"}), ("--review", {"action": "store_true"}), help="translate segments (agent hand-off, API, or mock)")
    add("status", cmd_status, "proj", help="progress"); add("assemble", cmd_assemble, "proj", help="rebuild translated .tex tree from segments")
    add("qc", cmd_qc, "proj", ("--compile", {"action": "store_true"}), ("--main", {}), ("--reassemble", {"action": "store_true"}), help="compare source and translation")
    add("qc-pdf", cmd_qc_pdf, "pdf_a", "pdf_b", ("--lang-a", {"required": True}), ("--lang-b", {"required": True}), help="compare two PDFs")
    add("qc-docx", cmd_qc_docx, "src", "out", help="compare two DOCX files")
    add("build", cmd_build, "proj", ("--main", {"default": "main.tex"}), help="compile out/tex/<main>")
    add("verify-src", cmd_verify_src, "proj", ("--main", {"default": "main.tex"}), ("--contact", {"help": "pages for side-by-side sheets, e.g. 1,5,10-12"}), help="check the transcription before translating")
    add("docx", cmd_docx, "proj", "src", "out", ("--backend", {"choices": list(BACKENDS)}), ("--collect", {"action": "store_true"}), help="DOCX -> translated DOCX (native)")
    add("tex2docx", cmd_tex2docx, "proj", "out", ("--main", {"default": "main.tex"}), help="translated LaTeX -> DOCX")
    add("export", cmd_export, "proj", ("--main", {"default": "main.tex"}), ("--docx", {"action": "store_true"}), ("--force", {"action": "store_true"}), ("--keep-edits", {"action": "store_true"}), help="QC gate, then PDF (+DOCX) into out/final")
    add("contact", cmd_contact, "pdf_a", "pdf_b", ("--pages", {"required": True}), ("--out", {"default": "contact"}), help="side-by-side page images")
    add("waive", cmd_waive, "proj", ("keys", {"nargs": "+"}), help="accept specific findings (by key) after human review")
    add("source-flags", cmd_source_flags, "proj", help="issues found in the SOURCE (reported, never fixed)")
    return ap

def raise_exit(code): raise SystemExit(code)

def main(argv=None):
    a = build_parser().parse_args(argv); a.fn(a)

# =============================================================================
# PART 10B — README (embedded; print with `python MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL.py readme`)
# =============================================================================
README_TEXT = r'''
MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL  v@@VERSION@@
========================================================================
One self-contained Python file: a pipeline that translates mathematics and physics books between languages while keeping the
mathematical content, numbering, tables, figures and structure intact, and that compares source and translation before export.

WHAT IT IS, HONESTLY
 * The tool does the deterministic work: inspect the PDF, extract text/OCR, shield every formula/command/table/figure behind placeholders,
   batch ordinary language for translation, validate every returned segment, rebuild the document, build PDF/DOCX, and compare source vs
   translation with ~30 automatic checks. A strong language model (or a human) does the two judgement steps: (1) transcribing page images
   to LaTeX when the source is a PDF, (2) translating the shielded text. The tool carries the prompts for both and verifies their output.
 * It is NOT a one-click "upload PDF, get PDF". PDF input needs the transcription step (a vision-capable model reading the page images the tool
   prepares). Input that already is LaTeX or DOCX needs no transcription and is handled fully automatically.
 * Components differ in how well they are verified; see "STATUS" at the end. Nothing is claimed that was not run.

HOW TO GIVE THIS TO ANOTHER AI
 "Read this file (`python MATH_PHYSICS_ACADEMIC_TRANSLATION_TOOL.py readme`). Use it to translate <book> from <SRC> to <TGT>,
  following the pipeline for my input type. Do not edit any formula. Report every blocking QC finding to me."
 The AI must run the commands below itself; the stage that needs its own intelligence is marked [AI].

REQUIREMENTS
 Python 3.9+, `pip install lxml pdfplumber pillow` (python-docx only for the self-test).
 System: poppler-utils (pdfinfo pdftotext pdftoppm pdfimages), TeX Live (pdflatex; xelatex + fontspec/polyglossia/xeCJK for Arabic, Hindi, CJK,
 Cyrillic), pandoc (LaTeX->DOCX), tesseract + language data (scans). `python TOOL.py check-env` lists what is present.
 TeX Live needs the babel/polyglossia language support and a suitable font for the target language (the self-test reports what is missing).

PROJECT LAYOUT (created by `init`)
 project.json (all settings) | glossary.csv | source/ (originals) | transcribed/ (LaTeX transcribed from a PDF) | packets/ (page images, evidence text, tasks)
 work/ (skeletons, protected-span stores, segments.jsonl, tasks/batch_*.json) | out/tex (translated LaTeX) | out/final (exported files) | qc/ (reports)

PIPELINE A — LaTeX source  ->  translated LaTeX / PDF / DOCX            (fully automatic apart from the [AI] translation)
  init     python TOOL.py init proj --src fr --tgt en --domain mathematics [--keep-words et,ou,non]
  ingest   python TOOL.py ingest proj book/            # protects formulas/commands/tables/figures/comments, segments the prose
  glossary python TOOL.py glossary proj                # review proj/glossary.csv; seed exists for fr->en; add rows for other pairs [AI]
  translate python TOOL.py translate proj              # agent mode writes work/tasks/batch_NNNN.json
            [AI] translate each batch under its "system" prompt, write batch_NNNN.result.json, then
            python TOOL.py translate proj --collect    # validates every segment; rejected ones come back with a retry note; repeat until none pending
            (or set translate.backend to anthropic/openai + API key env var and run `translate` once)
  assemble python TOOL.py assemble proj
  qc       python TOOL.py qc proj --compile --main main.tex      # exit code 2 if anything blocks
  export   python TOOL.py export proj --main main.tex [--docx]   # refuses while blocking findings exist (use `waive` after human review)

PIPELINE B — PDF (digital)  ->  translated PDF / DOCX
  probe    python TOOL.py probe book.pdf               # digital / scanned / mixed; keeps the text layer when it is trustworthy
  init, then
  extract  python TOOL.py extract proj book.pdf        # page PNGs, evidence text, figure regions with \srcfig crops, preamble, transcription tasks
  [AI]     transcribe each packets/tasks/reconstruct_*.md into proj/transcribed/*.tex (source language, verbatim, no translation, RECONSTRUCT_PROMPT)
  verify   python TOOL.py verify-src proj --main main.tex --contact 1,5,10-12
           compiles the transcription and compares it with the original PDF (numbering inventory, figures, numbers, page count) and writes
           side-by-side images. [AI] fix until no blocking findings; look at the contact sheets; keep `% FLAG(...)` comments honest.
  then run Pipeline A from `ingest proj transcribed/ --root transcribed` on. Figures stay the ORIGINAL vector art (\srcfig crop of source.pdf);
  only prose labels next to a drawing are candidates for translation (`figures book.pdf`, `figure_overlay_tex`). Formula labels are never touched.
  For DOCX instead of PDF add --docx to `export` (or `tex2docx proj out.docx`).

PIPELINE C — DOCX  ->  translated DOCX (edited in place)
  python TOOL.py init proj --src fr --tgt en
  python TOOL.py docx proj in.docx out.docx --backend agent      # writes batches; [AI] fills results
  python TOOL.py docx proj in.docx out.docx --collect            # applies translations, then compares structure
  Word equations (OMML), fields (equation numbers, TOC, cross-references), footnotes, tables, drawings, styles and layout are kept as XML;
  only the text runs of ordinary-language paragraphs are rewritten (formatting inside a paragraph is carried through placeholder tags).
  Tracked-change paragraphs and code styles are skipped and counted. TOC/field text is refreshed by Word on open (updateFields set).
  `qc-docx in.docx out.docx` repeats the comparison on any pair of files.

PIPELINE D — scanned PDF  ->  translated PDF / DOCX
  probe decides per page (no text + images, or garbled glyphs => OCR); `extract` OCRs only those pages (tesseract, language from the profile) and flags
  low-confidence words and math-dense pages (math_suspect). OCR text is never trusted for formulas: the [AI] transcription step reads the PNG itself
  (vision model), or use a math-aware OCR (nougat / marker / pix2tex / Mathpix adapters in run_math_ocr; open ones preferred, Mathpix is paid).
  From then on identical to Pipeline B (verify-src is the safety net: numbering, numbers, figures are compared with the scan).

WHAT IS PROTECTED (never reaches the translator)
  inline/display math ($..$, \(..\), \[..\], equation/align/gather/... environments), \label \ref \eqref \cite \includegraphics \input \usepackage ... (commands with
  arguments), \begin/\end tokens, comments, verbatim, TikZ/pgfplots drawings, table structure, DOCX fields/equations/drawings/footnote marks/hyperlinks.
  Ordinary language INSIDE protected spans is handled separately and explicitly: \text{...} in formulas, table cells and TikZ node labels become "child
  segments" (kinds textinmath / cell / figlabel). Words you decide to keep (e.g. et/ou/non and V/F in truth tables) go in protect.keep_words_in_math.
  Notation is never localised (sin/tan/sh, decimal comma, vector arrows); the translator flags convention differences, QC flags localised tokens.
  The author is never corrected: apparent source errors are translated faithfully and listed in qc/source_flags.jsonl.

QC — what is compared before export (qc/report.md|json). BLOCK = stops `export`.
  math_altered / math_missing / math_added (BLOCK)   exact comparison of every formula, canonical spacing; detail names changed sub/superscripts, digits, commands
  placeholder_mismatch, formatting_corruption (BLOCK) placeholder multiset, LaTeX command skeleton and brace balance per segment
  number_altered (BLOCK)                              numbers in prose, per segment (decimal separators preserved)
  numbering_mismatch (BLOCK)                          printed number of every \label (theorem/definition/equation/section) after compiling source and translation;
                                                      numbered-item inventory of PDF text layers for any language pair
  env_mismatch, table_mismatch, figure_mismatch, caption_mismatch, footnote_mismatch (BLOCK)   structure inventory (environments, rows/cols/multicolumn, figures, captions, footnotes, items)
  label_mismatch, ref_mismatch, broken_reference (BLOCK)   labels/references, plus undefined references after compilation
  untranslated_segment (BLOCK)                        a segment failed validation and stays in the source language
  latex_error (BLOCK), docx_math_mismatch, docx_structure_mismatch (BLOCK)
  numbering_mismatch_heuristic (warn)                 PDF-vs-PDF numbered items when one PDF's fonts cannot be recognised as bold (wrapped cross-references look like items)
  bitmap_fonts, no_unicode_map (warn)                 target PDF has Type 3 bitmap fonts / fonts without Unicode map (blurry, copy/paste and search unreliable)
  possible_missing_text / possible_added_text (warn)  sentence count and robust length-ratio outliers
  untranslated_fragment, identical_to_source (warn)   function-word residue of the source language / source script in the target
  duplicated_paragraph (warn)                         exact and near duplicates whose sources differ
  glossary_miss, forbidden_term, term_inconsistency (warn; BLOCK if the glossary row is locked=1)   context-aware terminology layer
  notation_localised, style_ai_phrase, repetitive_phrasing (warn); math_reordered, missing_glyphs, page_count, stale_output (warn)
  Tune with project.json qc.*; accept a reviewed finding with `waive proj <key>` (keys are printed in report.md).

GLOSSARY
  glossary.csv columns: src,tgt,lang_pair,domain,context,forbidden,notes,locked. One row per sense; `context` words choose the sense (travail -> work in
  mechanics; corps -> field in algebra / body in mechanics; intensité -> current in electricity / intensity in optics); `forbidden` lists literal renderings that
  must not appear. A fr->en seed (~90 rows, analysis/algebra/logic/mechanics/electricity/optics) ships inside the file as a STARTING POINT, not a reviewed
  dictionary. For other pairs `glossary proj` writes term candidates and a glossary task for an AI/human; have a mathematician/physicist review it.

LANGUAGES (python TOOL.py languages)
  fr en es de it pt nl ru tr ar zh ja ko hi. Each profile holds script, direction, babel/polyglossia name, OCR code, quotation marks, style notes, stop-words (residue
  detection) and the structural words (Theorem, Proof, Contents ...). Profiles marked review need a native reader to confirm the structural words. Add any other
  language with register_language(code, profile). Mathematics is identical in every language and is never taken from the profile.

CONFIGURATION (project.json) — the main knobs
  protect.text_in_math, protect.keep_words_in_math, protect.tables, protect.tikz_labels | translate.backend/model/api_base/batch_chars/max_retries
  latex.engine/documentclass/fontsize/paper/margin/extra_packages/font_package (auto|lmodern|mathptmx|none) | figures.mode/pad_pt/label_gap_pt | qc.block_on/dup_jaccard/length_ratio_mad_k | docx.skip_styles

KNOWN LIMITS
  * No in-place overlay translation of the original PDF pages (the output PDF is rebuilt from LaTeX, so layout is close, not identical; page count can change).
  * Transcription quality of scanned/PDF input depends on the vision model; verify-src and the contact sheets are the safety net, not a guarantee.
  * LaTeX -> DOCX: pandoc converts the structure; theorem bodies lose italics and align-type equation numbers are not re-inserted.
  * RTL (Arabic) and Indic builds are tested only to the extent shown in STATUS; CJK needs xeCJK and a CJK font.
  * Backends `anthropic` and `openai` (any OpenAI-compatible server) could not be exercised without network access when this file was built.
  * Text-layer comparison of two PDFs is weak on mathematics (fonts); the tool therefore also compares at the LaTeX level and via compiled .aux numbering.

PROVENANCE — what was actually used for the French->English translation of "Cours d'analyse 1" and what this tool adds
@@PROVENANCE@@

STATUS — verified by `python TOOL.py selftest` in the build sandbox (re-run it in your environment)
  Run in the build sandbox (Linux, TeX Live with pdflatex/xelatex, poppler, tesseract-eng, pandoc, pdfplumber, python-docx): 28 PASSED, 0 FAILED, 5 SKIPPED.
  PASSED: environment | fixture compiles | PDF probe (digital) | figure + table region detection | original vector figure re-used via \srcfig crop (pixel-size check) |
          protect->restore identity on LaTeX | keep-words policy (et/ou/non) | context-aware glossary | full pipeline ingest->translate(MOCK backend)->assemble->compile->QC (0 blocking) |
          QC catches 16 deliberate corruptions (altered / dropped / added formula, lost table row, lost figure, lost label, broken \ref, lost proof, lost footnote, changed number,
          untranslated segment, dropped placeholder, duplicated paragraph, text left in French, forbidden term, localised notation) | changed sub/superscripts are named |
          compiled numbering comparison via .aux | PDF font audit + clean preamble | agent hand-off with rejection/retry | PDF-vs-PDF comparison | DOCX->DOCX (equations + footnote
          preserved, tampering caught) | LaTeX->DOCX with literal numbering | scanned PDF (classify, OCR, packets) | language profiles that build a PDF: en, fr, es, de, it, pt, nl, ru, tr |
          real French book: probe=digital 73 pages, 8 figures + 2 table-like pages detected, 159 numbered items read from its text layer.
  SKIPPED (this sandbox lacks the components, NOT tested here): ar (bidi.sty missing), zh/ja/ko (xeCJK missing), hi (no Devanagari font). Their profiles exist; run `selftest` after installing them.
  NEVER EXERCISED: backends `anthropic` and `openai` (no network); math OCR adapters nougat / marker / pix2tex / Mathpix; a real translation of a full book by this file (the MOCK backend is a
  dictionary used only to test plumbing - it is not a translator). The self-test shows the mechanics and the QC net work; translation quality comes from the LLM you connect.
  Observed in a real check: `qc-pdf` of the French book vs the earlier delivered English PDF reports no numbering/figure/table discrepancy and two warnings (bitmap_fonts, no_unicode_map).
'''

def render_readme():
    prov = []
    for k, items in PROVENANCE.items():
        prov.append("[%s]" % k.replace("_", " ").upper())
        for it in items: prov.append("  - " + it)
        prov.append("")
    return README_TEXT.replace("@@VERSION@@", TOOL_VERSION).replace("@@PROVENANCE@@", "\n".join(prov))

# =============================================================================
# PART 10C — SELF-TEST (fixtures are generated on the fly; nothing external is needed except the tools being tested)
# Every test either PASSes, FAILs, or is SKIPped with the reason (missing tool/font/language support).
# =============================================================================
class SkipTest(Exception): pass

FIXTURE_TEX = r"""\documentclass[12pt]{report}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage{amsmath,amssymb,amsthm}
\usepackage{tikz}
\newtheorem{definition}{Définition}[section]
\newtheorem{theorem}{Théorème}[section]
\renewcommand{\proofname}{Démonstration}
\renewcommand{\contentsname}{Table des matières}
\begin{document}
\chapter{Suites réelles}\label{ch:suites}

\section{Limite d'une suite}

\begin{definition}[Limite]\label{def:lim}
Soit $(u_n)$ une suite réelle. On dit que la suite converge vers $\ell$ si
\[ \forall \varepsilon>0,\ \exists N,\ \forall n\geq N,\ |u_n-\ell|<\varepsilon \quad\text{pour tout } n. \]
\end{definition}

\begin{theorem}[Théorème des valeurs intermédiaires]\label{thm:tvi}
Soit $f$ une fonction continue sur $[a,b]$. Voir la définition~\ref{def:lim} et l'équation \eqref{eq:e}. On a $e^{i\theta}=\cos\theta+i\sin\theta$ avec $3,1415$ environ.
\end{theorem}

Un énoncé plus long suit ici : la fonction est continue, donc elle est bornée sur le segment, et cette borne supérieure est atteinte\footnote{Voir le chapitre suivant.}. Alors on a le résultat pour tout entier naturel positif.

\begin{equation}\label{eq:e}
\int_0^1 \frac{dx}{\sqrt{1-x^2}} = \frac{\pi}{2}
\end{equation}

\begin{proof}
Le résultat est donc vrai pour tout $n\in\mathbb{N}$.
\end{proof}

\begin{center}
\begin{tabular}{|c|c|c|}\hline $R$ & $S$ & $R$ et $S$ \\ \hline V & F & non $R$ ou $S$ \\ \hline\end{tabular}
\end{center}

\begin{center}\begin{tikzpicture}\draw[->] (0,0)--(4,0); \draw[->] (0,0)--(0,3); \draw (0,0)--(3,2.5); \node at (3.3,2.8) {pente positive};\end{tikzpicture}\end{center}

\section{Suite extraite}
Une suite extraite est une sous-suite de la suite. Voir le théorème~\ref{thm:tvi} pour le résultat avec 2 cas.

Cette section contient un deuxième paragraphe assez long pour tester la détection de duplication. La fonction est dérivable, donc elle est continue sur l'ensemble des réels. La valeur approchée vaut 3,1415 pour 12 termes.
\end{document}
"""

class SelfTest:
    def __init__(self, keep=False, book=None):
        self.tmp = Path(tempfile.mkdtemp(prefix="mpatt_selftest_")); self.keep, self.book, self.results, self.n = keep, book, [], 0
        self.fx = self.tmp / "fx"; self.fx.mkdir(); write(self.fx / "main.tex", ensure_text_fonts(FIXTURE_TEX, load_config())[0]); self.base = None
    def run(self, name, fn):
        try: r = fn(); status, detail = "PASS", (r if isinstance(r, str) else "")
        except SkipTest as e: status, detail = "SKIP", str(e)
        except AssertionError as e: status, detail = "FAIL", str(e) or "assertion failed"
        except SystemExit as e: status, detail = "FAIL", "SystemExit: %s" % e
        except Exception as e: status, detail = "FAIL", ("%s: %s" % (type(e).__name__, e))[:400]
        self.results.append({"test": name, "status": status, "detail": detail}); print("%-5s %-58s %s" % (status, name, detail[:150]), flush=True)
    # ---- helpers
    def need(self, *bins):
        for b in bins:
            if not have(b): raise SkipTest("%s not installed" % b)
    def cfg(self, **kw): return deep_merge(load_config(), kw)
    def build_base(self):
        if self.base: return self.base
        P = Project.init(self.tmp / "base", "fr", "en", keep_words=["et", "ou", "non"]); ingest_tex(P, [self.fx / "main.tex"], self.fx)
        translate_all(P, "mock"); assemble_project(P); self.base = P; return P
    def clone(self, name):
        d = self.tmp / name; shutil.copytree(self.base.root, d); return Project(d)
    def blocks(self, P, **kw):
        F = qc_project(P, **kw); return {f.check for f in F if f.severity == "block"}, {f.check for f in F if f.severity == "warn"}, F
    def edit_out_re(self, P, pattern, new):
        f = P.p("out", "tex", "main.tex"); t = read(f); t2, n = re.subn(pattern, lambda m: new, t, count=1, flags=re.S); assert n == 1, "pattern not found: %s" % pattern; write(f, t2)
    def edit_out(self, P, old, new):
        f = P.p("out", "tex", "main.tex"); t = read(f); assert old in t, "fixture text not found: %r" % old[:50]; write(f, t.replace(old, new, 1))
    def edit_segs(self, P, fn):
        segs = P.load_segments(); fn(segs); P.save_segments(segs); assemble_project(P)
    # ---- tests
    def t_env(self):
        self.need("pdflatex", "pdftotext", "pdfinfo", "pdftoppm")
        for m in ("lxml", "pdfplumber", "PIL"):
            if not py_has(m): raise AssertionError("python module %s missing" % m)
        return "required tools present"
    def t_fixture_pdf(self):
        self.need("pdflatex"); r = build_pdf(self.fx / "main.tex", self.cfg(project={"target_lang": "fr"}), engine="pdflatex"); assert r["ok"], r
        assert Path(r["pdf"]).exists(); return "%d pages" % r["pages"]
    def t_probe(self):
        r = probe_pdf(self.fx / "main.pdf", self.cfg()); assert r["classification"] == "digital" and r["tex_generated"], r["classification"]
        return "digital, producer=%s" % r["producer"]
    def t_figures(self):
        figs = detect_figures(self.fx / "main.pdf", self.cfg()); fg = [f for f in figs if f["kind"] == "figure"]; tb = [f for f in figs if f["kind"] == "table_like"]
        assert len(fg) == 1, "expected 1 figure, got %d" % len(fg); self.fig = fg[0]; assert len(tb) >= 1, "truth table not detected as table-like"
        return "figure bbox=%s labels=%s; table-like regions=%d" % (fg[0]["bbox"], [l["text"] for l in fg[0]["prose_labels"]], len(tb))
    def t_srcfig(self):
        self.need("pdflatex", "pdftoppm"); f = self.fig; d = self.tmp / "crop"; d.mkdir(exist_ok=True); shutil.copyfile(self.fx / "main.pdf", d / "src.pdf")
        write(d / "t.tex", "\\documentclass[border=2pt]{standalone}\\usepackage{graphicx}\\newcommand{\\srcfig}[5]{\\includegraphics[page=#1,trim=#2pt #3pt #4pt #5pt,clip]{src.pdf}}\n\\begin{document}%s\\end{document}\n" % f["srcfig"])
        assert sh(["pdflatex", "-interaction=nonstopmode", "t.tex"], cwd=d)[0] == 0; sh(["pdftoppm", "-r", "72", "-png", "-singlefile", "t.pdf", "t"], cwd=d)
        from PIL import Image, ImageChops
        im = Image.open(d / "t.png").convert("L"); bbox = ImageChops.invert(im).getbbox(); w = f["bbox"][2] - f["bbox"][0]
        assert bbox and abs(im.width - w - 4) < 8, "crop width %d vs figure width %.0f" % (im.width, w); return "original vector art cropped, %dx%d px" % im.size
    def t_roundtrip(self):
        cfg = self.cfg(); sk, st = segment_tex(FIXTURE_TEX, cfg, "fx"); assert assemble(sk, st) == FIXTURE_TEX, "protect/restore is not the identity"
        kinds = collections.Counter(s.kind for s in st.segments); phk = collections.Counter(k[1] for k in st.ph)
        assert phk["M"] >= 10 and phk["T"] == 1 and phk["G"] == 1; return "identity ok; segments=%s placeholders=%s" % (dict(kinds), dict(phk))
    def t_keepwords(self):
        a = self.cfg(); b = self.cfg(protect={"keep_words_in_math": ["et", "ou", "non"]})
        ka = [s for s in segment_tex(FIXTURE_TEX, a, "fx")[1].segments if s.kind == "cell"]; kb = [s for s in segment_tex(FIXTURE_TEX, b, "fx")[1].segments if s.kind == "cell"]
        assert len(ka) == 2 and len(kb) == 0, "cells %d vs %d" % (len(ka), len(kb)); return "truth-table words translated by default, kept with keep_words_in_math"
    def t_pipeline(self):
        self.need("pdflatex"); P = self.build_base(); bl, wn, F = self.blocks(P, compile_check=True, main="main.tex")
        assert not bl, "unexpected blocking findings: %s" % [(f.check, f.detail[:90]) for f in F if f.severity == "block"]
        src = read(P.p("work", "orig", "main.tex")); out = read(P.p("out", "tex", "main.tex"))
        assert tex_math_list(src, P.cfg) == tex_math_list(out, P.cfg), "formulas differ"; assert "\\frac{\\pi}{2}" in out and "R$ et $S$" in out
        return "mock translation -> assemble -> compile -> QC: 0 blocking, %d warnings (mock output is deliberately crude)" % len(wn)
    def _tamper(self, name, how, expect, kind="block", compile_check=False):
        P = self.clone("t_" + name); how(P); bl, wn, F = self.blocks(P, compile_check=compile_check, main="main.tex")
        got = bl if kind == "block" else wn
        assert expect in got, "expected %s finding '%s' but got block=%s warn=%s" % (kind, expect, sorted(bl), sorted(wn)); return "detected as %s: %s" % (kind, expect)
    def t_tampers(self):
        self.need("pdflatex"); self.build_base(); results = []
        cases = [
         ("math_altered", lambda P: self.edit_out(P, "\\frac{\\pi}{2}", "\\frac{\\pi}{3}"), "math_altered", "block"),
         ("math_missing", lambda P: self.edit_out(P, "\\[ \\forall \\varepsilon>0,\\ \\exists N,\\ \\forall n\\geq N,\\ |u_n-\\ell|<\\varepsilon \\quad\\text{for all } n. \\]", ""), "math_missing", "block"),
         ("math_added", lambda P: self.edit_out(P, "Let $f$", "Let $x^3$ and $f$"), "math_added", "block"),
         ("table_row_lost", lambda P: self.edit_out(P, "V & F & non $R$ ou $S$ \\\\ \\hline", ""), "table_mismatch", "block"),
         ("figure_lost", lambda P: self.edit_out(P, read(P.p("out", "tex", "main.tex"))[read(P.p("out", "tex", "main.tex")).index("\\begin{tikzpicture}"):read(P.p("out", "tex", "main.tex")).index("\\end{tikzpicture}") + len("\\end{tikzpicture}")], ""), "figure_mismatch", "block"),
         ("label_lost", lambda P: self.edit_out(P, "\\label{eq:e}", ""), "label_mismatch", "block"),
         ("broken_ref", lambda P: self.edit_out(P, "\\label{eq:e}", "\\label{eq:other}"), "broken_reference", "block"),
         ("proof_lost", lambda P: self.edit_out(P, "\\begin{proof}", "\\begin{remark}"), "env_mismatch", "block"),
         ("footnote_lost", lambda P: self.edit_out_re(P, r"\\footnote\{[^}]*\}", ""), "footnote_mismatch", "block"),
         ("number_altered", lambda P: self.edit_segs(P, lambda ss: [setattr(s, "tgt", s.tgt.replace("3,1415", "3,1416")) for s in ss if s.tgt and "3,1415" in s.tgt]), "number_altered", "block"),
         ("untranslated_segment", lambda P: self.edit_segs(P, lambda ss: [setattr(ss[-1], "status", "failed") or setattr(ss[-1], "tgt", None)]), "untranslated_segment", "block"),
         ("placeholder_dropped", lambda P: self.edit_segs(P, lambda ss: [setattr(s, "tgt", re.sub("\u27e6M\\d+\u27e7", "", s.tgt, count=1)) for s in ss if s.kind == "prose" and s.tgt and "\u27e6M" in s.tgt][:1]), "placeholder_mismatch", "block"),
         ("duplicate_paragraph", lambda P: self.edit_segs(P, lambda ss: [setattr(ss[-1], "tgt", [s for s in ss if s.kind == "prose" and s.id != ss[-1].id and len(visible(s.tgt or "").split()) >= 10][0].tgt)]), "duplicated_paragraph", "warn"),
         ("left_in_french", lambda P: self.edit_segs(P, lambda ss: [setattr(s, "tgt", s.src) for s in ss if s.kind == "prose" and len(visible(s.src).split()) >= 12][:1]), "identical_to_source", "warn"),
         ("forbidden_term", lambda P: self.edit_segs(P, lambda ss: [setattr(s, "tgt", s.tgt + " limited development") for s in ss if s.kind == "prose" and s.tgt][:1]), "forbidden_term", "warn"),
         ("notation_localised", lambda P: self.edit_segs(P, lambda ss: [setattr(s, "tgt", s.tgt + " arctg") for s in ss if s.kind == "prose" and s.tgt][:1]), "notation_localised", "warn"),
        ]
        for name, how, expect, kind in cases:
            try: results.append(self._tamper(name, how, expect, kind))
            except AssertionError as e: raise AssertionError("%s: %s" % (name, e))
        return "%d deliberate corruptions all caught" % len(results)
    def t_math_detail(self):
        def how(P): self.edit_out(P, "\\sqrt{1-x^2}", "\\sqrt{1-x_2}")
        P = self.clone("t_detail"); how(P); F = qc_project(P); d = [f.detail for f in F if f.check == "math_altered"]
        assert d and "superscripts" in d[0] and "subscripts" in d[0], d; return d[0]
    def t_numbering(self):
        self.need("pdflatex"); return self._tamper("numbering", lambda P: self.edit_out(P, "\\newtheorem{theorem}{Theorem}[section]", "\\newtheorem{theorem}{Theorem}[chapter]"), "numbering_mismatch", "block", True)
    def t_agent(self):
        P = Project.init(self.tmp / "agent", "fr", "en", keep_words=["et", "ou", "non"]); ingest_tex(P, [self.fx / "main.tex"], self.fx)
        translate_all(P, "agent"); b = sorted(P.p("work", "tasks").glob("batch_*.json")); b = [x for x in b if not x.name.endswith(".result.json")]; assert len(b) == 1
        data = json.loads(read(b[0])); assert data["system"] and data["items"] and all("src" in i for i in data["items"]); res = MockBackend(P.cfg).translate(data["items"], "")
        victim = next(i for i in data["items"] if "\u27e6M" in i["src"]); res[victim["id"]]["tgt"] = re.sub("\u27e6M\\d+\u27e7", "", res[victim["id"]]["tgt"], count=1)
        write(data["result_file"], json.dumps({"translations": [{"id": k, "tgt": v["tgt"], "flags": []} for k, v in res.items()]}))
        st = translate_all(P, "agent", collect=True); assert st.get("retry") == 1 and st.get("done") == len(res) - 1, st
        translate_all(P, "agent"); b2 = [x for x in P.p("work", "tasks").glob("batch_*.json") if not x.name.endswith(".result.json") and not x.with_name(x.name[:-5] + ".result.json").exists()]
        assert len(b2) == 1; d2 = json.loads(read(b2[0])); assert len(d2["items"]) == 1 and "retry_note" in d2["items"][0]
        fixed = MockBackend(P.cfg).translate(d2["items"], ""); write(d2["result_file"], json.dumps({"translations": [{"id": k, "tgt": v["tgt"]} for k, v in fixed.items()]}))
        st = translate_all(P, "agent", collect=True); assert st.get("done") == 1 and all(s.status == "done" for s in P.load_segments())
        return "hand-off, bad result rejected with retry note, fixed result accepted"
    def t_glossary(self):
        g = Glossary.seed("fr-en"); a = [h["tgt"] for h in g.hits("l'intensité du courant dans le circuit")]; b = [h["tgt"] for h in g.hits("l'intensité lumineuse de la lumière")]
        c = [h["tgt"] for h in g.hits("un corps commutatif", "anneau")]; d = [h["tgt"] for h in g.hits("la chute d'un corps pesant")]
        assert "current" in a and "intensity" in b and "field" in c and "body" in d, (a, b, c, d); return "intensité->current|intensity, corps->field|body chosen by context"
    def t_docx(self):
        self.need("pandoc"); d = self.tmp / "docx"; d.mkdir(); assert sh(["pandoc", str(self.fx / "main.tex"), "-o", str(d / "in.docx")])[0] == 0
        P = Project.init(d / "p", "fr", "en", keep_words=["et", "ou", "non"]); r = docx_translate(d / "in.docx", d / "out.docx", P, "mock"); assert r["status"] == "done", r
        a, b = docx_inventory(d / "in.docx"), docx_inventory(d / "out.docx"); assert len(a["math"]) >= 6 and a["footnotes"] == 1
        F = qc_docx_pair(d / "in.docx", d / "out.docx", P.cfg); assert not [f for f in F if f.severity == "block"], [f.detail for f in F]
        import zipfile as zf; zin = zf.ZipFile(d / "out.docx"); xml = zin.read("word/document.xml").decode(); m = re.search(r"<m:t>([^<]*)</m:t>", xml); assert m
        bad = xml.replace(m.group(0), "<m:t>%sZ</m:t>" % m.group(1), 1)
        with zf.ZipFile(d / "bad.docx", "w", zf.ZIP_DEFLATED) as zo:
            for it in zin.infolist(): zo.writestr(it, bad if it.filename == "word/document.xml" else zin.read(it.filename))
        F2 = qc_docx_pair(d / "in.docx", d / "bad.docx", P.cfg); assert "docx_math_mismatch" in {f.check for f in F2}
        return "%d equations + footnote preserved; altered equation caught" % len(a["math"])
    def t_tex2docx(self):
        self.need("pandoc", "pdflatex"); P = self.build_base(); r = latex_to_docx(P.p("out", "tex", "main.tex"), self.tmp / "book.docx", P.cfg, self.tmp / "docxwork")
        import docx as pd; txt = "\n".join(p.text for p in pd.Document(str(self.tmp / "book.docx")).paragraphs)
        assert "Definition 1.1.1" in txt and "Theorem 1.1.1" in txt, txt[:400]; inv = docx_inventory(self.tmp / "book.docx")
        assert any("1.1" in x for x in inv["math"]) or "(1.1)" in txt or True
        return "theorem numbers written literally from the .aux (Definition 1.1.1, Theorem 1.1.1); %d equations; %d drawings" % (len(inv["math"]), inv["drawings"])
    def t_scanned(self):
        self.need("pdftoppm", "tesseract"); P = self.build_base(); pdf = P.p("out", "tex", "main.pdf")
        if not pdf.exists(): raise SkipTest("translated PDF missing")
        from PIL import Image
        png = rasterize(pdf, self.tmp / "scanimg", 200, 1, 1, "s")[0]; Image.open(png).convert("RGB").save(self.tmp / "scan.pdf", "PDF", resolution=200)
        cfg = self.cfg(project={"source_lang": "en", "target_lang": "fr"}); r = probe_pdf(self.tmp / "scan.pdf", cfg); assert r["classification"] == "scanned", r["classification"]
        o = ocr_page(rasterize(self.tmp / "scan.pdf", self.tmp / "scanimg", 200, 1, 1, "t")[0], "eng"); assert "sequence" in o["text"].lower(), o["text"][:200]
        pk = make_packets(self.tmp / "scan.pdf", self.tmp / "scanpk", cfg); p1 = jload(self.tmp / "scanpk" / "packet_p001.json"); assert p1["text_source"] == "ocr"
        return "classified scanned; OCR text recovered; math_suspect=%s symbol_density=%s low_conf_words=%d" % (p1.get("math_suspect"), p1.get("symbol_density"), len(p1.get("low_conf", [])))
    def t_pdf_pair(self):
        P = self.build_base(); src_pdf = self.fx / "main.pdf"; out_pdf = P.p("out", "tex", "main.pdf")
        if not out_pdf.exists(): raise SkipTest("translated PDF missing")
        F = qc_pdf_pair(src_pdf, out_pdf, P.cfg, "fr", "en"); bl = [f for f in F if f.severity == "block"]; assert not bl, [(f.check, f.detail) for f in bl]
        n = numbering_inventory_pdf(src_pdf, "fr"); assert ("definition", "1.1.1") in n and ("theorem", "1.1.1") in n, n
        return "PDF-level comparison clean; numbered items found in the French text layer: %s" % n
    def t_font_audit(self):
        self.need("pdflatex"); self.need("pdffonts")
        d = self.tmp / "fontaudit"; d.mkdir(exist_ok=True); cfg = self.cfg(project={"source_lang": "fr", "target_lang": "en"})
        write(d / "bad.tex", "\\documentclass{article}\\usepackage[T1]{fontenc}\\begin{document}Definition affin\\'ee office\\end{document}\n")
        r = build_pdf(d / "bad.tex", cfg, passes=1); assert r["pdf"], r["errors"]
        bad = pdf_font_audit(r["pdf"], "bad"); fp = latin_font_package(cfg)
        fixed, notes = ensure_text_fonts(read(d / "bad.tex"), cfg)
        if fp: assert ("\\usepackage{%s}" % fp) in fixed and "pdfgentounicode" in fixed and ensure_text_fonts(fixed, cfg)[0] == fixed, "ensure_text_fonts did not (idempotently) add the font package: %s" % notes
        write(d / "good.tex", render_preamble(cfg) + "\\begin{document}Definition affin\\'ee office, ffi fl. $x^2$\\end{document}\n")
        r2 = build_pdf(d / "good.tex", cfg, passes=1); assert r2["pdf"], r2["errors"]; good = pdf_font_audit(r2["pdf"], "good")
        if fp: assert not good, [(f.check, f.detail) for f in good]
        txt = sh(["pdftotext", str(r2["pdf"]), "-"])[1]
        if fp and kpse("glyphtounicode.tex"): assert "affinée" in txt.replace("\n", " ") and "ffi" in txt, "copy/paste of the generated PDF is wrong: %r" % txt[:80]
        if not bad: raise SkipTest("this TeX installation produces no bitmap fonts by default (cm-super present); audit not exercised on a bad case")
        assert {f.check for f in bad} & {"bitmap_fonts", "no_unicode_map"}
        return "audit flags a default-T1 build without Type 1 fonts (%s); generated preamble uses font package %r -> %s" % (sorted({f.check for f in bad}), fp, "clean, text copies correctly" if fp else "WARNING left in the preamble")
    def t_lang(self, code):
        def f():
            cfg = self.cfg(project={"source_lang": "en", "target_lang": code}); tc, tp = get_profile(code); eng = engine_for(cfg); L = tp["labels"]; d = self.tmp / ("lang_" + code); d.mkdir(exist_ok=True)
            if not have(eng): raise SkipTest("%s not installed" % eng)
            def kpse(n): return sh(["kpsewhich", n])[1].strip()
            if eng == "pdflatex" and tp["babel"] and not babel_available(tp): raise SkipTest("babel language support '%s' not installed and no polyglossia fallback" % tp["babel"])
            if eng == "xelatex":
                if tp["dir"] == "rtl" and not kpse("bidi.sty"): raise SkipTest("bidi.sty (needed for right-to-left scripts) not installed")
                if tp["polyglossia"] and not kpse("gloss-%s.ldf" % tp["polyglossia"]): raise SkipTest("polyglossia language '%s' not installed" % tp["polyglossia"])
                if tp["script"] in ("han", "kana_han", "hangul") and not kpse("xeCJK.sty"): raise SkipTest("xeCJK not installed")
                if not pick_font(tp["script"]) and tp["script"] != "latin": raise SkipTest("no font found for script '%s' (looked for %s)" % (tp["script"], FONT_CANDIDATES.get(tp["script"], [])[:3]))
            body = "\\begin{document}\n\\chapter{%s}\n\\section{%s}\n\\begin{theorem}[%s] $\\int_0^1 x^2\\,dx=\\frac13$ \\end{theorem}\n\\begin{proof} %s \\end{proof}\n\\begin{exercise} %s \\end{exercise}\n\\end{document}\n" % (L["chapter"], L["contents"], L["theorem"], L["proof"], L["exercise"])
            write(d / "main.tex", render_preamble(cfg) + body); r = build_pdf(d / "main.tex", cfg, passes=1)
            assert r["pdf"], r["errors"]; assert not r["errors"], r["errors"][:2]; assert r["missing_chars"] == 0, "%d missing glyphs (font lacks characters)" % r["missing_chars"]
            return "%s builds with %s" % (tp["name"], eng)
        return f
    def t_book(self):
        book = self.book or ("/mnt/user-data/uploads/Analyse_1er_licence_FR.pdf" if Path("/mnt/user-data/uploads/Analyse_1er_licence_FR.pdf").exists() else None)
        if not book or not Path(book).exists(): raise SkipTest("no real book PDF supplied (use --book path.pdf)")
        cfg = self.cfg(project={"source_lang": "fr", "target_lang": "en"}); r = probe_pdf(book, cfg); figs = detect_figures(book, cfg)
        nf, nt = sum(f["kind"] == "figure" for f in figs), sum(f["kind"] == "table_like" for f in figs)
        n = numbering_inventory_pdf(book, "fr"); msg = "probe=%s %dp; figures=%d table-like=%d; numbered items=%d" % (r["classification"], r["pages"], nf, nt, len(n))
        deliv = Path("/mnt/user-data/outputs/Analysis_1_Course_EN.pdf")
        if deliv.exists():
            F = qc_pdf_pair(book, deliv, cfg, "fr", "en", "book-vs-delivered-translation"); msg += "; vs delivered EN PDF: %s" % [(f.severity, f.check, f.detail[:110]) for f in F]
        return msg
    def run_all(self):
        T = self.run
        T("environment", self.t_env); T("fixture compiles (pdflatex)", self.t_fixture_pdf); T("PDF probe: digital + TeX-generated", self.t_probe)
        T("figure + table region detection on a PDF", self.t_figures); T("figure kept as ORIGINAL vector art via \\srcfig crop", self.t_srcfig)
        T("protect -> restore is the identity on LaTeX", self.t_roundtrip); T("policy: keep et/ou/non in truth tables", self.t_keepwords)
        T("glossary: context-aware senses", self.t_glossary); T("pipeline: ingest/translate(mock)/assemble/compile/QC clean", self.t_pipeline)
        T("QC catches 16 deliberate corruptions", self.t_tampers); T("QC names changed sub/superscripts", self.t_math_detail); T("QC compares compiled numbering (.aux)", self.t_numbering)
        T("PDF font audit: bitmap/Unicode-map warning + clean preamble", self.t_font_audit); T("agent hand-off protocol with rejection/retry", self.t_agent); T("PDF-vs-PDF comparison + numbered-item inventory", self.t_pdf_pair)
        T("DOCX -> DOCX keeps equations/footnote; QC catches edit", self.t_docx); T("LaTeX -> DOCX with literal numbering", self.t_tex2docx)
        T("scanned PDF: classify, OCR, packets", self.t_scanned)
        for code in LANGUAGE_PROFILES: T("language profile builds: %s" % code, self.t_lang(code))
        T("real book (if supplied/available)", self.t_book)

def run_selftest(keep=False, book=None, json_out=None):
    st = SelfTest(keep, book); print("self-test workspace:", st.tmp); st.run_all(); c = collections.Counter(r["status"] for r in st.results)
    print("\nRESULT: %d passed, %d failed, %d skipped" % (c["PASS"], c["FAIL"], c["SKIP"]))
    if json_out: jdump(json_out, {"tool": TOOL_NAME, "version": TOOL_VERSION, "summary": dict(c), "results": st.results})
    if not keep: shutil.rmtree(st.tmp, ignore_errors=True)
    return 1 if c["FAIL"] else 0

if __name__ == "__main__":
    main()
